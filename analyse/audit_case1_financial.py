"""Read-only paired Case 1 audit. No model requests or workbook writes."""
import collections
import glob
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / 'benchmarks/results'
NAMES = {
    'financial': 'spreadsheetbench-v1-harness-representative-200-seed42-planner-contract-rerun-20260914b',
    'basic': 'spreadsheetbench-v1-harness-basic-representative-200-seed42-planner-contract-rerun-20260914b',
}

def rows(path):
    return [r for f in sorted(path.glob('workers/*/results.json')) for r in json.loads(f.read_text())]

def cohorts():
    groups = {}
    for arm, name in NAMES.items():
        original = rows(RESULTS / name)
        result = {r['task_id']: r for r in original}
        assert len(result) == len(original) == 200
        for retry in rows(RESULTS / 'planner-contract-rerun-retries-20260914' / arm):
            old = result[retry['task_id']]
            if old.get('error_type') == 'ProviderError' and retry.get('status') == 'completed':
                result[retry['task_id']] = retry
        groups[arm] = result
    return groups

def case_pass(r):
    cases = r.get('case_results', [])
    return cases[0].get('passed') if cases and cases[0].get('status') == 'scored' else None

def trajectory(r):
    path = Path(r['run_dir']) / 'case-1/trajectory.jsonl'
    return [json.loads(line) for line in path.read_text().splitlines() if line.strip()]

def main():
    import sys
    if '--cells' in sys.argv:
        return cells()
    groups = cohorts()
    summary = {}
    for arm, group in groups.items():
        scored = [r for r in group.values() if r.get('status') == 'completed']
        summary[arm] = dict(scored=len(scored), soft=sum(r['soft'] for r in scored)/len(scored),
                            hard=sum(r['hard'] for r in scored)/len(scored))
    paired = collections.defaultdict(list)
    for tid in sorted(groups['financial']):
        f, b = (case_pass(groups[a][tid]) for a in ('financial', 'basic'))
        paired[f'financial={f},basic={b}'].append(tid)
    print(json.dumps({'summary': summary, 'pairs': dict(paired)}, indent=2))
    for tid in paired['financial=False,basic=True']:
        print('\nTASK', tid)
        for arm in ('financial', 'basic'):
            r = groups[arm][tid]
            events = trajectory(r)
            print(arm, 'run', r['run_dir'])
            print('stages', [s.get('name') for s in r.get('agent', {}).get('stages', [])])
            for event in events:
                name, payload = event['event'], event.get('payload', {})
                if name == 'model.responded' and payload.get('stage') == 'plan':
                    print('PLAN', payload.get('text'))
                elif name.startswith('harness.') and name not in ('harness.plugin.activated', 'harness.composition.resolved', 'harness.skills.routed'):
                    print(name, json.dumps(payload, ensure_ascii=False)[:5000])

def cells():
    import sys
    from openpyxl import load_workbook
    from spreadsheet_harness.spreadsheetbench_v1 import load_spreadsheetbench_v1, _official_cells, _official_equal
    groups = cohorts()
    tasks = {t.task_id: t for t in load_spreadsheetbench_v1(ROOT / 'benchmarks/data/spreadsheetbench_912_v0.1')}
    for tid in sorted(groups['financial']):
        f, b = groups['financial'][tid], groups['basic'][tid]
        reverse = '--reverse' in sys.argv
        if case_pass(f) is not (True if reverse else False) or case_pass(b) is not (False if reverse else True):
            continue
        if '--task' in sys.argv and tid != sys.argv[sys.argv.index('--task') + 1]:
            continue
        task = tasks[tid]
        print('\nTASK', tid, task.instruction, '\nANSWER', task.answer_position)
        paths = {'input': task.cases[0].input_path, 'gold': task.cases[0].golden_path,
                 'financial': Path(f['run_dir']) / 'case-1/artifacts/output.xlsx',
                 'basic': Path(b['run_dir']) / 'case-1/artifacts/output.xlsx'}
        books = {k: load_workbook(v, data_only=True) for k, v in paths.items()}
        forms = {k: load_workbook(v, data_only=False) for k, v in paths.items()}
        failed, unchanged, total = [], 0, 0
        for raw in task.answer_position.split(','):
            if '!' in raw:
                sheet, region = raw.split('!')
                sheet = sheet.strip("'")
            else:
                sheet, region = books['gold'].sheetnames[0], raw
            for cell in _official_cells(region.strip("'")):
                total += 1
                vals = {k: w[sheet][cell].value if sheet in w.sheetnames else '<missing sheet>' for k,w in books.items()}
                losing_arm = 'basic' if reverse else 'financial'
                if not _official_equal(vals['gold'], vals[losing_arm]):
                    unchanged += int(vals[losing_arm] == vals['input'])
                    if len(failed) < (2 if '--brief' in sys.argv else 12):
                        failed.append({'cell': sheet+'!'+cell, 'values': vals,
                                       'formulas': {k: w[sheet][cell].value if sheet in w.sheetnames else None for k,w in forms.items()}})
        print('EVALUATED', total, 'FAILED_UNCHANGED', unchanged)
        for item in failed:
            print(json.dumps(item, ensure_ascii=False, default=str))
        for w in [*books.values(), *forms.values()]:
            w.close()

if __name__ == '__main__':
    main()
