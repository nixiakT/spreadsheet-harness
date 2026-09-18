"""Offline gate audit; rules never read gold, labels come from prior scoring."""
import collections
import json
from pathlib import Path
from openpyxl import load_workbook
from audit_case1_financial import cohorts, case_pass, trajectory


def main():
    for arm, records in cohorts().items():
        counts = collections.Counter()
        examples = []
        for tid, r in sorted(records.items()):
            passed = case_pass(r)
            if passed is None:
                continue
            events = trajectory(r)
            execute = any(e['event'] == 'model.requested' and e['payload'].get('stage') == 'execute' for e in events)
            actions = [a for e in events if e['event'] == 'harness.planner_actions.verified' for a in e['payload'].get('actions', [])]
            if execute or not actions:
                continue
            book = load_workbook(Path(r['run_dir']) / 'case-1/artifacts/output.xlsx', data_only=True)
            errors = []
            for action in actions:
                sheet, cell = action.get('sheet'), action.get('target')
                if sheet in book.sheetnames and isinstance(cell, str) and ':' not in cell:
                    current = book[sheet][cell]
                    if current.data_type == 'e':
                        errors.append([sheet, cell, current.value])
            book.close()
            counts[f'all_no_executor_verified_{passed}'] += 1
            if errors:
                counts[f'gate_flagged_{passed}'] += 1
                examples.append({'task_id': tid, 'official_pass': passed, 'error_count': len(errors), 'example': errors[:2]})
        print(arm, json.dumps(counts), flush=True)
        print(json.dumps(examples, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
