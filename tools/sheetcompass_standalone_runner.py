"""Standalone SheetCompass-style Explorer/Programmer/Reflector runner.
This intentionally does not import spreadsheet_harness.
"""
import json, os, shutil
from pathlib import Path
from openpyxl import load_workbook
from openai import OpenAI

root=Path(os.environ.get('SC_ROOT','benchmarks/results/sheetcompass-standalone-qwen36-90-20260906')); root.mkdir(parents=True,exist_ok=True)
client=OpenAI(base_url=os.environ.get('OPENAI_BASE_URL','http://47.96.153.159:8010/v1'),api_key=os.environ['OPENAI_API_KEY'],timeout=180.0,max_retries=2)
limit=int(os.environ.get('SC_LIMIT','30'))
for cat in ('Debugging','Financial_Model','Template'):
 data=json.loads((Path('benchmarks/data/spreadsheetbench-v2')/cat/'dataset.json').read_text())[:limit]
 for row in data:
  try:
   tid=f'{cat}/{row["id"]}'; out=root/tid.replace('/','_'); out.mkdir(parents=True,exist_ok=True)
   src=Path('benchmarks/data/spreadsheetbench-v2')/cat/row['spreadsheet_path']; dst=out/'output.xlsx'; shutil.copy2(src,dst)
   wb=load_workbook(dst, data_only=False); graph={'sheets':[]}
   for ws in wb.worksheets:
    graph['sheets'].append({'name':ws.title,'dimension':ws.calculate_dimension(),'nodes':[{'cell':c.coordinate,'value':c.value} for rowc in ws.iter_rows(max_row=min(ws.max_row,30),max_col=min(ws.max_column,20)) for c in rowc if c.value is not None]})
   prompt=f'You are SheetCompass Explorer/Programmer/Reflector. Build graph relations, execute the task, then reflect. Return ONLY compact JSON edits. Task: {row.get("instruction", row.get("task", ""))}\nGRAPH:{json.dumps(graph, default=str)[:6000]}\nJSON: {{"edits":[{{"cell":"A1","value":"..."}}]}}'
   r=client.chat.completions.create(model='qwen36-35b-a3b',messages=[{'role':'user','content':prompt}],temperature=1,max_tokens=2048,extra_body={'enable_thinking':False})
   text=r.choices[0].message.content
   try:
    payload=json.loads(text[text.find('{'):text.rfind('}')+1])
    edits=payload if isinstance(payload,list) else payload.get('edits',[])
    for edit in edits:
     cell=edit.get('cell') if isinstance(edit,dict) else None
     if cell and isinstance(edit,dict) and edit.get('value') is not None: wb.active[cell]=edit['value']
    wb.save(dst)
   except Exception:
    pass
   (out/'trajectory.json').write_text(json.dumps({'task_id':tid,'graph':graph,'response':text},ensure_ascii=False,indent=2,default=str))
  except Exception as e:
   try: (out/'error.txt').write_text(repr(e))
   except Exception: pass
  print(f'processed {cat}/{row["id"]}', flush=True)
