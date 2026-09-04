"""High-precision, auditable repairs for local spreadsheet formula patterns."""

from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from typing import Any

from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.formula.translate import Translator, TranslatorError


@dataclass(frozen=True)
class FormulaPatternRepair:
    """A formula replacement independently implied by adjacent formulas."""

    sheet: str
    cell: str
    current: str
    replacement: str
    directions: tuple[str, ...]
    neighbors: tuple[str, ...]

    def to_dict(self) -> dict[str, Any]:
        return {
            "kind": "formula_pattern_conflict",
            "sheet": self.sheet,
            "cell": self.cell,
            "current": self.current,
            "neighbor_derived": self.replacement,
            "directions": list(self.directions),
            "neighbors": list(self.neighbors),
        }


def _translated_pair(
    first: Any,
    target: Any,
    second: Any,
) -> str | None:
    if not all(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for cell in (first, target, second)
    ):
        return None
    try:
        from_first = Translator(
            first.value, origin=first.coordinate
        ).translate_formula(target.coordinate)
        from_second = Translator(
            second.value, origin=second.coordinate
        ).translate_formula(target.coordinate)
    except (TokenizerError, TranslatorError, TypeError, ValueError):
        return None
    if from_first != from_second or target.value == from_first:
        return None
    return from_first


def detect_formula_pattern_repairs(
    workbook: Any,
    *,
    sheet_names: Iterable[str] | None = None,
    max_rows: int = 500,
    max_columns: int = 100,
) -> list[FormulaPatternRepair]:
    """Find formula conflicts supported by two opposing translated neighbors.

    A proposal is emitted only when both neighbors translate to the exact same
    formula at the target coordinate. Conflicting horizontal and vertical
    proposals for the same cell are discarded.
    """

    worksheet_names = {worksheet.title for worksheet in workbook.worksheets}
    selected_names = (
        [name for name in sheet_names if name in worksheet_names]
        if sheet_names is not None
        else [worksheet.title for worksheet in workbook.worksheets]
    )
    repairs: list[FormulaPatternRepair] = []
    for sheet_name in selected_names:
        if sheet_name not in worksheet_names:
            continue
        worksheet = workbook[sheet_name]
        row_limit = min(int(worksheet.max_row or 0), max_rows)
        column_limit = min(int(worksheet.max_column or 0), max_columns)
        proposals: dict[str, list[tuple[str, str, tuple[str, str]]]] = {}
        populated_cells = list(getattr(worksheet, "_cells", {}).values())
        for cell in populated_cells:
            row_number = int(cell.row)
            column_number = int(cell.column)
            if (
                row_number > row_limit
                or column_number > column_limit
                or not isinstance(cell.value, str)
                or not cell.value.startswith("=")
            ):
                continue
            if 1 < column_number < column_limit:
                left = worksheet.cell(row_number, column_number - 1)
                right = worksheet.cell(row_number, column_number + 1)
                replacement = _translated_pair(left, cell, right)
                if replacement is not None:
                    proposals.setdefault(cell.coordinate, []).append(
                        ("horizontal", replacement, (left.coordinate, right.coordinate))
                    )
            if 1 < row_number < row_limit:
                above = worksheet.cell(row_number - 1, column_number)
                below = worksheet.cell(row_number + 1, column_number)
                replacement = _translated_pair(above, cell, below)
                if replacement is not None:
                    proposals.setdefault(cell.coordinate, []).append(
                        ("vertical", replacement, (above.coordinate, below.coordinate))
                    )
        for coordinate, items in proposals.items():
            replacements = {item[1] for item in items}
            if len(replacements) != 1:
                continue
            repairs.append(
                FormulaPatternRepair(
                    sheet=sheet_name,
                    cell=coordinate,
                    current=str(worksheet[coordinate].value),
                    replacement=replacements.pop(),
                    directions=tuple(sorted({item[0] for item in items})),
                    neighbors=tuple(
                        dict.fromkeys(neighbor for item in items for neighbor in item[2])
                    ),
                )
            )
    return sorted(repairs, key=lambda item: (item.sheet.casefold(), item.cell))


def select_safe_formula_pattern_repairs(
    repairs: Iterable[FormulaPatternRepair],
    *,
    task_hint: str = "",
) -> list[FormulaPatternRepair]:
    """Keep only proposals strong enough for deterministic automatic repair."""

    candidates = [
        repair
        for repair in repairs
        if len(repair.replacement) > 1
        and repair.replacement.startswith("=")
        and "#REF!" not in repair.replacement.upper()
    ]
    if len(candidates) == 1:
        return candidates
    normalized_hint = task_hint.casefold().replace("_", " ")
    if "sign convention" in normalized_hint:
        def without_leading_sign(formula: str) -> str:
            body = formula[1:] if formula.startswith("=") else formula
            return body[1:] if body.startswith(("+", "-")) else body

        sign_candidates = [
            repair
            for repair in candidates
            if without_leading_sign(repair.current)
            == without_leading_sign(repair.replacement)
        ]
        if len(sign_candidates) == 1:
            return sign_candidates
    return [
        repair
        for repair in candidates
        if len(repair.directions) >= 2 or "#REF!" in repair.current.upper()
    ]


__all__ = [
    "FormulaPatternRepair",
    "detect_formula_pattern_repairs",
    "select_safe_formula_pattern_repairs",
]
