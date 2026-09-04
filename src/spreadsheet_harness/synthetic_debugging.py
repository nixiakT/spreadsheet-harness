"""Build and score exact-answer debugging cases from unlabeled real workbooks."""

from __future__ import annotations

import hashlib
import json
import random
import shutil
from collections.abc import Iterable
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile, ZipInfo

from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.formula.translate import Translator, TranslatorError

from .benchmark import _atomic_write_json, _sha256
from .errors import HarnessError
from .formula_patterns import (
    detect_formula_pattern_repairs,
    select_safe_formula_pattern_repairs,
)
from .openpyxl_compat import load_workbook

SYNTHETIC_DEBUGGING_SCHEMA = "synthetic-debugging-corpus-v1"
DEFAULT_INSTRUCTION = (
    "Please audit and fix this file thoroughly. Errors may include formula logic errors or "
    "inconsistencies, broken cell references, and potential calculation errors. Preserve all "
    "unrelated workbook content."
)


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _safe_xlsx_members(archive: ZipFile, *, max_member_bytes: int) -> list[ZipInfo]:
    members: list[ZipInfo] = []
    for member in archive.infolist():
        path = Path(member.filename)
        if (
            member.is_dir()
            or path.suffix.casefold() != ".xlsx"
            or path.is_absolute()
            or ".." in path.parts
            or member.file_size <= 0
            or member.file_size > max_member_bytes
        ):
            continue
        members.append(member)
    return members


def _metadata_by_filename(archive: ZipFile) -> dict[str, dict[str, Any]]:
    try:
        raw = archive.read("xlsx_metadata.json")
    except KeyError:
        return {}
    try:
        document = json.loads(raw)
    except (UnicodeDecodeError, json.JSONDecodeError):
        return {}
    files = document.get("files", []) if isinstance(document, dict) else []
    if not isinstance(files, list):
        return {}
    return {
        str(item["filename"]): dict(item)
        for item in files
        if isinstance(item, dict) and isinstance(item.get("filename"), str)
    }


def _coherent_formula_targets(
    workbook: Any,
    *,
    max_rows: int,
    max_columns: int,
) -> list[tuple[str, str, str, tuple[str, ...]]]:
    targets: list[tuple[str, str, str, tuple[str, ...]]] = []
    for worksheet in workbook.worksheets:
        row_limit = min(int(worksheet.max_row or 0), max_rows)
        column_limit = min(int(worksheet.max_column or 0), max_columns)
        for cell in list(getattr(worksheet, "_cells", {}).values()):
            row_number = int(cell.row)
            column_number = int(cell.column)
            formula = cell.value
            if (
                row_number > row_limit
                or column_number > column_limit
                or not isinstance(formula, str)
                or not formula.startswith("=")
                or formula == "=0"
            ):
                continue
            directions: list[str] = []
            if 1 < column_number < column_limit:
                left = worksheet.cell(row_number, column_number - 1)
                right = worksheet.cell(row_number, column_number + 1)
                if all(
                    isinstance(item.value, str) and item.value.startswith("=")
                    for item in (left, right)
                ):
                    try:
                        from_left = Translator(
                            left.value, origin=left.coordinate
                        ).translate_formula(cell.coordinate)
                        from_right = Translator(
                            right.value, origin=right.coordinate
                        ).translate_formula(cell.coordinate)
                    except (TokenizerError, TranslatorError, TypeError, ValueError):
                        pass
                    else:
                        if from_left == from_right == formula:
                            directions.append("horizontal")
            if 1 < row_number < row_limit:
                above = worksheet.cell(row_number - 1, column_number)
                below = worksheet.cell(row_number + 1, column_number)
                if all(
                    isinstance(item.value, str) and item.value.startswith("=")
                    for item in (above, below)
                ):
                    try:
                        from_above = Translator(
                            above.value, origin=above.coordinate
                        ).translate_formula(cell.coordinate)
                        from_below = Translator(
                            below.value, origin=below.coordinate
                        ).translate_formula(cell.coordinate)
                    except (TokenizerError, TranslatorError, TypeError, ValueError):
                        pass
                    else:
                        if from_above == from_below == formula:
                            directions.append("vertical")
            if directions:
                targets.append(
                    (worksheet.title, cell.coordinate, formula, tuple(directions))
                )
    return targets


def build_synthetic_debugging_corpus(
    archive_path: str | Path,
    output_dir: str | Path,
    *,
    limit: int = 100,
    seed: int = 20_260_825,
    max_member_bytes: int = 25 * 1024 * 1024,
    max_rows: int = 500,
    max_columns: int = 100,
) -> dict[str, Any]:
    """Create formula-corruption cases whose original workbook is the golden answer."""

    if limit <= 0:
        raise ValueError("limit must be positive")
    source = Path(archive_path).expanduser().resolve(strict=True)
    output = Path(output_dir).expanduser().resolve()
    if output.exists():
        raise HarnessError(f"Synthetic corpus output already exists: {output}")
    output.mkdir(parents=True)
    cases_dir = output / "cases"
    cases_dir.mkdir()
    rows: list[dict[str, Any]] = []
    skipped: dict[str, int] = {}
    try:
        with ZipFile(source) as archive:
            metadata = _metadata_by_filename(archive)
            members = _safe_xlsx_members(archive, max_member_bytes=max_member_bytes)
            random.Random(seed).shuffle(members)
            for member in members:
                if len(rows) >= limit:
                    break
                try:
                    raw = archive.read(member)
                    workbook = load_workbook(
                        BytesIO(raw), data_only=False, read_only=False, keep_links=False
                    )
                except Exception:
                    skipped["unreadable"] = skipped.get("unreadable", 0) + 1
                    continue
                try:
                    existing = detect_formula_pattern_repairs(
                        workbook, max_rows=max_rows, max_columns=max_columns
                    )
                    if existing:
                        skipped["preexisting_conflict"] = (
                            skipped.get("preexisting_conflict", 0) + 1
                        )
                        continue
                    targets = _coherent_formula_targets(
                        workbook, max_rows=max_rows, max_columns=max_columns
                    )
                    if not targets:
                        skipped["no_supported_formula"] = (
                            skipped.get("no_supported_formula", 0) + 1
                        )
                        continue
                    target = targets[0]
                    sheet_name, coordinate, original_formula, directions = target
                    workbook[sheet_name][coordinate] = "=0"
                    injected = select_safe_formula_pattern_repairs(
                        detect_formula_pattern_repairs(
                            workbook, max_rows=max_rows, max_columns=max_columns
                        )
                    )
                    if len(injected) != 1 or (
                        injected[0].sheet,
                        injected[0].cell,
                        injected[0].replacement,
                    ) != (sheet_name, coordinate, original_formula):
                        skipped["not_uniquely_recoverable"] = (
                            skipped.get("not_uniquely_recoverable", 0) + 1
                        )
                        continue
                    case_id = f"formula_pattern_{len(rows) + 1:04d}"
                    case_dir = cases_dir / case_id
                    case_dir.mkdir()
                    golden_path = case_dir / "golden.xlsx"
                    input_path = case_dir / "input.xlsx"
                    golden_path.write_bytes(raw)
                    workbook.save(input_path)
                    source_metadata = metadata.get(member.filename, {})
                    rows.append(
                        {
                            "id": case_id,
                            "instruction": DEFAULT_INSTRUCTION,
                            "input_path": str(input_path.relative_to(output)),
                            "golden_path": str(golden_path.relative_to(output)),
                            "source": {
                                "filename": member.filename,
                                "sha256": _sha256_bytes(raw),
                                "title": source_metadata.get("title"),
                                "snippet": source_metadata.get("snippet"),
                                "query": source_metadata.get("query"),
                            },
                            "mutation": {
                                "kind": "formula_pattern_conflict",
                                "sheet": sheet_name,
                                "cell": coordinate,
                                "original_formula": original_formula,
                                "injected_formula": "=0",
                                "support_directions": list(directions),
                            },
                        }
                    )
                finally:
                    workbook.close()
    except BadZipFile as exc:
        shutil.rmtree(output)
        raise HarnessError(f"Invalid synthetic workbook archive: {source}") from exc
    if not rows:
        shutil.rmtree(output)
        raise HarnessError("No uniquely recoverable formula-pattern cases were found")
    manifest = {
        "schema_version": SYNTHETIC_DEBUGGING_SCHEMA,
        "archive": {
            "path": str(source),
            "sha256": _sha256(source),
        },
        "generation": {
            "seed": seed,
            "requested_limit": limit,
            "case_count": len(rows),
            "max_member_bytes": max_member_bytes,
            "max_rows": max_rows,
            "max_columns": max_columns,
            "skipped": skipped,
        },
        "cases": rows,
    }
    _atomic_write_json(output / "manifest.json", manifest)
    return manifest


def load_synthetic_debugging_manifest(path: str | Path) -> tuple[Path, dict[str, Any]]:
    manifest_path = Path(path).expanduser().resolve(strict=True)
    document = json.loads(manifest_path.read_text(encoding="utf-8"))
    if not isinstance(document, dict) or document.get("schema_version") != SYNTHETIC_DEBUGGING_SCHEMA:
        raise HarnessError("Unsupported synthetic debugging manifest")
    cases = document.get("cases")
    if not isinstance(cases, list) or not cases:
        raise HarnessError("Synthetic debugging manifest has no cases")
    return manifest_path.parent, document


def score_synthetic_debugging_outputs(
    manifest_path: str | Path,
    output_paths: Iterable[str | Path],
) -> dict[str, Any]:
    """Score target formulas exactly and report cell-level unrelated regressions."""

    root, manifest = load_synthetic_debugging_manifest(manifest_path)
    resolved = [Path(path).expanduser().resolve(strict=True) for path in output_paths]
    outputs = [
        workbook
        for path in resolved
        for workbook in (
            sorted(path.glob("*_output.xlsx")) if path.is_dir() else [path]
        )
    ]
    by_id = {path.stem.removesuffix("_output"): path for path in outputs}
    results: list[dict[str, Any]] = []
    for row in manifest["cases"]:
        case_id = str(row["id"])
        output = by_id.get(case_id)
        if output is None:
            results.append({"id": case_id, "scored": False, "reason": "missing_output"})
            continue
        baseline = load_workbook(root / row["input_path"], data_only=False, read_only=False)
        golden = load_workbook(root / row["golden_path"], data_only=False, read_only=False)
        candidate = load_workbook(output, data_only=False, read_only=False)
        try:
            mutation = row["mutation"]
            sheet_name = str(mutation["sheet"])
            coordinate = str(mutation["cell"])
            modification_correct = (
                sheet_name in candidate.sheetnames
                and candidate[sheet_name][coordinate].value
                == golden[sheet_name][coordinate].value
            )
            regressions = 0
            compared = 0
            if candidate.sheetnames != baseline.sheetnames:
                regressions += 1
            for worksheet in baseline.worksheets:
                if worksheet.title not in candidate.sheetnames:
                    regressions += 1
                    continue
                other = candidate[worksheet.title]
                max_row = max(int(worksheet.max_row or 0), int(other.max_row or 0))
                max_column = max(int(worksheet.max_column or 0), int(other.max_column or 0))
                for row_number in range(1, max_row + 1):
                    for column_number in range(1, max_column + 1):
                        left = worksheet.cell(row_number, column_number)
                        right = other.cell(row_number, column_number)
                        if worksheet.title == sheet_name and left.coordinate == coordinate:
                            continue
                        compared += 1
                        if left.value != right.value:
                            regressions += 1
            results.append(
                {
                    "id": case_id,
                    "scored": True,
                    "modification_correct": modification_correct,
                    "regression_cells": regressions,
                    "compared_regression_cells": compared,
                    "passed": modification_correct and regressions == 0,
                }
            )
        finally:
            baseline.close()
            golden.close()
            candidate.close()
    scored = [row for row in results if row.get("scored")]
    return {
        "schema_version": "synthetic-debugging-score-v1",
        "case_count": len(results),
        "scored_count": len(scored),
        "passed_count": sum(bool(row.get("passed")) for row in scored),
        "results": results,
    }


def run_deterministic_debugging_repair(
    manifest_path: str | Path,
    output_dir: str | Path,
) -> dict[str, Any]:
    """Run the production high-confidence formula repair policy on synthetic cases."""

    root, manifest = load_synthetic_debugging_manifest(manifest_path)
    output = Path(output_dir).expanduser().resolve()
    if output.exists():
        raise HarnessError(f"Synthetic repair output already exists: {output}")
    output.mkdir(parents=True)
    results: list[dict[str, Any]] = []
    for row in manifest["cases"]:
        case_id = str(row["id"])
        workbook = load_workbook(root / row["input_path"], data_only=False, read_only=False)
        try:
            repairs = select_safe_formula_pattern_repairs(
                detect_formula_pattern_repairs(workbook)
            )
            for repair in repairs:
                cell = workbook[repair.sheet][repair.cell]
                if cell.value == repair.current:
                    cell.value = repair.replacement
            destination = output / f"{case_id}_output.xlsx"
            workbook.save(destination)
            results.append(
                {
                    "id": case_id,
                    "repair_count": len(repairs),
                    "output": str(destination),
                }
            )
        finally:
            workbook.close()
    return {
        "schema_version": "synthetic-debugging-deterministic-run-v1",
        "case_count": len(results),
        "results": results,
    }


__all__ = [
    "SYNTHETIC_DEBUGGING_SCHEMA",
    "build_synthetic_debugging_corpus",
    "load_synthetic_debugging_manifest",
    "run_deterministic_debugging_repair",
    "score_synthetic_debugging_outputs",
]
