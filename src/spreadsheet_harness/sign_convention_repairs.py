"""High-confidence financial sign-convention repairs.

The rules in this module are deliberately semantic and task gated.  They require both an
explicit sign-convention hint and a recognizable financial row identity before proposing a
formula change.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

_LOCAL_CELL = r"\$?[A-Z]{1,3}\$?\d+"
_MULTIPLY_GROWTH_RE = re.compile(
    rf"^=(?P<base>{_LOCAL_CELL})\*\(1-(?P<rate>{_LOCAL_CELL})\)$",
    re.IGNORECASE,
)
_NEGATIVE_PRODUCT_RE = re.compile(
    rf"^=-(?P<first>{_LOCAL_CELL})\*(?P<second>{_LOCAL_CELL})$",
    re.IGNORECASE,
)
_BINARY_ADD_RE = re.compile(
    rf"^=(?P<first>{_LOCAL_CELL})\+(?P<second>{_LOCAL_CELL})$",
    re.IGNORECASE,
)
_AFTER_TAX_RE = re.compile(
    rf"^=(?P<cost>{_LOCAL_CELL})\*\(1\+(?P<tax>{_LOCAL_CELL})\)$",
    re.IGNORECASE,
)
_BINARY_SUBTRACT_RE = re.compile(
    rf"^=(?P<first>{_LOCAL_CELL})-(?P<second>{_LOCAL_CELL})$",
    re.IGNORECASE,
)
_ADDITIVE_LOCAL_REFERENCE_RE = re.compile(
    rf"(?P<operator>[+-])(?P<reference>{_LOCAL_CELL})",
    re.IGNORECASE,
)
_ONE_PLUS_LOCAL_REFERENCE_RE = re.compile(
    rf"\(1\+(?P<reference>{_LOCAL_CELL})\)",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class SignConventionRepair:
    sheet: str
    cell: str
    current: str
    replacement: str
    kind: str
    evidence: tuple[str, ...]

    def to_dict(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet,
            "target": self.cell,
            "current": self.current,
            "formula": self.replacement,
            "kind": self.kind,
            "evidence": list(self.evidence),
        }


def _normalized(value: Any) -> str:
    return " ".join(re.sub(r"[^a-z0-9%]+", " ", str(value or "").casefold()).split())


def _row_label(worksheet: Any, row: int, before_column: int) -> str:
    return _normalized(_raw_row_label(worksheet, row, before_column))


def _raw_row_label(worksheet: Any, row: int, before_column: int) -> str:
    for column in range(min(before_column - 1, 8), 0, -1):
        value = worksheet.cell(row, column).value
        if value is not None and not (isinstance(value, str) and value.startswith("=")):
            return str(value).strip()
    return ""


def _negative_line_item_formula(value: Any) -> bool:
    if not isinstance(value, str) or not value.startswith("="):
        return False
    compact = re.sub(r"\s+", "", value).upper()
    return compact.startswith("=-") or compact.endswith("*-1")


def _is_additive_subtotal_label(label: str) -> bool:
    """Return whether a row label denotes an additive financial subtotal.

    A negative source line is not sufficient evidence on its own: roll-forward rows may
    intentionally subtract an already-negative amount (for example, deriving beginning PP&E
    from ending PP&E and D&A).  Restrict the rewrite to rows whose public label identifies a
    conventional profit, earnings, cash-flow, margin, or explicit total/subtotal measure.
    """

    return any(
        marker in label
        for marker in (
            "ebit",
            "income",
            "profit",
            "earnings",
            "cash flow",
            "margin",
            "subtotal",
            "total",
        )
    )


def _add_already_negative_components(
    worksheet: Any,
    formula: str,
    *,
    target_column: int,
) -> tuple[str, tuple[str, ...]]:
    evidence: list[str] = []

    def replace(match: re.Match[str]) -> str:
        if match.group("operator") != "-":
            return match.group(0)
        # Only rewrite binary subtraction in an additive expression.  The same token pattern
        # also matches a leading unary minus (``=-A1``) and a negated denominator
        # (``=A1/-A2``); both carry different semantics and must remain untouched.
        if match.start() == 1 or formula[match.start() - 1] in "*/^(,":
            return match.group(0)
        reference = match.group("reference")
        referenced_row, referenced_column = _reference_parts(reference)
        raw_reference_label = _raw_row_label(
            worksheet,
            referenced_row,
            target_column,
        )
        referenced_value = worksheet.cell(referenced_row, referenced_column).value
        if not re.match(r"^\s*\(-\)", raw_reference_label) or not (
            _negative_line_item_formula(referenced_value)
        ):
            return match.group(0)
        evidence.append(
            f"{worksheet.cell(referenced_row, referenced_column).coordinate}:"
            f"{raw_reference_label}"
        )
        return f"+{reference}"

    return _ADDITIVE_LOCAL_REFERENCE_RE.sub(replace, formula), tuple(evidence)


def _subtract_transaction_cost_components(
    worksheet: Any,
    formula: str,
    *,
    target_column: int,
) -> tuple[str, tuple[str, ...]]:
    evidence: list[str] = []

    def replace(match: re.Match[str]) -> str:
        if match.group("operator") != "+":
            return match.group(0)
        reference = match.group("reference")
        referenced_row, referenced_column = _reference_parts(reference)
        reference_label = _row_label(
            worksheet,
            referenced_row,
            target_column,
        )
        if "transaction cost" not in reference_label:
            return match.group(0)
        evidence.append(
            f"{worksheet.cell(referenced_row, referenced_column).coordinate}:"
            f"{reference_label}"
        )
        return f"-{reference}"

    return _ADDITIVE_LOCAL_REFERENCE_RE.sub(replace, formula), tuple(evidence)


def _reference_parts(reference: str) -> tuple[int, int]:
    coordinate = reference.replace("$", "")
    column, row = coordinate_from_string(coordinate)
    return int(row), column_index_from_string(column)


def _same_cell(reference: str, *, row: int, column: int) -> bool:
    return _reference_parts(reference) == (row, column)


def _nearby_labeled_row(
    worksheet: Any,
    *,
    target_row: int,
    target_column: int,
    phrases: tuple[str, ...],
    lookback: int = 6,
) -> int | None:
    for row in range(target_row - 1, max(0, target_row - lookback - 1), -1):
        label = _row_label(worksheet, row, target_column)
        if any(phrase in label for phrase in phrases):
            return row
    return None


def detect_sign_convention_repairs(
    workbook: Any,
    *,
    task_hint: str,
) -> list[SignConventionRepair]:
    """Detect financial sign errors supported by labels and exact formula structure."""

    if "sign convention" not in _normalized(task_hint):
        return []

    repairs: list[SignConventionRepair] = []
    for worksheet in workbook.worksheets:
        for cell in list(getattr(worksheet, "_cells", {}).values()):
            formula = cell.value
            if not isinstance(formula, str) or not formula.startswith("="):
                continue
            label = _row_label(worksheet, int(cell.row), int(cell.column))

            if "revenue" in label and (match := _MULTIPLY_GROWTH_RE.fullmatch(formula)):
                rate_row, _ = _reference_parts(match.group("rate"))
                rate_label = _row_label(worksheet, rate_row, int(cell.column))
                if "growth" in rate_label:
                    repairs.append(
                        SignConventionRepair(
                            worksheet.title,
                            cell.coordinate,
                            formula,
                            f"={match.group('base')}*(1+{match.group('rate')})",
                            "revenue_growth_addition",
                            (label, rate_label),
                        )
                    )
                    continue

            if any(phrase in label for phrase in ("capex", "capital expenditure")) and (
                match := _NEGATIVE_PRODUCT_RE.fullmatch(formula)
            ):
                first_row, _ = _reference_parts(match.group("first"))
                second_row, _ = _reference_parts(match.group("second"))
                referenced_labels = {
                    _row_label(worksheet, first_row, int(cell.column)),
                    _row_label(worksheet, second_row, int(cell.column)),
                }
                if any("% of revenue" in item for item in referenced_labels) and any(
                    item == "revenue" or item.endswith(" revenue") for item in referenced_labels
                ):
                    repairs.append(
                        SignConventionRepair(
                            worksheet.title,
                            cell.coordinate,
                            formula,
                            f"={match.group('first')}*{match.group('second')}",
                            "capex_positive_schedule",
                            tuple(sorted(referenced_labels | {label})),
                        )
                    )
                    continue

            if any(
                phrase in label for phrase in ("change in nwc", "change in net working capital")
            ) and (match := _BINARY_SUBTRACT_RE.fullmatch(formula)):
                first_row, first_column = _reference_parts(match.group("first"))
                second_row, second_column = _reference_parts(match.group("second"))
                if (
                    first_row == second_row == int(cell.row) - 1
                    and first_column == int(cell.column) - 1
                    and second_column == int(cell.column)
                ):
                    repairs.append(
                        SignConventionRepair(
                            worksheet.title,
                            cell.coordinate,
                            formula,
                            f"={match.group('second')}-{match.group('first')}",
                            "working_capital_current_less_prior",
                            (label,),
                        )
                    )
                    continue

            if any(
                phrase in label for phrase in ("change in nwc", "change in net working capital")
            ) and (match := _NEGATIVE_PRODUCT_RE.fullmatch(formula)):
                first_row, _ = _reference_parts(match.group("first"))
                first_label = _row_label(worksheet, first_row, int(cell.column))
                if first_row == int(cell.row) + 1 and "% of" in first_label:
                    repairs.append(
                        SignConventionRepair(
                            worksheet.title,
                            cell.coordinate,
                            formula,
                            f"={match.group('first')}*{match.group('second')}",
                            "working_capital_preserve_signed_ratio",
                            (label, first_label),
                        )
                    )
                    continue

            if "unlevered free cash flow" in label:
                component_rows = {
                    "nopat": _nearby_labeled_row(
                        worksheet,
                        target_row=int(cell.row),
                        target_column=int(cell.column),
                        phrases=("nopat",),
                    ),
                    "depreciation": _nearby_labeled_row(
                        worksheet,
                        target_row=int(cell.row),
                        target_column=int(cell.column),
                        phrases=("depreciation",),
                    ),
                    "capex": _nearby_labeled_row(
                        worksheet,
                        target_row=int(cell.row),
                        target_column=int(cell.column),
                        phrases=("capex", "capital expenditure"),
                    ),
                    "working_capital": _nearby_labeled_row(
                        worksheet,
                        target_row=int(cell.row),
                        target_column=int(cell.column),
                        phrases=("change in working capital", "change in nwc"),
                    ),
                }
                if all(row is not None for row in component_rows.values()):
                    column = cell.column_letter
                    expected_current = (
                        f"={column}{component_rows['nopat']}+{column}{component_rows['depreciation']}"
                        f"-{column}{component_rows['capex']}+{column}{component_rows['working_capital']}"
                    )
                    replacement = (
                        f"={column}{component_rows['nopat']}+{column}{component_rows['depreciation']}"
                        f"-{column}{component_rows['capex']}-{column}{component_rows['working_capital']}"
                    )
                    if formula == expected_current:
                        repairs.append(
                            SignConventionRepair(
                                worksheet.title,
                                cell.coordinate,
                                formula,
                                replacement,
                                "ufcf_subtract_working_capital_increase",
                                tuple(component_rows),
                            )
                        )
                        continue

            if "implied equity value" in label and (match := _BINARY_ADD_RE.fullmatch(formula)):
                ev_row = _nearby_labeled_row(
                    worksheet,
                    target_row=int(cell.row),
                    target_column=int(cell.column),
                    phrases=("implied ev", "enterprise value"),
                    lookback=4,
                )
                debt_row = _nearby_labeled_row(
                    worksheet,
                    target_row=int(cell.row),
                    target_column=int(cell.column),
                    phrases=("net debt",),
                    lookback=3,
                )
                if (
                    ev_row is not None
                    and debt_row is not None
                    and _same_cell(match.group("first"), row=ev_row, column=int(cell.column))
                    and _same_cell(match.group("second"), row=debt_row, column=int(cell.column))
                ):
                    repairs.append(
                        SignConventionRepair(
                            worksheet.title,
                            cell.coordinate,
                            formula,
                            f"={match.group('first')}-{match.group('second')}",
                            "equity_value_less_net_debt",
                            (label, "implied ev", "net debt"),
                        )
                    )
                    continue

            # Financial models commonly store expense/cash-outflow line items as already-negative
            # formulas and then add them into EBIT/UFCF subtotals. Subtracting those cells again
            # reverses the sign. Require both an explicit ``(-)`` row marker and a negative source
            # formula before changing the subtotal operator.
            if _is_additive_subtotal_label(label):
                subtotal_replacement, negative_component_evidence = (
                    _add_already_negative_components(
                        worksheet,
                        formula,
                        target_column=int(cell.column),
                    )
                )
                if subtotal_replacement != formula:
                    repairs.append(
                        SignConventionRepair(
                            worksheet.title,
                            cell.coordinate,
                            formula,
                            subtotal_replacement,
                            "add_already_negative_line_items",
                            tuple([label, *negative_component_evidence]),
                        )
                    )
                    continue

            # Transaction costs reduce combined equity value. Require an explicit target label and
            # a referenced row labeled Transaction Costs; this also transfers across repeated deal
            # blocks without relying on fixed coordinates.
            if "value of newco equity" in label:
                transaction_replacement, transaction_evidence = (
                    _subtract_transaction_cost_components(
                        worksheet,
                        formula,
                        target_column=int(cell.column),
                    )
                )
                if transaction_replacement != formula:
                    repairs.append(
                        SignConventionRepair(
                            worksheet.title,
                            cell.coordinate,
                            formula,
                            transaction_replacement,
                            "newco_equity_less_transaction_costs",
                            tuple([label, *transaction_evidence]),
                        )
                    )
                    continue

            if label == "wacc":
                wacc_repaired = False
                for tax_match in _ONE_PLUS_LOCAL_REFERENCE_RE.finditer(formula):
                    tax_reference = tax_match.group("reference")
                    tax_row, _ = _reference_parts(tax_reference)
                    tax_label = _row_label(worksheet, tax_row, int(cell.column))
                    if "tax rate" not in tax_label:
                        continue
                    replacement = (
                        formula[: tax_match.start()]
                        + f"(1-{tax_reference})"
                        + formula[tax_match.end() :]
                    )
                    repairs.append(
                        SignConventionRepair(
                            worksheet.title,
                            cell.coordinate,
                            formula,
                            replacement,
                            "wacc_after_tax_debt_cost",
                            (label, tax_label),
                        )
                    )
                    wacc_repaired = True
                    break
                if wacc_repaired:
                    continue

            if "after tax cost of debt" in label and (match := _AFTER_TAX_RE.fullmatch(formula)):
                cost_row = _nearby_labeled_row(
                    worksheet,
                    target_row=int(cell.row),
                    target_column=int(cell.column),
                    phrases=("cost of debt",),
                    lookback=4,
                )
                tax_row = _nearby_labeled_row(
                    worksheet,
                    target_row=int(cell.row),
                    target_column=int(cell.column),
                    phrases=("tax rate",),
                    lookback=3,
                )
                if (
                    cost_row is not None
                    and tax_row is not None
                    and _same_cell(match.group("cost"), row=cost_row, column=int(cell.column))
                    and _same_cell(match.group("tax"), row=tax_row, column=int(cell.column))
                ):
                    repairs.append(
                        SignConventionRepair(
                            worksheet.title,
                            cell.coordinate,
                            formula,
                            f"={match.group('cost')}*(1-{match.group('tax')})",
                            "after_tax_cost_of_debt",
                            (label, "cost of debt", "tax rate"),
                        )
                    )

    return sorted(repairs, key=lambda item: (item.sheet.casefold(), item.cell))


def repair_sign_conventions(path: str | Path, *, task_hint: str) -> list[dict[str, Any]]:
    """Apply a complete set of task-gated semantic sign repairs to one workbook."""

    workbook_path = Path(path)
    workbook = load_workbook(
        workbook_path,
        data_only=False,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        repairs = detect_sign_convention_repairs(workbook, task_hint=task_hint)
        for repair in repairs:
            cell = workbook[repair.sheet][repair.cell]
            if cell.value != repair.current:
                return []
        for repair in repairs:
            workbook[repair.sheet][repair.cell] = repair.replacement
        if repairs:
            workbook.save(workbook_path)
        return [repair.to_dict() for repair in repairs]
    finally:
        workbook.close()


__all__ = [
    "SignConventionRepair",
    "detect_sign_convention_repairs",
    "repair_sign_conventions",
]
