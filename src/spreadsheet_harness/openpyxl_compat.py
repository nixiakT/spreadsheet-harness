"""Local compatibility helpers for supported openpyxl releases."""

from __future__ import annotations

import io
import re
import zipfile
from os import PathLike
from typing import Any
from xml.etree.ElementTree import ParseError

from openpyxl.chartsheet import Chartsheet
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.packaging.relationship import RelationshipList, get_rels_path
from openpyxl.reader.drawings import find_images
from openpyxl.reader.excel import ExcelReader
from openpyxl.xml.functions import fromstring

_OOXML_NAMESPACE_URIS = {
    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    "dc": "http://purl.org/dc/elements/1.1/",
    "dcterms": "http://purl.org/dc/terms/",
    "dcmitype": "http://purl.org/dc/dcmitype/",
    "xsi": "http://www.w3.org/2001/XMLSchema-instance",
}
_XML_PREFIX_USE_RE = re.compile(r"(?:<|</|\s)(?P<prefix>[A-Za-z_][\w.-]*):[A-Za-z_][\w.-]*")
_XMLNS_DECL_RE = re.compile(r"\sxmlns:(?P<prefix>[A-Za-z_][\w.-]*)=")
_ROOT_TAG_RE = re.compile(r"<(?P<tag>[A-Za-z_][\w.-]*(?::[A-Za-z_][\w.-]*)?)(?P<attrs>[^>]*)>")
_OOXML_METADATA_XML = frozenset({"docProps/core.xml"})


class _ChartsheetCompatibleExcelReader(ExcelReader):
    """Keep empty chartsheets loadable without changing worksheet handling."""

    def read_chartsheet(self, sheet: Any, rel: Any) -> None:
        sheet_path = rel.target
        rels_path = get_rels_path(sheet_path)
        if rels_path in self.valid_files:
            # Preserve openpyxl's drawing/image handling whenever relationships exist.
            super().read_chartsheet(sheet, rel)
            self.wb._sheets[-1].sheet_state = sheet.state
            return

        # openpyxl 3.1.5 initializes this as ``[]`` and then calls ``find`` on it.
        # Use the same empty relationship container as the worksheet reader.
        rels = RelationshipList()
        with self.archive.open(sheet_path, "r") as source:
            node = fromstring(source.read())
        chartsheet = Chartsheet.from_tree(node)
        chartsheet._parent = self.wb
        chartsheet.title = sheet.name
        chartsheet.sheet_state = sheet.state
        self.wb._add_sheet(chartsheet)

        for drawing_rel in rels.find(SpreadsheetDrawing._rel_type):
            charts, _images = find_images(self.archive, drawing_rel.target)
            for chart in charts:
                chartsheet.add_chart(chart)


def _load_workbook_once(
    filename: Any,
    read_only: bool = False,
    keep_vba: bool = False,
    data_only: bool = False,
    keep_links: bool = True,
    rich_text: bool = False,
) -> Any:
    reader = _ChartsheetCompatibleExcelReader(
        filename,
        read_only,
        keep_vba,
        data_only,
        keep_links,
        rich_text,
    )
    reader.read()
    return reader.wb


def _repair_unbound_prefix_xml(text: str) -> str:
    match = _ROOT_TAG_RE.search(text)
    if match is None:
        return text
    declared = {item.group("prefix") for item in _XMLNS_DECL_RE.finditer(match.group(0))}
    used = {item.group("prefix") for item in _XML_PREFIX_USE_RE.finditer(text)}
    missing = sorted(
        prefix
        for prefix in used - declared
        if prefix in _OOXML_NAMESPACE_URIS
    )
    if not missing:
        return text
    injected = "".join(f' xmlns:{prefix}="{_OOXML_NAMESPACE_URIS[prefix]}"' for prefix in missing)
    root = match.group(0)
    replacement = f"{root[:-1]}{injected}>"
    return text[: match.start()] + replacement + text[match.end() :]


def _repaired_archive_bytes(source: str | PathLike[str]) -> bytes | None:
    changed = False
    buffer = io.BytesIO()
    with zipfile.ZipFile(source) as archive:
        with zipfile.ZipFile(buffer, "w") as repaired:
            for info in archive.infolist():
                payload = archive.read(info.filename)
                if info.filename in _OOXML_METADATA_XML:
                    try:
                        fromstring(payload)
                    except ParseError as exc:
                        if "unbound prefix" in str(exc):
                            candidate = _repair_unbound_prefix_xml(
                                payload.decode("utf-8", errors="ignore")
                            ).encode("utf-8")
                            if candidate != payload:
                                payload = candidate
                                changed = True
                repaired.writestr(info, payload)
    if not changed:
        return None
    return buffer.getvalue()


def repair_workbook_archive_in_place(path: str | PathLike[str]) -> bool:
    repaired = _repaired_archive_bytes(path)
    if repaired is None:
        return False
    with open(path, "wb") as handle:
        handle.write(repaired)
    return True


def load_workbook(
    filename: Any,
    read_only: bool = False,
    keep_vba: bool = False,
    data_only: bool = False,
    keep_links: bool = True,
    rich_text: bool = False,
) -> Any:
    """Load a workbook with an isolated, idempotent chartsheet reader fix."""

    parse_error: ParseError | None = None
    try:
        return _load_workbook_once(
            filename,
            read_only=read_only,
            keep_vba=keep_vba,
            data_only=data_only,
            keep_links=keep_links,
            rich_text=rich_text,
        )
    except ParseError as exc:
        if "unbound prefix" not in str(exc) or not isinstance(filename, (str, PathLike)):
            raise
        parse_error = exc
    repaired = _repaired_archive_bytes(filename)
    if repaired is None:
        assert parse_error is not None
        raise parse_error
    return _load_workbook_once(
        io.BytesIO(repaired),
        read_only=read_only,
        keep_vba=keep_vba,
        data_only=data_only,
        keep_links=keep_links,
        rich_text=rich_text,
    )


__all__ = ["load_workbook", "repair_workbook_archive_in_place"]
