BYTES=4000
CELLS=4096
SHEETS=6
LABELS=6
import json
import re
from itertools import islice
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
def _emit(recs):
 out=[]
 for record in recs[:16]:
  record=dict(record)
  for attempt in range(128):
   line=json.dumps(record,ensure_ascii=False,separators=(",",":"))
   trial={"ok":True,"context":out+[line]}
   if len(line)<=4000 and len(json.dumps(trial,ensure_ascii=False).encode())<=BYTES:
    out.append(line); break
   lists=[value for value in record.values() if isinstance(value,list) and value]
   if not lists: break
   max(lists,key=lambda value:len(json.dumps(value,ensure_ascii=False))).pop()
   record["context_trimmed"]=True
 return {"ok":True,"context":out}
def _targets(text,names):
 addr=r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6})?"
 fnd=[]
 spans=[]
 for name in names:
  pfx="(?:"+re.escape("'"+name.replace("'","''")+"'!")+"|"+re.escape(name+"!")+")"
  for m in islice(re.finditer(r"(?<![A-Za-z0-9_])"+pfx+"("+addr+r")(?![A-Za-z0-9_])",text,re.I),16):
   value=m.group(1).replace("$","").upper()
   a,b,c,d=range_boundaries(value)
   if 1<=a<=c<=16384 and 1<=b<=d<=1048576:
    fnd.append([name,value])
    spans.append(m.span())
 bare=[]
 for m in islice(re.finditer(r"(?<![A-Za-z0-9_])"+addr+r"(?![A-Za-z0-9_])",text),32):
  if not any(a<=m.start()<b for a,b in spans):
   ref=m.group().replace("$","").upper()
   a,b,c,d=range_boundaries(ref)
   if not (1<=a<=c<=16384 and 1<=b<=d<=1048576): continue
   if re.fullmatch(r"(?:FY|CY)(?:19|20|21)\d{2}",ref): continue
   if re.fullmatch(r"Q[1-4]",ref) and not re.search(r"(?:cell|range)\s*$",text[:m.start()],re.I): continue
   bare.append(ref)
 ment=[name for name in names if name.casefold() in text.casefold()]
 return fnd[:16],bare[:8],ment[:12]
def _runs(values):
 out=[]
 for value in sorted(set(values)):
  if out and value==out[-1][1]+1:
   out[-1][1]=value
  else:
   out.append([value,value])
 return out
def _span(a,b,c,d):
 return get_column_letter(a)+str(b)+":"+get_column_letter(c)+str(d)
def _check(context):
 return set(context)=={"hook","instruction","task_category","workbook_path","config"} and context["hook"]=="before_task"
def _sheets(wb,tgs,ment):
 wanted={name for name,ref in tgs}|set(ment)
 return sorted(wb.worksheets,key=lambda sheet:sheet.title not in wanted)[:SHEETS]
def _label(labs,cell,text):
 value=str(cell.value)
 words=[word for word in re.findall(r"[A-Za-z0-9_%]+",value.casefold()) if len(word)>1]
 relevant=any(word in text.casefold() for word in words)
 rank=2*int(relevant)+int("%" in value or "%" in str(cell.number_format))
 labs.append([rank,cell.coordinate,value[:64]])
 labs.sort(key=lambda item:-item[0])
 del labs[LABELS:]
def on_hook(context):
 if not _check(context): return {"ok":False,"error":"before_task context required"}
 wb=load_workbook(context["workbook_path"],data_only=False,keep_links=False)
 try:
  text=str(context["instruction"])[:12000]
  tgs,bare,ment=_targets(text,wb.sheetnames)
  recs=[{"observer":"task-boundaries","category":str(context["task_category"])[:64],"explicit_targets":tgs,"unqualified_refs":bare,"mentioned_sheets":ment,"sheets_capped":len(wb.sheetnames)>SHEETS,"notes":"Nonempty bands are not permission to fill rectangles; preserve array anchors and unrelated cells."}]
  for sheet in _sheets(wb,tgs,ment):
   cells=list(islice(sheet._cells.values(),CELLS))
   occ=[]; labs=[]; arrays=[]; types={}; dates=[]; pct=[]; errors=[]
   for cell in cells:
    value=cell.value
    if value is None: continue
    occ.append((cell.row,cell.column))
    kind="formula" if cell.data_type=="f" else str(cell.data_type)
    types[kind]=types.get(kind,0)+1
    if hasattr(value,"ref") and cell.data_type=="f": arrays.append([cell.coordinate,str(value.ref)[:48]])
    if cell.is_date: dates.append(cell.coordinate)
    if "%" in str(cell.number_format): pct.append(cell.coordinate)
    if cell.data_type=="e": errors.append([cell.coordinate,str(value)[:24]])
    if isinstance(value,str) and cell.data_type not in {"f","e"}: _label(labs,cell,text)
   regions=[]
   gs=[]; pts=[]; prev=None
   for y,x in sorted(occ):
    if prev is not None and y>prev+1:
     gs.append(pts); pts=[]
    pts.append((y,x)); prev=y
   if pts: gs.append(pts)
   for pts in gs[:8]:
    lo,hi=pts[0][0],pts[-1][0]
    for left,right in _runs(x for y,x in pts):
     if len(regions)<8: regions.append(_span(left,lo,right,hi))
   row={"sheet":sheet.title[:100],"sampled_cells":len(cells),"scan_capped":len(sheet._cells)>CELLS,"nonempty_types":types,"labels":[label[1:] for label in labs],"nonempty_bands":regions,"array_anchors":arrays[:8],"percentage_count":len(pct),"date_count":len(dates),"merged_ranges":[str(x) for x in islice(sheet.merged_cells.ranges,4)]}
   recs.append(row)
  return _emit(recs)
 finally:
  wb.close()
