BYTES=6000
CELLS=20000
SHEETS=12
LABELS=12
import json
import re
from itertools import islice
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
def _emit(recs):
 out=[]
 for record in recs[:16]:
  record=dict(record)
  for attempt in range(128):
   line=json.dumps(record,ensure_ascii=False,separators=(",",":"))
   trial={"ok":True,"context":out+[line]}
   if len(line)<=4000 and len(json.dumps(trial,ensure_ascii=False).encode())<=BYTES:
    out.append(line); break
   lists=[value for value in record.values() if isinstance(value,list) and value]
   if not lists: break
   max(lists,key=lambda value:len(json.dumps(value,ensure_ascii=False))).pop()
   record["context_trimmed"]=True
 return {"ok":True,"context":out}
def _targets(text,names):
 addr=r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6})?"
 fnd=[]
 spans=[]
 for name in names:
  pfx="(?:"+re.escape("'"+name.replace("'","''")+"'!")+"|"+re.escape(name+"!")+")"
  for m in islice(re.finditer(r"(?<![A-Za-z0-9_])"+pfx+"("+addr+r")(?![A-Za-z0-9_])",text,re.I),16):
   value=m.group(1).replace("$","").upper()
   a,b,c,d=range_boundaries(value)
   if 1<=a<=c<=16384 and 1<=b<=d<=1048576:
    fnd.append([name,value])
    spans.append(m.span())
 bare=[]
 for m in islice(re.finditer(r"(?<![A-Za-z0-9_])"+addr+r"(?![A-Za-z0-9_])",text),32):
  if not any(a<=m.start()<b for a,b in spans):
   ref=m.group().replace("$","").upper()
   a,b,c,d=range_boundaries(ref)
   if not (1<=a<=c<=16384 and 1<=b<=d<=1048576): continue
   if re.fullmatch(r"(?:FY|CY)(?:19|20|21)\d{2}",ref): continue
   if re.fullmatch(r"Q[1-4]",ref) and not re.search(r"(?:cell|range)\s*$",text[:m.start()],re.I): continue
   bare.append(ref)
 ment=[name for name in names if name.casefold() in text.casefold()]
 return fnd[:16],bare[:8],ment[:12]
def _runs(values):
 out=[]
 for value in sorted(set(values)):
  if out and value==out[-1][1]+1:
   out[-1][1]=value
  else:
   out.append([value,value])
 return out
def _span(a,b,c,d):
 return get_column_letter(a)+str(b)+":"+get_column_letter(c)+str(d)
def _check(context):
 return set(context)=={"hook","instruction","task_category","workbook_path","config"} and context["hook"]=="before_task"
def _sheets(wb,tgs,ment):
 wanted={name for name,ref in tgs}|set(ment)
 return sorted(wb.worksheets,key=lambda sheet:sheet.title not in wanted)[:SHEETS]
def _label(labs,cell,text):
 value=str(cell.value)
 words=[word for word in re.findall(r"[A-Za-z0-9_%]+",value.casefold()) if len(word)>1]
 relevant=any(word in text.casefold() for word in words)
 rank=2*int(relevant)+int("%" in value or "%" in str(cell.number_format))
 labs.append([rank,cell.coordinate,value[:64]])
 labs.sort(key=lambda item:-item[0])
 del labs[LABELS:]
from datetime import date, datetime
def _period(value):
 if isinstance(value,(datetime,date)): return "date-header-candidate"
 if type(value) in (int,float) and 1900<=value<=2200 and value==int(value): return "year-like-number"
 if not isinstance(value,str): return None
 text=value.strip()
 if re.fullmatch(r"(?:FY|CY)?\s*(?:19|20|21)\d{2}[AE]?",text,re.I): return "year-header-candidate"
 if re.fullmatch(r"(?:Q[1-4](?:[- /]*(?:19|20|21)\d{2})?|(?:19|20|21)\d{2}[- /]*Q[1-4])",text,re.I): return "quarter-header-candidate"
 if re.fullmatch(r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)(?:[- /]+\d{2,4})?",text,re.I): return "month-header-candidate"
 if re.fullmatch(r"(?:19|20|21)\d{2}[-/](?:0?[1-9]|1[0-2])",text): return "month-like-text"
 return None
def on_hook(context):
 if not _check(context): return {"ok":False,"error":"before_task context required"}
 wb=load_workbook(context["workbook_path"],data_only=False,keep_links=False)
 try:
  text=str(context["instruction"])[:12000]
  tgs,bare,ment=_targets(text,wb.sheetnames)
  recs=[{"observer":"financial-targets","category":str(context["task_category"])[:64],"explicit_targets":tgs,"unqualified_refs":bare,"target_ambiguity":bool(bare) or not bool(tgs),"mentioned_sheets":ment,"notes":"Period-like cells are unconfirmed. Same column is not a cross-block/sheet join key; check metric, period, units and explicit task scope."}]
  for sheet in _sheets(wb,tgs,ment):
   cells=list(islice(sheet._cells.values(),CELLS))
   prs=[]; labs=[]; pct=[]
   for cell in cells:
    value=cell.value
    if value is None or cell.data_type in {"f","e"}: continue
    kind=_period(value)
    if kind and len(prs)<LABELS*3: prs.append([cell.row,cell.column,cell.coordinate,str(value)[:40],kind])
    elif isinstance(value,str): _label(labs,cell,text)
    if "%" in str(cell.number_format) and len(pct)<8: pct.append(cell.coordinate)
   row={"sheet":sheet.title[:100],"sampled_cells":len(cells),"scan_capped":len(sheet._cells)>CELLS,"period_candidates":[p[2:5] for p in prs[:LABELS]],"metric_labels":[label[1:] for label in labs],"percentage_cells":pct}
   gs=[]
   for y in sorted({p[0] for p in prs}):
    items=[p for p in prs if p[0]==y]
    for left,right in _runs(p[1] for p in items):
     if len(gs)<6:
      gs.append({"row":y,"span":_span(left,y,right,y),"headers":[p[2:5] for p in items if left<=p[1]<=right][:8]})
   row["independent_period_groups"]=gs
   repeated={}
   for y,x,at,label,kind in prs: repeated.setdefault(label,set()).add(y)
   row["same_label_multiple_rows"]=[[label,sorted(rows)[:4]] for label,rows in repeated.items() if len(rows)>1][:6]
   row["array_anchors"]=[[c.coordinate,str(c.value.ref)[:48]] for c in cells if c.data_type=="f" and hasattr(c.value,"ref")][:4]
   recs.append(row)
  return _emit(recs)
 finally:
  wb.close()
