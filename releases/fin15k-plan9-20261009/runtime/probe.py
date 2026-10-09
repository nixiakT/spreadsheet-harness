#!/usr/bin/env python3
"""Live end-to-end launch canary using an artificial workbook, not v2 answers.

The frozen baseline and 50-joint artifacts each solve the same tiny task. This
checks source isolation, tool replay, generated hooks and final workbook I/O;
it is not a SpreadsheetBench score or evidence for further candidate tuning.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook

_CHILD = r'''
import json, sys, traceback
from pathlib import Path
from openpyxl import load_workbook
from spreadsheet_harness import arms
from spreadsheet_harness.budget import RunBudget
from spreadsheet_harness.config import ProviderConfig
from spreadsheet_harness.errors import AgentExecutionFailure, ProviderError
from spreadsheet_harness.plugins import CompositionSpec
from spreadsheet_harness.session import WorkbookSession
from spreadsheet_harness.skills import SkillRegistry

artifact=Path(sys.argv[1]).resolve()
output=Path(sys.argv[2])
settings=json.loads(sys.argv[3])
assert Path(arms.__file__).resolve().is_relative_to(artifact/'src')
config=ProviderConfig(base_url=settings['base_url'],api_key=Path(settings['api_key_file']).read_text().strip(),
    model=settings['model'],api_protocol='chat-completions',reasoning_effort='medium',
    temperature=0,top_p=1,seed=20261009,enable_thinking=settings['enable_thinking'],
    timeout_seconds=settings['request_timeout'],litellm_timeout_seconds=settings['litellm_timeout'],
    max_retries=settings['request_retries'],request_interval_seconds=settings['request_interval'])
doc=json.loads((artifact.parent/'composition.json').read_text())
composition=CompositionSpec.create(doc['name'],doc['plugins'],doc.get('overrides',{}))
session=WorkbookSession.create(output.parent/'artificial-input.xlsx',output/'session',recorder_secrets=(config.api_key,))
report={'artifact':str(artifact),'loaded_arms':arms.__file__,'model':config.model,
        'artificial_not_benchmark':True,'normal_completion':False}
try:
    result=arms.run_arm('spreadsheet-harness-financial',config,session,SkillRegistry([artifact/'skills']).freeze(),
        'On Forecast, write the formula =C2+C3 into C4 and =D2+D3 into D4. '
        'These are the ONLY two cells to edit. Preserve every other cell, including blank C6 and D6. '
        'Save, recalculate, verify the results C4=80 and D4=85, and submit.',
        max_output_tokens=settings['max_output_tokens'],max_elapsed_seconds=settings['task_timeout'],
        budget=RunBudget(max_model_calls=settings['max_model_calls'],max_total_tokens=None,
                         max_elapsed_seconds=settings['task_timeout']),
        max_turns_per_arm=50,composition=composition,task_category='Financial_Model')
    report.update(normal_completion=True,agent=result.to_dict())
except AgentExecutionFailure as exc:
    report.update(error_type=type(exc).__name__,error=str(exc).replace(config.api_key,'[REDACTED]'))
    if exc.agent_result is not None:
        report['agent']=exc.agent_result.to_dict()
except Exception as exc:
    report.update(error_type=type(exc).__name__,error=str(exc).replace(config.api_key,'[REDACTED]'))
book=load_workbook(session.workbook_path,data_only=False)
sheet=book['Forecast']
report['cells']={cell:sheet[cell].value for cell in ('C4','D4','C6','D6')}
report['target_formulas_ok']=sheet['C4'].value=='=C2+C3' and sheet['D4'].value=='=D2+D3'
report['untouched_blanks_ok']=sheet['C6'].value is None and sheet['D6'].value is None
book.close()
report['ok']=report['normal_completion'] and report['target_formulas_ok'] and report['untouched_blanks_ok']
print(json.dumps(report,ensure_ascii=False,default=str))
'''


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--bundle", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--model", required=True)
    parser.add_argument("--base-url", required=True)
    parser.add_argument("--api-key-file", type=Path, required=True)
    parser.add_argument("--request-timeout", type=float, default=1800)
    parser.add_argument("--litellm-timeout", type=float, default=1800)
    parser.add_argument("--task-timeout", type=float, default=21600)
    parser.add_argument("--request-retries", type=int, default=1)
    parser.add_argument("--request-interval", type=float, default=1.1)
    parser.add_argument("--max-model-calls", type=int, default=50)
    parser.add_argument("--max-output-tokens", type=int, default=32768)
    parser.add_argument("--workers", type=int, choices=(1, 2), default=2)
    parser.add_argument("--global-limiter-file", type=Path)
    parser.add_argument("--global-requests-per-minute", type=int)
    parser.add_argument("--global-tokens-per-minute", type=int)
    thinking = parser.add_mutually_exclusive_group(required=True)
    thinking.add_argument("--enable-thinking", dest="enable_thinking", action="store_true")
    thinking.add_argument("--disable-thinking", dest="enable_thinking", action="store_false")
    args = parser.parse_args()
    limiter_values = (args.global_limiter_file, args.global_requests_per_minute,
                      args.global_tokens_per_minute)
    if any(value is not None for value in limiter_values):
        if not all(value is not None for value in limiter_values):
            parser.error("All three --global-* limiter options must be specified together")
        if args.global_requests_per_minute <= 0 or args.global_tokens_per_minute <= 0:
            parser.error("Global limiter request/token limits must be positive")
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Forecast"
    sheet.append(["Metric", None, 2024, 2025])
    sheet.append(["Revenue", None, 100, 110])
    sheet.append(["Expenses", None, -20, -25])
    sheet.append(["Net Profit"])
    sheet["A6"] = "Not requested"
    workbook.save(output / "artificial-input.xlsx")
    workbook.close()
    settings = {
        "model": args.model, "base_url": args.base_url,
        "api_key_file": str(args.api_key_file.resolve()), "enable_thinking": args.enable_thinking,
        "request_timeout": args.request_timeout, "litellm_timeout": args.litellm_timeout,
        "task_timeout": args.task_timeout, "request_retries": args.request_retries,
        "request_interval": args.request_interval, "max_model_calls": args.max_model_calls,
        "max_output_tokens": args.max_output_tokens,
        "workers": args.workers,
        "global_limiter": None if args.global_limiter_file is None else {
            "file": str(args.global_limiter_file.resolve()),
            "requests_per_minute": args.global_requests_per_minute,
            "tokens_per_minute": args.global_tokens_per_minute,
        },
    }
    bundle = args.bundle.resolve()
    document = json.loads((bundle / "candidate-manifest.json").read_text())
    joint = next(row for row in document["candidates"]
                 if row["evidence_size"] == 50 and row["mechanism"] == "joint")
    joint_path = Path(joint["artifact"])
    artifacts = {"baseline": bundle / "baseline/artifact", "50-joint": joint_path if joint_path.is_absolute() else bundle / joint_path}

    def execute(item: tuple[str, Path]) -> dict:
        label, artifact = item
        destination = output / label
        destination.mkdir()
        proxy_names = {"http_proxy", "https_proxy", "HTTP_PROXY", "HTTPS_PROXY",
                       "all_proxy", "ALL_PROXY", "no_proxy", "NO_PROXY"}
        environment = {key: value for key, value in os.environ.items()
                       if not key.startswith(("SHEET_HARNESS_", "SHEET_AGENT_"))
                       and key not in proxy_names}
        environment["PYTHONPATH"] = str(artifact / "src")
        environment["PYTHONDONTWRITEBYTECODE"] = "1"
        environment["PYTHONPYCACHEPREFIX"] = str(destination / "private-pycache")
        environment["NO_PROXY"] = "*"
        environment["no_proxy"] = "*"
        limiter = settings["global_limiter"]
        if limiter is not None:
            environment["SHEET_AGENT_GLOBAL_LIMITER_FILE"] = limiter["file"]
            environment["SHEET_AGENT_GLOBAL_REQUESTS_PER_MINUTE"] = str(limiter["requests_per_minute"])
            environment["SHEET_AGENT_GLOBAL_TOKENS_PER_MINUTE"] = str(limiter["tokens_per_minute"])
        with (destination / "runner.log").open("x", encoding="utf-8") as log:
            result = subprocess.run(
                [sys.executable, "-c", _CHILD, str(artifact), str(destination), json.dumps(settings)],
                cwd=artifact, env=environment, stdout=log, stderr=subprocess.STDOUT,
                timeout=settings["task_timeout"] + settings["request_timeout"] + 300,
                check=False,
            )
        report = json.loads((destination / "runner.log").read_text().splitlines()[-1])
        report.update(label=label, returncode=result.returncode)
        (destination / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
        print(json.dumps({key: report[key] for key in
                          ("label", "returncode", "ok", "normal_completion", "cells", "loaded_arms")}), flush=True)
        return report

    with concurrent.futures.ThreadPoolExecutor(max_workers=args.workers) as pool:
        rows = list(pool.map(execute, artifacts.items()))
    report = {"artificial_not_benchmark": True, "settings": settings, "results": rows,
              "ok": all(row["ok"] for row in rows),
              "probe_source_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest()}
    (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
