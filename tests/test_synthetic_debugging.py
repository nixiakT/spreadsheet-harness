from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from spreadsheet_harness.synthetic_debugging import (
    build_synthetic_debugging_corpus,
    run_deterministic_debugging_repair,
    score_synthetic_debugging_outputs,
)


def _formula_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Forecast"
    worksheet["B2"] = "=B1*2"
    worksheet["C2"] = "=C1*2"
    worksheet["D2"] = "=D1*2"
    workbook.save(path)
    workbook.close()


def test_build_and_score_synthetic_debugging_corpus(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _formula_workbook(source)
    metadata = {
        "files": [
            {
                "filename": "source.xlsx",
                "title": "Forecast workbook",
                "query": "financial forecast",
            }
        ]
    }
    archive = tmp_path / "corpus.zip"
    with ZipFile(archive, "w") as bundle:
        bundle.write(source, "source.xlsx")
        bundle.writestr("xlsx_metadata.json", json.dumps(metadata))

    output = tmp_path / "synthetic"
    manifest = build_synthetic_debugging_corpus(archive, output, limit=1, seed=7)

    assert manifest["generation"]["case_count"] == 1
    case = manifest["cases"][0]
    assert case["source"]["title"] == "Forecast workbook"
    assert case["mutation"]["original_formula"] == "=C1*2"
    synthetic_input = load_workbook(output / case["input_path"], data_only=False)
    assert synthetic_input["Forecast"]["C2"].value == "=0"
    synthetic_input.close()

    candidate = tmp_path / f"{case['id']}_output.xlsx"
    golden = load_workbook(output / case["golden_path"], data_only=False)
    golden.save(candidate)
    golden.close()
    report = score_synthetic_debugging_outputs(output / "manifest.json", [candidate])

    assert report["passed_count"] == 1
    assert report["results"][0]["regression_cells"] == 0

    repaired_dir = tmp_path / "repaired"
    run = run_deterministic_debugging_repair(output / "manifest.json", repaired_dir)
    repaired_report = score_synthetic_debugging_outputs(
        output / "manifest.json", [repaired_dir]
    )

    assert run["case_count"] == 1
    assert repaired_report["passed_count"] == 1
