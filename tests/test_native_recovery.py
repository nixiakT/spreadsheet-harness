from __future__ import annotations

import json

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from spreadsheet_harness.agent import _responses_input_to_chat_messages
from spreadsheet_harness.render import (
    _restore_ooxml_array_formulas,
    _seed_ooxml_formula_cached_values,
    repair_mistyped_error_caches,
)
from spreadsheet_harness.spreadsheetbench_v1 import (
    replay_v1_calls,
    successful_v1_replay_calls,
)


@pytest.mark.parametrize("model", ["DeepSeek-V4-Flash", "dashscope/qwen3-coder-480b-a35b-instruct"])
def test_text_only_routes_explicitly_omit_images(model):
    result = _responses_input_to_chat_messages(None, [{"role": "user", "content": [
        {"type": "input_text", "text": "Workbook preview"},
        {"type": "input_image", "image_url": "data:image/png;base64,secret"},
    ]}], model=model)
    assert "Workbook preview" in result[0]["content"]
    assert "Image not delivered" in result[0]["content"]
    assert "base64" not in result[0]["content"]


def test_other_routes_keep_images():
    result = _responses_input_to_chat_messages(None, [{"role": "user", "content": [
        {"type": "input_image", "image_url": "data:image/png;base64,abc"},
    ]}], model="vision-test")
    assert result[0]["content"][0]["type"] == "image_url"


@pytest.mark.parametrize("value", ["#NAME?", "#VALUE!", True, 42, "hello"])
def test_array_restore_keeps_cache_type(tmp_path, value):
    source, converted = tmp_path / "source.xlsx", tmp_path / "converted.xlsx"
    wb = Workbook()
    wb.active["A1"] = ArrayFormula(ref="A1", text="=UNKNOWN(A2)")
    wb.save(source)
    wb.active["A1"] = value
    wb.save(converted)
    wb.close()
    _restore_ooxml_array_formulas(source, converted)
    result = load_workbook(converted, data_only=True)
    assert result.active["A1"].value == value
    result.close()
    result = load_workbook(converted, data_only=False)
    assert isinstance(result.active["A1"].value, ArrayFormula)
    result.close()


def test_cache_seed_allows_new_or_deleted_sheet(tmp_path):
    source, target = tmp_path / "source.xlsx", tmp_path / "target.xlsx"
    wb = Workbook()
    wb.active["A1"] = "=1+1"
    wb.create_sheet("Deleted")
    wb.save(source)
    del wb["Deleted"]
    wb.create_sheet("Added")["A1"] = "=2+2"
    wb.save(target)
    wb.close()
    _seed_ooxml_formula_cached_values(target, source)
    result = load_workbook(target, data_only=False)
    assert result.sheetnames == ["Sheet", "Added"]
    assert result["Added"]["A1"].value == "=2+2"
    result.close()


def test_v1_bash_mutation_is_replayed(tmp_path):
    source = tmp_path / "source.xlsx"
    wb = Workbook()
    wb.active["A1"] = 1
    wb.save(source)
    wb.close()
    events = []
    for changed, command in [(False, "pwd"), (True, "python -c 'import sheet_harness; w=sheet_harness.load_workbook(); w.active[\"A1\"]=7; sheet_harness.save_workbook(w)'")]:
        events.extend([
            {"event": "tool.called", "payload": {"name": "bash", "arguments": {"command": command}}},
            {"event": "tool.returned", "payload": {"name": "bash", "result": {"ok": True, "workbook_changed": changed}}},
        ])
    trajectory = tmp_path / "trajectory.jsonl"
    trajectory.write_text("\n".join(json.dumps(e) for e in events))
    calls = successful_v1_replay_calls(trajectory)
    assert len(calls) == 1
    replay = replay_v1_calls(source, tmp_path / "replay", calls)
    assert replay["status"] == "scored", replay
    result = load_workbook(replay["output_workbook"])
    assert result.active["A1"].value == 7
    result.close()


def test_error_cache_repair_preserves_error_value(tmp_path):
    import zipfile

    from spreadsheet_harness.render import _replace_ooxml_parts

    path = tmp_path / "output.xlsx"
    wb = Workbook()
    wb.active["A1"] = "#NAME?"
    wb.active["A2"] = "ordinary text"
    wb.save(path)
    wb.close()
    with zipfile.ZipFile(path) as z:
        xml = z.read("xl/worksheets/sheet1.xml").replace(b't="e"', b't="n"')
    _replace_ooxml_parts(path, {"xl/worksheets/sheet1.xml": xml})
    assert repair_mistyped_error_caches(path) == 1
    wb = load_workbook(path, data_only=True)
    assert wb.active["A1"].value == "#NAME?"
    assert wb.active["A1"].data_type == "e"
    assert wb.active["A2"].value == "ordinary text"
    wb.close()
    assert repair_mistyped_error_caches(path) == 0


def test_v1_official_evaluator_exceptions_are_false_not_dropped(tmp_path):
    import importlib.util
    from pathlib import Path

    from spreadsheet_harness.spreadsheetbench_v1 import (
        SpreadsheetBenchV1Case,
        SpreadsheetBenchV1Instruction,
    )

    spec = importlib.util.spec_from_file_location("recovery_test", Path(__file__).parents[1] / "tools/complete_native_benchmarks.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.save(path)
    wb.close()
    task = SpreadsheetBenchV1Instruction("x", "edit", "Cell-Level Manipulation", "A:G", None,
        tuple(SpreadsheetBenchV1Case(i, path, path) for i in (1, 2, 3)))
    score = module.score_v1(task, {i: path for i in (1, 2, 3)}, tmp_path / "eval.log")
    assert score["soft"] == score["hard"] == 0
    assert len(score["case_results"]) == 3
    assert all("official_evaluator_exception" in c["reason"] for c in score["case_results"])
