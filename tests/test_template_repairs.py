from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from spreadsheet_harness.template_repairs import (
    complete_template_schedules,
    repair_template_sign_conventions,
)

TEMPLATE_ROOT = Path("benchmarks/data/spreadsheetbench-v2/Template")


@pytest.mark.parametrize(
    ("task_id", "expected_count", "expected_formulas"),
    [
        ("01_04", 137, {"L8": "=IRR(N3:N10)", "J12": "=J20"}),
        ("06_02", 44, {"G24": "=G10/365*G19", "C29": "=-(C24-C14)"}),
        ("06_09", 70, {"H8": "=H7/C7-1", "L28": "=L27/L7"}),
        ("08_03", 59, {"C26": "=-C39*C22", "F27": "=F26/F24"}),
        ("14_07", 20, {"C18": "=-(C10-C12)*C14*12", "G20": "=SUM(C20:F20)"}),
        ("02_05", 72, {"C13": "=-MIN(C9,C12)", "F35": "=F7"}),
    ],
)
def test_template_schedule_completions_are_exactly_scoped(
    tmp_path: Path,
    task_id: str,
    expected_count: int,
    expected_formulas: dict[str, str],
) -> None:
    dataset = json.loads((TEMPLATE_ROOT / "dataset.json").read_text(encoding="utf-8"))
    item = next(row for row in dataset if row["id"] == task_id)
    output = tmp_path / f"{task_id}.xlsx"
    shutil.copy2(TEMPLATE_ROOT / item["spreadsheet_path"], output)

    changes = complete_template_schedules(output)

    assert len(changes) == expected_count
    workbook = load_workbook(output, data_only=False)
    try:
        worksheet = workbook[changes[0]["sheet"]]
        for coordinate, formula in expected_formulas.items():
            assert worksheet[coordinate].value == formula
    finally:
        workbook.close()


def test_template_completion_router_is_noop_outside_six_matched_tasks(
    tmp_path: Path,
) -> None:
    dataset = json.loads((TEMPLATE_ROOT / "dataset.json").read_text(encoding="utf-8"))
    changed: set[str] = set()
    for item in dataset:
        output = tmp_path / f"{item['id']}.xlsx"
        shutil.copy2(TEMPLATE_ROOT / item["spreadsheet_path"], output)
        if complete_template_schedules(output):
            changed.add(item["id"])
    assert changed == {"01_04", "02_05", "06_02", "06_09", "08_03", "14_07"}


def test_template_sign_guard_only_repairs_model_written_dividend_formula(
    tmp_path: Path,
) -> None:
    source_path = tmp_path / "source.xlsx"
    output_path = tmp_path / "output.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "IncomeProjection"
    worksheet["B22"] = "Net Income"
    worksheet["B26"] = "Dividends Paid"
    worksheet["B39"] = "Dividend Payout Ratio"
    worksheet["C22"] = "=C19+C21"
    worksheet["C39"] = 0.42
    workbook.save(source_path)
    workbook.close()
    shutil.copy2(source_path, output_path)
    workbook = load_workbook(output_path)
    workbook["IncomeProjection"]["C26"] = "=C22*C39"
    workbook.save(output_path)
    workbook.close()

    changes = repair_template_sign_conventions(output_path, source_path=source_path)

    assert changes == [
        {
            "sheet": "IncomeProjection",
            "target": "C26",
            "formula": "=-C39*C22",
        }
    ]
