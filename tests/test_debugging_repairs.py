from __future__ import annotations

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from spreadsheet_harness.debugging_repairs import (
    detect_debugging_repair_candidates,
    repair_broken_sheet_qualifiers,
    restore_deleted_scenario_selector_row,
    restore_missing_rate_driver_rows,
    repair_semantic_broken_references,
    restore_structural_error_rows,
    restore_missing_assumption_rows,
)


def test_deleted_scenario_selector_row_is_inferred_across_layout_offsets() -> None:
    for offset in (0, 7):
        workbook = Workbook()
        scenario = workbook.active
        scenario.title = "Scenario Engine"
        for row in (offset + 1, offset + 2):
            for column in range(5, 8):
                scenario.cell(row, column).value = row * column
        content_row = offset + 5
        inserted_at = content_row - 1
        active_row = offset + 10
        scenario.cell(content_row, 4).value = '="Acquisition Target"'
        scenario.cell(active_row, 4).value = "Active Case"
        broken_rows = (offset + 12, offset + 15, offset + 18)
        for row in broken_rows:
            scenario.cell(row, 5).value = f"=CHOOSE(#REF!,H{row},I{row},J{row})"
        links = workbook.create_sheet("Summary")
        links["B3"] = f"='Scenario Engine'!E{broken_rows[0]}"

        actions = restore_deleted_scenario_selector_row(
            workbook,
            instruction=(
                "Please restore deleted rows. The model uses 3 scenario cases. "
                "The active scenario case selector value is 2."
            ),
        )

        assert len(actions) == 1
        assert actions[0]["sheet"] == "Scenario Engine"
        assert actions[0]["target"] == f"{inserted_at}:{inserted_at}"
        assert scenario.cell(inserted_at, 4).value == "Case"
        assert scenario.cell(inserted_at, 5).value == 2
        assert scenario.cell(content_row + 1, 4).value == '="Acquisition Target"'
        for old_row in broken_rows:
            new_row = old_row + 1
            assert scenario.cell(new_row, 5).value == (
                f"=CHOOSE($E${inserted_at},H{new_row},I{new_row},J{new_row})"
            )
        assert links["B3"].value == f"='Scenario Engine'!E{broken_rows[0] + 1}"
        workbook.close()


def test_deleted_scenario_selector_row_fails_closed_without_explicit_structure_task() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["D10"] = "Active Case"
    for row in (12, 15, 18):
        worksheet.cell(row, 5).value = f"=CHOOSE(#REF!,H{row},I{row},J{row})"

    assert (
        restore_deleted_scenario_selector_row(
            workbook,
            instruction="Please audit this scenario model thoroughly.",
        )
        == []
    )
    workbook.close()


def test_broken_sheet_qualifier_repair_requires_a_unique_inventory_match() -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    workbook.create_sheet("North Ops")
    workbook.create_sheet("North Plan")
    summary["C7"] = "='[1]North 0ps'!D9+'North Ops'!D10"
    summary["C8"] = "='[Budget.xlsx]North Ops'!D9"
    summary["C9"] = "='North'!D9"

    actions = repair_broken_sheet_qualifiers(
        workbook,
        instruction="Fix broken cross-sheet references and typos in sheet names.",
    )

    assert len(actions) == 1
    assert summary["C7"].value == "='North Ops'!D9+'North Ops'!D10"
    assert summary["C8"].value == "='[Budget.xlsx]North Ops'!D9"
    assert summary["C9"].value == "='North'!D9"
    workbook.close()


def test_broken_sheet_qualifier_repair_is_instruction_gated() -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    workbook.create_sheet("Operations")
    summary["B2"] = "='0perations'!C4"

    assert repair_broken_sheet_qualifiers(workbook, instruction="Audit formulas.") == []
    assert summary["B2"].value == "='0perations'!C4"
    workbook.close()


def test_broken_sheet_qualifier_repair_accepts_wrong_sheet_name_wording() -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    workbook.create_sheet("Financial Performance")
    summary["C7"] = "='[1]Financial Perf'!D9"

    actions = repair_broken_sheet_qualifiers(
        workbook,
        instruction="Fix wrong sheet name references and cascading errors.",
    )

    assert len(actions) == 1
    assert summary["C7"].value == "='Financial Performance'!D9"
    workbook.close()


def test_structural_error_row_uses_component_and_cross_sheet_witnesses() -> None:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview["B4"] = "Energy"
    overview["B5"] = "% Growth"
    overview["B6"] = "Engineering"
    overview["B7"] = "% Growth"
    overview["B8"] = "% Growth"
    overview["C4"] = 10
    overview["D4"] = 12
    overview["C6"] = 3
    overview["D6"] = 4
    overview["C8"] = "=(#REF!/#REF!)-1"
    model = workbook.create_sheet("Model")
    model["B10"] = "Total Contract Revenue"
    model["C10"] = "='Overview'!#REF!"

    actions = restore_structural_error_rows(
        workbook,
        instruction="Audit deleted rows and broken #REF! references.",
    )

    assert len(actions) == 1
    assert overview["B8"].value == "Total Contract Revenue"
    assert overview["C8"].value == "=SUM(C4,C6)"
    assert overview["D8"].value == "=SUM(D4,D6)"
    assert model["C10"].value == "='Overview'!#REF!"
    workbook.close()


def test_structural_error_row_fails_closed_without_unique_witness() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["B4"] = "Energy"
    worksheet["B5"] = "% Growth"
    worksheet["B6"] = "Engineering"
    worksheet["B7"] = "% Growth"
    worksheet["B8"] = "% Growth"
    worksheet["C8"] = "=(#REF!/#REF!)-1"
    model = workbook.create_sheet("Model")
    model["B10"] = "One"
    model["C10"] = "='Sheet'!#REF!"
    model["B11"] = "Two"
    model["C11"] = "='Sheet'!#REF!"

    assert restore_structural_error_rows(
        workbook,
        instruction="Audit deleted rows and broken #REF! references.",
    ) == []
    workbook.close()


def test_missing_assumption_row_uses_parallel_control_block_and_formula_witness() -> None:
    workbook = Workbook()
    dcf = workbook.active
    dcf.title = "DCF"
    dcf["B8"] = "Effective Tax Rate"
    dcf["C8"] = "=WACC!C17"
    dcf["B9"] = "Terminal Value Growth Rate"
    dcf["C9"] = "=Assumptions!C10"
    dcf["E9"] = "Total Revenue"
    dcf["I29"] = "=I26/(1+#REF!)^I2"
    dcf["R27"] = "=R26*(1+C9)/(#REF!-C9)"
    wacc = workbook.create_sheet("WACC")
    wacc["B22"] = "WACC"
    wacc["C22"] = 0.08
    assumptions = workbook.create_sheet("Assumptions")
    assumptions["B10"] = "Terminal Value Growth Rate"
    assumptions["C10"] = 0.02

    actions = restore_missing_assumption_rows(
        workbook,
        instruction="Audit deleted rows and repair #REF! formulas.",
    )

    assert len(actions) == 1
    assert dcf["B9"].value == "WACC"
    assert dcf["C9"].value == "=WACC!C22"
    assert dcf["B10"].value == "Terminal Value Growth Rate"
    assert dcf["I29"].value == "=I26/(1+$C$9)^I2"
    assert dcf["R27"].value == "=R26*(1+C10)/(C9-C10)"
    workbook.close()


def test_margin_denominator_reanchors_after_structural_insertion() -> None:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview["B3"] = "Energy"
    overview["B4"] = "% Growth"
    overview["B5"] = "Engineering"
    overview["B6"] = "% Growth"
    overview["B7"] = "% Growth"
    overview["B8"] = "Total Net Revenue"
    overview["B9"] = "Gross Profit"
    overview["B10"] = "% Margin (Contract Rev.)"
    overview["C3"] = 10
    overview["C5"] = 3
    overview["C7"] = "=(#REF!/#REF!)-1"
    overview["C9"] = 5
    overview["C10"] = "=C9/C11"
    model = workbook.create_sheet("Model")
    model["B12"] = "Total Contract Revenue"
    model["C12"] = "='Overview'!#REF!"

    actions = restore_structural_error_rows(
        workbook,
        instruction="Audit deleted rows and broken #REF! references.",
    )
    assert actions
    repair_semantic_broken_references(
        workbook,
        instruction="Audit deleted rows and broken #REF! references.",
    )
    assert overview["C11"].value == "=C10/C7"
    workbook.close()


def test_missing_rate_driver_row_uses_unique_curve_header_and_years() -> None:
    workbook = Workbook()
    model = workbook.active
    model.title = "Model"
    model["B10"] = "Debt Schedule"
    model["B12"] = "Revolver"
    model["B16"] = "Interest"
    model["J16"] = "=SUM($I16,#REF!)*AVERAGE(J13,J15)"
    model["K16"] = "=SUM($I16,#REF!)*AVERAGE(K13,K15)"
    model["L16"] = "=SUM($I16,#REF!)*AVERAGE(L13,L15)"
    model["B20"] = "Interest"
    model["J20"] = "=SUM($I20,#REF!)*AVERAGE(J17,J19)"
    model["K20"] = "=SUM($I20,#REF!)*AVERAGE(K17,K19)"
    model["L20"] = "=SUM($I20,#REF!)*AVERAGE(L17,L19)"
    model["B5"] = "Fiscal Year"
    model["J5"] = 2026
    model["K5"] = "=J5+1"
    model["L5"] = "=K5+1"
    curve = workbook.create_sheet("3-month Term SOFR")
    curve["P9"] = "3-month Term SOFR"
    curve["O10"] = 2026
    curve["O11"] = 2027
    curve["O12"] = 2028
    curve["P10"] = 0.03
    curve["P11"] = 0.031
    curve["P12"] = 0.032

    actions = restore_missing_rate_driver_rows(
        workbook,
        instruction="Repair deleted rows and broken #REF! references.",
    )

    assert len(actions) == 1
    assert model["B11"].value == "SOFR"
    assert model["J11"].value == "='3-month Term SOFR'!P10"
    assert model["J17"].value == "=SUM($I17,J$11)*AVERAGE(J14,J16)"
    workbook.close()


def test_semantic_broken_reference_uses_source_label_and_aggregate_row() -> None:
    workbook = Workbook()
    model = workbook.active
    model.title = "Model"
    model["B2"] = "Total Revenue"
    model["C2"] = "='Overview'!#REF!"
    overview = workbook.create_sheet("Overview")
    overview["B4"] = "Energy"
    overview["B6"] = "Engineering"
    overview["B8"] = "Total Revenue"
    overview["C8"] = "=SUM(C4,C6)"
    model["C1"] = "=SUM(C5,C6)"
    model["B3"] = "Gross Profit"
    model["C3"] = "=#REF!-10"

    actions = repair_semantic_broken_references(
        workbook,
        instruction="Repair broken #REF! references after deleted rows.",
    )

    assert len(actions) == 2
    assert model["C2"].value == "='Overview'!C8"
    assert model["C3"].value == "=C1-10"
    workbook.close()


def test_semantic_broken_reference_uses_labelled_rate_row_for_weighted_formula() -> None:
    workbook = Workbook()
    model = workbook.active
    model.title = "Model"
    model["B5"] = "SOFR"
    model["J5"] = "='Curve'!P10"
    model["I6"] = 0.04
    model["I7"] = 0.08
    overview = workbook.create_sheet("Overview")
    overview["B2"] = "Total Debt"
    overview["C2"] = "=((Model!I6+Model!#REF!)*Model!I7)"
    actions = repair_semantic_broken_references(
        workbook,
        instruction="Repair broken #REF! references after a deleted row.",
    )
    assert len(actions) == 1
    assert overview["C2"].value == "=((Model!I6+Model!J5)*Model!I7)"
    workbook.close()


def test_cross_sheet_semantics_align_entity_metric_and_header_columns() -> None:
    workbook = Workbook()
    target = workbook.active
    target.title = "Model"
    source = workbook.create_sheet("Source")
    source["A1"] = "Source table"
    source["B4"] = "Treasury"
    source["C4"] = "Expected"
    source["B5"] = "Yields"
    source["C5"] = "Inflation"
    source["A7"] = "Net income"
    source["B7"] = 8
    source["A8"] = "Operating income"
    source["B8"] = 10
    source["C8"] = 20
    source["A9"] = "Exploration"
    source["B9"] = 2
    source["A10"] = "Depreciation"
    source["B10"] = 3
    target["B2"] = "Terminal Value Growth Rate"
    target["C2"] = "=+'Source'!B8"
    target["B3"] = "EBITDAX"
    target["C3"] = "='Source'!B7+'Source'!B9+'Source'!B10"
    target["B4"] = "Occidental"
    source["D4"] = "Chevron"
    source["E4"] = "Occidental"
    source["E8"] = 30
    target["C4"] = "='Source'!D8"

    candidates = detect_debugging_repair_candidates(
        workbook,
        task_hint="Incorrect Cross Sheet References_input.xlsx",
        max_candidates=10_000,
    )
    observed = {(item.cell, item.kind, item.replacement) for item in candidates}
    assert ("C2", "cross_sheet_semantic_alignment", "=+'Source'!C8") in observed
    assert ("C3", "cross_sheet_semantic_alignment", "='Source'!B8+'Source'!B9+'Source'!B10") in observed
    assert ("C4", "cross_sheet_semantic_alignment", "='Source'!E8") in observed
    workbook.close()


def test_cross_sheet_parallel_block_requires_repeated_source_witness() -> None:
    workbook = Workbook()
    target = workbook.active
    target.title = "Model"
    source = workbook.create_sheet("WACC")
    source["C22"] = 0.08
    source["D22"] = 0.09
    target["B2"] = "WACC"
    target["C2"] = "=WACC!$D$22"
    target["J2"] = "WACC"
    target["K2"] = "=WACC!C22"
    candidates = detect_debugging_repair_candidates(
        workbook,
        task_hint="Incorrect Cross Sheet References_input.xlsx",
        max_candidates=500,
    )
    assert any(
        item.cell == "C2"
        and item.kind == "cross_sheet_parallel_block"
        and item.replacement == "=WACC!$C$22"
        for item in candidates
    )
    workbook.close()


def test_incorrect_average_candidates_are_exact_and_task_gated() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"
    worksheet["C3"] = "=AVERAGE(C1:C2)"
    worksheet["D3"] = "=SUM(360,365)"

    assert detect_debugging_repair_candidates(workbook, task_hint="other.xlsx") == []
    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Average_input.xlsx"
    )

    assert any(item.cell == "D3" and item.replacement == "=AVERAGE(360,365)" for item in candidates)
    assert any(item.cell == "C3" and item.replacement == "=AVERAGE(C1:C3)" for item in candidates)
    assert all(item.target.startswith("'Model'!") for item in candidates)


def test_embedded_hardcode_sparse_label_scan_does_not_mutate_iterator() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"
    worksheet["A3"] = "Interest expense"
    worksheet["D3"] = "=B3*0.04"
    worksheet["C5"] = 0.04

    candidates = detect_debugging_repair_candidates(
        workbook,
        task_hint="Embedded Hardcodes_input.xlsx",
        max_candidates=500,
    )

    assert isinstance(candidates, list)


def test_average_candidates_cover_whole_range_shift_and_neighbor_translation() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"
    worksheet["I7"] = "=SUM($I7,I$1)*AVERAGE(I4,I6)"
    worksheet["J7"] = "=SUM($I7,J$1)*AVERAGE(J6)"
    worksheet["D9"] = "=AVERAGE(D3:E3)"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Average_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "J7"
        and item.kind == "average_neighbor_translation"
        and item.replacement == "=SUM($I7,J$1)*AVERAGE(J4,J6)"
        for item in candidates
    )
    assert any(
        item.cell == "D9"
        and item.kind == "aggregate_range_shift"
        and item.replacement == "=AVERAGE(C3:D3)"
        for item in candidates
    )


def test_average_vertical_period_windows_extend_only_with_repeated_source_boundaries() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"
    source = workbook.create_sheet("SOFR")
    for row in range(8, 68):
        source.cell(row, 3).value = float(row)
    worksheet["Q52"] = "=AVERAGE('SOFR'!C8:C18)"
    worksheet["R52"] = "=AVERAGE('SOFR'!C20:C30)"
    worksheet["S52"] = "=AVERAGE('SOFR'!C32:C42)"
    worksheet["T52"] = "=AVERAGE('SOFR'!C44:C54)"
    worksheet["U52"] = "=AVERAGE('SOFR'!C56:C66)"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Average_input.xlsx", max_candidates=500
    )
    observed = {
        (item.cell, item.kind, item.replacement)
        for item in candidates
        if item.kind == "average_vertical_period_extension"
    }
    assert ("Q52", "average_vertical_period_extension", "=AVERAGE('SOFR'!C8:C19)") in observed
    assert ("U52", "average_vertical_period_extension", "=AVERAGE('SOFR'!C56:C67)") in observed


def test_average_context_excludes_blank_and_subject_company() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "WACC"
    source = workbook.create_sheet("Exhibit 6")
    worksheet["C6"] = "Anadarko"
    worksheet["D6"] = "Comparables"
    worksheet["D7"] = "=AVERAGE('Exhibit 6'!D36:J36)"
    source["D5"] = "Anadarko"
    source["E5"] = "Chevron"
    source["D36"] = 0.4
    source["E36"] = 0.1
    source["J36"] = 0.2
    worksheet["I15"] = "=AVERAGE($F$15:$H$15)"
    worksheet["G15"] = 0.2
    worksheet["H15"] = 0.3

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Average_input.xlsx", max_candidates=500
    )
    assert any(
        item.cell == "D7"
        and item.kind == "average_exclude_subject"
        and item.replacement == "=AVERAGE('Exhibit 6'!E36:J36)"
        for item in candidates
    )
    assert any(
        item.cell == "I15"
        and item.kind == "average_exclude_blank_endpoint"
        and item.replacement == "=AVERAGE($G$15:$H$15)"
        for item in candidates
    )


def test_average_range_shift_exposes_duplicate_neighbor_for_execution_guard() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "WACC"
    worksheet["D9"] = "=AVERAGE(H10:M10)"
    worksheet["D10"] = "=AVERAGE(H9:M9)"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Average_input.xlsx", max_candidates=500
    )
    assert any(
        item.cell == "D10"
        and item.replacement == worksheet["D9"].value
        and item.kind in {"aggregate_range_shift", "aggregate_boundary"}
        for item in candidates
    )


def test_average_array_formula_stops_before_summary_rows_after_blank_gap() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "WACC"
    source = workbook.create_sheet("Exhibit 9")
    worksheet["D14"] = ArrayFormula(
        ref="D14", text="=AVERAGE('Exhibit 9'!M8:M26/100)"
    )
    for row in range(8, 23):
        source.cell(row, 13).value = float(row)
    source["M25"] = 7.3
    source["M26"] = 8.7

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Average_input.xlsx", max_candidates=500
    )
    assert any(
        item.cell == "D14"
        and item.kind == "average_trim_summary_after_gap"
        and item.replacement == "=AVERAGE('Exhibit 9'!M8:M22/100)"
        for item in candidates
    )


def test_average_context_excludes_self_and_uses_ending_balance_for_interest() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "LBO"
    worksheet["B12"] = "Average"
    worksheet["E12"] = "=AVERAGE(E7:E12)"
    worksheet["B132"] = "Beginning Balance"
    worksheet["C132"] = "Cap:"
    worksheet["B134"] = "Paydown"
    worksheet["B135"] = "Ending Balance"
    worksheet["B136"] = "Interest"
    worksheet["G136"] = "=AVERAGE(G132,G134)*G127"
    worksheet["B180"] = "Beginning Balance"
    worksheet["B181"] = "Increase / (Decrease)"
    worksheet["B182"] = "Ending Balance"
    worksheet["B183"] = "Interest on Cash"
    worksheet["I183"] = "=H183*AVERAGE(I180:I182)"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Average_input.xlsx", max_candidates=500
    )
    assert any(
        item.cell == "E12"
        and item.kind == "average_exclude_self_reference"
        and item.replacement == "=AVERAGE(E7:E11)"
        for item in candidates
    )
    assert any(
        item.cell == "I183"
        and item.kind == "average_balance_endpoints"
        and item.replacement == "=H183*AVERAGE(I180,I182)"
        for item in candidates
    )
    assert any(
        item.cell == "G136"
        and item.kind == "average_beginning_ending_balance"
        and item.replacement == "=AVERAGE(G132,G135)*G127"
        for item in candidates
    )


def test_incorrect_average_repairs_cagr_semantics_and_header_period() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "P&L Summary"
    worksheet["M5"] = "'23A-'25A"
    worksheet["N5"] = "'26E-'30E"
    worksheet["M10"] = "=(E10-C10)/(2*C10)"
    worksheet["N13"] = "=(J13/F13)^(1/5)-1"
    worksheet["P12"] = '=+IFERROR((G12/E12-1)/(YEAR(G$7)-YEAR(E$7)),"NA")'
    worksheet["Q12"] = "=_xludf.RRI(5,W12,AA12)"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Average_input.xlsx", max_candidates=500
    )
    expected = {
        ("M10", "cagr_fixed_period", "=(E10/C10)^(1/2)-1"),
        ("N13", "cagr_period_alignment", "=(J13/F13)^(1/4)-1"),
        (
            "P12",
            "cagr_compound_years",
            '=+IFERROR((G12/E12)^(1/(YEAR(G$7)-YEAR(E$7)))-1,"NA")',
        ),
        ("Q12", "cagr_rri_native_function", "=rri(5,W12,AA12)"),
    }
    observed = {(item.cell, item.kind, item.replacement) for item in candidates}
    assert expected <= observed


def test_incorrect_average_uses_financial_table_semantics() -> None:
    workbook = Workbook()
    comps = workbook.active
    comps.title = "Comps + Beta"
    comps["B29"] = "Relevered Beta"
    comps["C30"] = "=AVERAGE(I22:I26)"
    comps["I20"] = "3yr Levered Beta"
    comps["K20"] = "Unlevered Beta"
    comps["H5"] = "N/A"
    comps["H6"] = "=E6/K6"
    comps["H7"] = "=E7/K7"
    comps["H15"] = "=MEDIAN(H7:H10)"
    dcf = workbook.create_sheet("DCF Model")
    dcf["G7"] = "Projected"
    dcf["D15"] = 0.1
    dcf["E15"] = 0.2
    dcf["F15"] = 0.3
    dcf["G15"] = "=AVERAGE(D15:E15)"
    dcf["Q94"] = 100
    dcf["R94"] = 120
    dcf["Q95"] = 0.9
    dcf["R95"] = 0.8
    lbo = workbook.create_sheet("LBO & Lenders")
    lbo["B36"] = "Exit Proceeds (Average of Exit Low / High DCF Values)"
    lbo["X36"] = "=AVERAGE('DCF Model'!Q94:R95)"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Average_input.xlsx", max_candidates=500
    )
    expected = {
        ("Comps + Beta", "C30", "=AVERAGE(K22:K26)"),
        ("Comps + Beta", "H15", "=MEDIAN(H6:H10)"),
        ("DCF Model", "G15", "=AVERAGE(D15:F15)"),
        ("LBO & Lenders", "X36", "=AVERAGE('DCF Model'!Q94:R94)"),
    }
    observed = {(item.sheet, item.cell, item.replacement) for item in candidates}
    assert expected <= observed


def test_sign_candidates_cover_two_operator_flip_and_row_peer_translation() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "LBO"
    worksheet["B4"] = "Enterprise Value"
    worksheet["C4"] = "=C1-C2+C3"
    worksheet["H7"] = "=SUM(H2:H4)"
    worksheet["I7"] = "=I2-I3-I4"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Sign Conventions_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "C4" and item.kind == "sign_operator_flip" and item.replacement == "=C1+C2-C3"
        for item in candidates
    )
    assert any(
        item.cell == "I7"
        and item.kind == "row_peer_translation"
        and item.replacement == "=SUM(I2:I4)"
        for item in candidates
    )
    assert any(
        item.cell == "I7"
        and item.kind == "sign_contiguous_sum"
        and item.replacement == "=SUM(I2:I4)"
        for item in candidates
    )


def test_sign_candidate_uses_explicit_plus_minus_line_item_labels() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "LBO"
    worksheet["B10"] = "Equity Value"
    worksheet["B11"] = "(+) Debt"
    worksheet["B12"] = "(-) Cash"
    worksheet["B13"] = "Enterprise Value"
    worksheet["C13"] = "=+C10-C11+C12"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Sign Conventions_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "C13"
        and item.kind == "label_sign_alignment"
        and item.replacement == "=+C10+C11-C12"
        for item in candidates
    )


def test_embedded_hardcode_candidate_uses_bidirectional_formula_consensus() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "DCF"
    worksheet["H5"] = "=H2*$C$1"
    worksheet["I5"] = 12.4
    worksheet["J5"] = "=J2*$C$1"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Embedded Hardcodes_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "I5"
        and item.kind == "embedded_hardcode_consensus"
        and item.current == 12.4
        and item.replacement == "=I2*$C$1"
        for item in candidates
    )


def test_embedded_hardcode_candidate_uses_two_same_direction_row_peers() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Cases"
    worksheet["F11"] = 0.12
    worksheet["G11"] = "=Financials!G10/Financials!F10-1"
    worksheet["H11"] = "=Financials!H10/Financials!G10-1"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Embedded Hardcodes_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "F11"
        and item.kind == "embedded_row_consensus"
        and item.replacement == "=Financials!F10/Financials!E10-1"
        for item in candidates
    )


def test_embedded_hardcode_candidate_restores_flat_forecast_chain() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "DCF"
    worksheet["I15"] = "=AVERAGE($G$15:$H$15)"
    for coordinate in ("J15", "K15", "L15", "M15"):
        worksheet[coordinate] = 0.124

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Embedded Hardcodes_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "J15" and item.kind == "embedded_flat_run_chain" and item.replacement == "=I15"
        for item in candidates
    )
    assert any(
        item.cell == "M15" and item.kind == "embedded_flat_run_chain" and item.replacement == "=L15"
        for item in candidates
    )


def test_double_counting_peer_translation_removes_all_repeated_terms() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Operating Model + DCF"
    worksheet["N28"] = "=N20-N23-N23-N26"
    worksheet["O28"] = "=O20-O26"
    worksheet["P28"] = "=P20-P26"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Double Counting_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "N28"
        and item.kind == "double_count_peer_translation"
        and item.replacement == "=N20-N26"
        for item in candidates
    )


def test_double_counting_peer_translation_requires_duplicate_signal() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"
    worksheet["N28"] = "=N20-N23-N26"
    worksheet["O28"] = "=O20-O26"
    worksheet["P28"] = "=P20-P26"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Double Counting_input.xlsx", max_candidates=500
    )

    assert not any(
        item.cell == "N28" and item.kind == "double_count_peer_translation"
        for item in candidates
    )


def test_embedded_hardcode_semantic_valuation_links_use_workbook_labels() -> None:
    workbook = load_workbook(
        "benchmarks/data/spreadsheetbench-v2/Debugging/spreadsheet/08_Debugging/input_files/Embedded Hardcodes_input.xlsx",
        data_only=False,
    )
    try:
        candidates = detect_debugging_repair_candidates(
            workbook, task_hint="Embedded Hardcodes_input.xlsx", max_candidates=10_000
        )
    finally:
        workbook.close()
    observed = {(item.sheet, item.cell, item.kind, item.replacement) for item in candidates}
    assert (
        "Operating Model + DCF",
        "C108",
        "embedded_exit_ebitda_multiple",
        "=INDEX('Comps + WACC'!H4:H12,MATCH(\"Median\",'Comps + WACC'!B4:B12,0))",
    ) in observed
    assert (
        "Operating Model + DCF",
        "C113",
        "embedded_net_debt_lookup",
        "=INDEX('Comps + WACC'!C39:C41,MATCH(\"Net Debt\",'Comps + WACC'!B39:B41,0))",
    ) in observed
    assert (
        "Revenue Build",
        "S6",
        "embedded_forecast_case_link",
        "=S29",
    ) in observed


def test_embedded_hardcode_candidate_shifts_absolute_source_entity_column() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Merger Model"
    worksheet["B42"] = "Stock Price"
    worksheet["C42"] = 62.4
    worksheet["D42"] = "='Exhibit 8b'!$C$14"
    exhibit = workbook.create_sheet("Exhibit 8b")
    exhibit["B14"] = 62.36000061035156
    exhibit["C14"] = 63.9900016784668

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Embedded Hardcode_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "C42"
        and item.kind == "embedded_absolute_source_column"
        and item.replacement == "='Exhibit 8b'!$B$14"
        for item in candidates
    )


def test_embedded_hardcode_candidate_reuses_label_aligned_literal_reference() -> None:
    workbook = Workbook()
    synergies = workbook.active
    synergies.title = "Synergies"
    synergies["B11"] = "(x) OXY Share Price"
    synergies["D10"] = "='Exhibit 6'!$D$33"
    synergies["D11"] = 62.4
    synergies["D12"] = "=D9*D10*D11"
    merger = workbook.create_sheet("Merger Model")
    merger["B42"] = "Stock Price"
    merger["C42"] = 62.4
    merger["D42"] = "='Exhibit 8b'!$C$14"
    exhibit = workbook.create_sheet("Exhibit 8b")
    exhibit["B14"] = 62.36000061035156
    exhibit["C14"] = 63.9900016784668

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Embedded Hardcode_input.xlsx", max_candidates=500
    )

    assert any(
        item.sheet == "Synergies"
        and item.cell == "D11"
        and item.kind == "embedded_shared_literal_reference"
        and item.replacement == "='Exhibit 8b'!$B$14"
        for item in candidates
    )


def test_embedded_formula_literal_candidate_links_same_column_assumption() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"
    worksheet["J45"] = "=-0.25*J44"
    worksheet["J46"] = 0.25

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Embedded Hardcodes_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "J45"
        and item.kind == "embedded_literal_same_column"
        and item.replacement == "=-J46*J44"
        for item in candidates
    )


def test_embedded_hardcode_candidate_extrapolates_nonunit_row_stride() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "LBO"
    worksheet["Q41"] = 0.04
    worksheet["R41"] = 0.05
    worksheet["Q42"] = "=+I22"
    worksheet["Q43"] = "=+I27"
    worksheet["R42"] = "=+J22"
    worksheet["R43"] = "=+J27"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Embedded Hardcode_input.xlsx", max_candidates=500
    )
    assert any(
        item.cell == "Q41"
        and item.kind == "embedded_matching_sequence"
        and item.replacement == "=+I17"
        for item in candidates
    )


def test_embedded_hardcode_candidate_links_lbo_multiple_and_cross_sheet_driver() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "LBO"
    worksheet["B18"] = "Entry EBITDA Multiple"
    worksheet["C18"] = "=C38"
    worksheet["B20"] = "Exit EBITDA Multiple"
    worksheet["C20"] = "=C38"
    worksheet["W18"] = "(x) Exit Multiple"
    worksheet["X17"] = "=P23"
    worksheet["Y17"] = "=Q23"
    worksheet["Z17"] = "=R23"
    worksheet["X18"] = 14.2
    worksheet["Y18"] = 14.2
    worksheet["Z18"] = 14.2
    worksheet["B78"] = "Senior Secured TLB Interest Expense"
    worksheet["B47"] = "Senior Debt Rate"
    worksheet["C47"] = 0.04
    worksheet["Q78"] = "=Q73*(0.04+Q52)"
    worksheet["R78"] = "=R73*(0.04+R52)"
    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Embedded Hardcode_input.xlsx", max_candidates=500
    )
    assert any(
        item.cell == "X18"
        and item.kind == "embedded_exit_multiple_anchor"
        and item.replacement == "=C18"
        for item in candidates
    )
    assert any(
        item.cell == "Y18"
        and item.kind == "embedded_exit_multiple_anchor"
        and item.replacement == "=$C$20"
        for item in candidates
    )
    assert any(
        item.cell == "Q78"
        and item.kind == "embedded_literal_assumption_match"
        and item.replacement == "=Q73*($C$47+Q52)"
        for item in candidates
    )


def test_double_counting_candidates_remove_direct_duplicates() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"
    worksheet["M88"] = "=SUM(M83:M87)+M84"
    worksheet["I8"] = "=SUM(I4,I6,I4)"
    worksheet["C83"] = "=C71+C74+C71-SUM(C77,C80)"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Double Counting_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "M88"
        and item.kind == "double_count_range_member"
        and item.replacement == "=SUM(M83:M87)"
        for item in candidates
    )
    assert any(
        item.cell == "C83"
        and item.kind == "double_count_direct_term"
        and item.replacement == "=C71+C74-SUM(C77,C80)"
        for item in candidates
    )
    assert any(
        item.cell == "I8"
        and item.kind == "double_count_duplicate"
        and item.replacement == "=SUM(I6,I4)"
        for item in candidates
    )


def test_double_counting_candidates_use_financial_subtotal_labels() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "LBO"
    worksheet["B64"] = "Core EBIT"
    worksheet["B57"] = "Adj. EBITDA"
    worksheet["B66"] = "EBIT"
    worksheet["B69"] = "Net Income"
    worksheet["B158"] = "Cash Interest"
    worksheet["B175"] = "Total Interest"
    worksheet["H46"] = "Revolver Interest Expense"
    worksheet["H54"] = "Ending Balance"
    worksheet["G64"] = "=G57-G62"
    worksheet["G66"] = "=+G57+G75"
    worksheet["G69"] = "=+SUM(G64:G68)"
    worksheet["G175"] = "=+G136+G137+G145+G152+G158"
    worksheet["J54"] = "=SUM(J51:J53)"
    worksheet["K54"] = "=SUM(K51:K53)+J46"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Double Counting_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "G66"
        and item.kind == "double_count_subtotal_chain"
        and item.replacement == "=+G64+G75"
        for item in candidates
    )
    assert any(
        item.cell == "G69"
        and item.kind == "double_count_subtotal_chain"
        and item.replacement == "=+SUM(G66:G68)"
        for item in candidates
    )
    assert any(
        item.cell == "G175"
        and item.kind == "double_count_derived_interest"
        and item.replacement == "=+G136+G137+G145+G152"
        for item in candidates
    )
    assert any(
        item.cell == "K54"
        and item.kind == "double_count_rollforward_interest"
        and item.replacement == "=SUM(K51:K53)"
        for item in candidates
    )


def test_double_counting_candidates_remove_cross_row_fee_reuse() -> None:
    workbook = Workbook()
    lbo = workbook.active
    lbo.title = "LBO"
    bridge = workbook.create_sheet("Valuation Bridge")
    projected = workbook.create_sheet("Projected IS")
    lbo["B13"] = "Transaction Fees"
    lbo["C13"] = 110
    lbo["B15"] = "Minimum Cash Balance"
    lbo["C15"] = 20
    lbo["E14"] = "Excess Cash"
    lbo["F14"] = "=C10-C15-C13"
    lbo["E20"] = "Transaction Fees"
    lbo["F20"] = "=+C13"
    lbo["H13"] = "EBITDA"
    lbo["J13"] = "=+'Projected IS'!D29+'Projected IS'!D32"
    projected["B29"] = "EBITDA"
    projected["B32"] = "Depreciation and amortization"
    bridge["B3"] = "Starting Equity"
    bridge["C3"] = "=LBO!F15-LBO!C13"
    bridge["B4"] = "Transaction Fees"
    bridge["C4"] = "=-LBO!C13"
    bridge["B8"] = "Exit Equity"
    bridge["C8"] = "=SUM(C3:C7)"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Double Counting_input.xlsx", max_candidates=500
    )
    observed = {
        (item.sheet, item.cell, item.kind, item.replacement) for item in candidates
    }

    assert (
        "LBO",
        "F14",
        "double_count_fee_in_cash",
        "=C10-C15",
    ) in observed
    assert (
        "Valuation Bridge",
        "C3",
        "double_count_cross_row_component",
        "=LBO!F15",
    ) in observed
    assert (
        "LBO",
        "J13",
        "double_count_embedded_subtotal_component",
        "=+'Projected IS'!D29",
    ) in observed


def test_double_counting_candidate_expands_verified_cross_sheet_total() -> None:
    workbook = Workbook()
    target = workbook.active
    target.title = "DCF"
    source = workbook.create_sheet("Balance Sheet")
    source["A20"] = "Accounts payable"
    source["A21"] = "Short-term debt"
    source["A22"] = "Other current liabilities"
    source["A23"] = "Total current liabilities"
    source["C20"] = 2164
    source["C21"] = 947
    source["C22"] = 1547
    source["C23"] = 4658
    source["C24"] = 15470
    source["C9"] = 1295
    target["R32"] = (
        "=('Balance Sheet'!C23+'Balance Sheet'!C24+'Balance Sheet'!C21"
        "-'Balance Sheet'!C20-'Balance Sheet'!C22)-'Balance Sheet'!C9"
    )
    comparables = workbook.create_sheet("WACC")
    comparables["C6"] = "Subject Co"
    comparables["D6"] = "Comparables"
    for column, company in zip(
        range(7, 14),
        ("Subject Co", "Peer A", "Peer B", "Peer C", "Peer D", "Peer E", "Peer F"),
        strict=True,
    ):
        comparables.cell(6, column).value = company
        comparables.cell(9, column).value = column / 10
    comparables["D10"] = "=AVERAGE(G9:M9)"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Double Counting_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "R32"
        and item.kind == "double_count_total_component"
        and item.replacement
        == "=('Balance Sheet'!C23+'Balance Sheet'!C24-'Balance Sheet'!C20"
        "-'Balance Sheet'!C22)-'Balance Sheet'!C9"
        for item in candidates
    )
    assert any(
        item.target == "'WACC'!D10"
        and item.kind == "average_exclude_subject"
        and item.replacement == "=AVERAGE(H9:M9)"
        for item in candidates
    )
    source["C23"] = 9999
    assert not any(
        item.kind == "double_count_total_component"
        for item in detect_debugging_repair_candidates(
            workbook, task_hint="Double Counting_input.xlsx", max_candidates=500
        )
    )


def test_double_counting_candidates_align_parallel_and_repeated_blocks() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Merger Model"
    worksheet["B21"] = "Premium Paid"
    worksheet["G21"] = "Premium Paid"
    worksheet["D21"] = "=D20-D18"
    worksheet["I21"] = "=I20-I18"
    worksheet["J21"] = "=H21+I21"
    worksheet["B25"] = "Break-up Fees"
    worksheet["D25"] = 1000
    worksheet["B29"] = "Value of NewCo equity"
    worksheet["G29"] = "Value of NewCo equity"
    worksheet["E29"] = "=+E18-E21+E23-E25-E26"
    worksheet["J29"] = "=J18+J23-J26"
    worksheet["B49"] = "Premium Paid"
    worksheet["D49"] = "=D48-D46"
    worksheet["B53"] = "Break-up Fees"
    worksheet["D53"] = 1000
    worksheet["B57"] = "Value of NewCo equity"
    worksheet["E57"] = "=+E46-E49+E51-E53-E54"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Double Counting_input.xlsx", max_candidates=500
    )
    observed = {
        (item.cell, item.kind, item.replacement)
        for item in candidates
        if item.kind == "double_count_parallel_block_term"
    }

    assert (
        "E29",
        "double_count_parallel_block_term",
        "=+E18+E23-E25-E26",
    ) in observed
    assert (
        "E57",
        "double_count_parallel_block_term",
        "=+E46+E51-E53-E54",
    ) in observed


def test_cross_sheet_candidates_use_matching_source_row_label() -> None:
    workbook = Workbook()
    target = workbook.active
    target.title = "DCF"
    source = workbook.create_sheet("Income Statement")
    target["B15"] = "EBITDA"
    target["C15"] = "='Income Statement'!J23"
    source["B23"] = "Revenue"
    source["J23"] = "=1"
    source["B25"] = "EBITDA"
    source["J25"] = "=2"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Cross Sheet References_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "C15"
        and item.kind == "cross_sheet_label_alignment"
        and item.replacement == "='Income Statement'!J25"
        for item in candidates
    )


def test_cross_sheet_candidates_keep_an_equally_specific_existing_label() -> None:
    workbook = Workbook()
    target = workbook.active
    target.title = "Cash Flow"
    source = workbook.create_sheet("Income Statement")
    target["B14"] = "Net Income"
    target["D14"] = "='Income Statement'!J37"
    source["B37"] = "Net Income to Company"
    source["J37"] = "=1"
    source["B39"] = "Net Income to Common Shareholders"
    source["J39"] = "=2"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Cross Sheet References_input.xlsx", max_candidates=500
    )

    assert not any(
        item.cell == "D14" and item.kind == "cross_sheet_label_alignment"
        for item in candidates
    )


def test_index_match_candidates_cover_exact_mode_and_source_whitespace() -> None:
    workbook = Workbook()
    model = workbook.active
    model.title = "DCF"
    source = workbook.create_sheet("Exhibit 5")
    source["A8"] = "  10 Years"
    model["C7"] = "=INDEX('Exhibit 5'!B:B,MATCH(\"10 Years\",'Exhibit 5'!A:A,1))"
    model["C8"] = "=INDEX('Exhibit 5'!B1:B12,MATCH(\"10 Years\",'Exhibit 5'!A1:A12,1))"
    model["C9"] = "=INDEX('Exhibit 5'!B1:B12,OpCase-1)"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Index Match_input.xlsx", max_candidates=500
    )

    assert any(
        item.kind == "index_match_exact_mode" and item.replacement.endswith(",0))")
        for item in candidates
    )
    assert any(
        item.cell == "C8"
        and item.kind == "index_match_exact_mode"
        and item.replacement.endswith(",0))")
        for item in candidates
    )
    assert any(
        item.cell == "C9"
        and item.kind == "index_selector_offset"
        and item.replacement == "=INDEX('Exhibit 5'!B1:B12,OpCase)"
        for item in candidates
    )
    assert any(
        item.kind == "index_match_exact_label" and 'MATCH("  10 Years"' in item.replacement
        for item in candidates
    )


def test_index_match_candidate_skips_repeated_blank_spacer_before_data() -> None:
    workbook = Workbook()
    model = workbook.active
    model.title = "WACC"
    source = workbook.create_sheet("Exhibit 6")
    for row, label in enumerate(("Debt", "Equity", "MV Leverage (%)"), start=7):
        source.cell(row, 2).value = label
        source.cell(row, 3).value = None
        source.cell(row, 4).value = row / 10
    model["C7"] = (
        "=INDEX('Exhibit 6'!C:C,MATCH(\"MV Leverage (%)\",'Exhibit 6'!B:B,0))"
    )

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Index Match_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "C7"
        and item.kind == "index_return_column_after_blank_spacer"
        and item.replacement
        == "=INDEX('Exhibit 6'!D:D,MATCH(\"MV Leverage (%)\",'Exhibit 6'!B:B,0))"
        for item in candidates
    )


def test_index_match_candidate_keeps_populated_adjacent_return_column() -> None:
    workbook = Workbook()
    model = workbook.active
    model.title = "WACC"
    source = workbook.create_sheet("Exhibit 6")
    for row, label in enumerate(("Debt", "Equity", "MV Leverage (%)"), start=7):
        source.cell(row, 2).value = label
        source.cell(row, 3).value = row / 20
        source.cell(row, 4).value = row / 10
    model["C7"] = (
        "=INDEX('Exhibit 6'!C:C,MATCH(\"MV Leverage (%)\",'Exhibit 6'!B:B,0))"
    )

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Index Match_input.xlsx", max_candidates=500
    )

    assert not any(
        item.cell == "C7" and item.kind == "index_return_column_after_blank_spacer"
        for item in candidates
    )


def test_index_match_semantic_alignment_uses_metric_labels_and_header_span() -> None:
    workbook = Workbook()
    lbo = workbook.active
    lbo.title = "Ex 1 - LBO"
    source = workbook.create_sheet("Ex 4 - Organic Operating Model")
    for column, year in enumerate(range(2022, 2031), start=3):
        source.cell(1, column).value = year
        source.cell(18, column).value = column
        source.cell(40, column).value = column * 10
    source["B18"] = "Total Revenue"
    source["B40"] = "Adjusted EBITDA"
    lbo["P1"] = 2025
    lbo["O17"] = "Total Revenue"
    lbo["P17"] = (
        "=INDEX('Ex 4 - Organic Operating Model'!C14:K18,"
        "MATCH(P1-1,'Ex 4 - Organic Operating Model'!C1:K1,0))"
    )
    lbo["B17"] = "Entry EBITDA - FY2025"
    lbo["C17"] = (
        "=INDEX('Ex 4 - Organic Operating Model'!D40:K40,"
        "MATCH(2025,'Ex 4 - Organic Operating Model'!C1:K1,0))"
    )

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Incorrect Index Match_input.xlsx", max_candidates=500
    )
    semantic = {
        item.cell: item.replacement
        for item in candidates
        if item.kind == "index_semantic_alignment"
    }
    assert semantic["P17"] == (
        "=INDEX('Ex 4 - Organic Operating Model'!C18:K18,"
        "MATCH(P1,'Ex 4 - Organic Operating Model'!C1:K1,0))"
    )
    assert semantic["C17"] == (
        "=INDEX('Ex 4 - Organic Operating Model'!C40:K40,"
        "MATCH(2025,'Ex 4 - Organic Operating Model'!C1:K1,0))"
    )


def test_relative_reference_candidates_anchor_drifting_series_to_first_peer() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "DCF"
    worksheet["G18"] = "=-G17*C8"
    worksheet["H18"] = "=-H17*D8"
    worksheet["I18"] = "=-I17*E8"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Relative vs Absolute References_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "I18"
        and item.kind == "relative_series_anchor"
        and item.replacement == "=-I17*$C$8"
        for item in candidates
    )


def test_relative_reference_candidates_release_overanchored_parallel_series() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "DCF"
    for column, year in zip("CDEFGH", range(1, 7), strict=True):
        worksheet[f"{column}97"] = year
        worksheet[f"{column}98"] = 0.1
        worksheet[f"{column}99"] = f"=(1+{column}98)^$C$97"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Relative vs Absolute References_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "D99"
        and item.kind == "relative_release_series_anchor"
        and item.replacement == "=(1+D98)^D$97"
        for item in candidates
    )


def test_unit_mismatch_candidates_cover_percent_days_and_growth() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"
    worksheet["B4"] = "Tax Rate"
    worksheet["C4"] = 5.3
    worksheet["C4"].number_format = "0.0%"
    worksheet["B8"] = "Accounts Receivable Days"
    worksheet["J8"] = "=(J7*J6)/12"
    worksheet["B14"] = "Revenue Growth"
    worksheet["D14"] = "=D13/C13"

    candidates = detect_debugging_repair_candidates(
        workbook, task_hint="Unit Mismatch_input.xlsx", max_candidates=500
    )

    assert any(
        item.cell == "C4" and item.kind == "unit_percent_scale" and item.replacement == "=5.3/100"
        for item in candidates
    )
    assert any(
        item.cell == "J8" and item.kind == "unit_days_per_year" and "/365" in item.replacement
        for item in candidates
    )
    assert any(
        item.cell == "D14" and item.kind == "unit_growth_rate" and item.replacement.endswith("-1")
        for item in candidates
    )
