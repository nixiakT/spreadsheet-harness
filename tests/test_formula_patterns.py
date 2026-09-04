from __future__ import annotations

from openpyxl import Workbook

from spreadsheet_harness.formula_patterns import (
    detect_formula_pattern_repairs,
    select_safe_formula_pattern_repairs,
)


def test_detects_single_horizontal_formula_pattern_conflict() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"
    worksheet["B2"] = "=B1*2"
    worksheet["C2"] = "=0"
    worksheet["D2"] = "=D1*2"

    repairs = detect_formula_pattern_repairs(workbook)

    assert len(repairs) == 1
    assert repairs[0].sheet == "Model"
    assert repairs[0].cell == "C2"
    assert repairs[0].replacement == "=C1*2"
    assert repairs[0].directions == ("horizontal",)
    assert select_safe_formula_pattern_repairs(repairs) == repairs


def test_rejects_conflicting_horizontal_and_vertical_replacements() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["B2"] = "=B1*2"
    worksheet["C2"] = "=0"
    worksheet["D2"] = "=D1*2"
    worksheet["C1"] = "=A1"
    worksheet["C3"] = "=A3"

    repairs = detect_formula_pattern_repairs(workbook)

    assert repairs == []


def test_multiple_single_direction_conflicts_are_not_selected() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    for row in (2, 4):
        worksheet.cell(row, 2, f"=B{row - 1}*2")
        worksheet.cell(row, 3, "=0")
        worksheet.cell(row, 4, f"=D{row - 1}*2")

    repairs = detect_formula_pattern_repairs(workbook)

    assert len(repairs) == 2
    assert select_safe_formula_pattern_repairs(repairs) == []


def test_out_of_range_neighbor_translation_is_ignored() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["B2"] = "=A1"
    worksheet["C2"] = "=0"
    worksheet["D2"] = "=A1"

    assert detect_formula_pattern_repairs(workbook) == []


def test_chartsheets_are_ignored() -> None:
    workbook = Workbook()
    workbook.create_chartsheet("Dashboard")

    assert detect_formula_pattern_repairs(workbook) == []


def test_invalid_ref_translation_replacement_is_never_selected() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["B2"] = "='Data'!#REF!"
    worksheet["C2"] = "='Data'!#REF!"
    worksheet["D2"] = "='Data'!#REF!"

    repairs = detect_formula_pattern_repairs(workbook)

    assert all(repair.replacement == "=" for repair in repairs)
    assert select_safe_formula_pattern_repairs(repairs) == []


def test_unique_sign_only_repair_requires_matching_task_hint() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["B2"] = "=-B1"
    worksheet["C2"] = "=C1"
    worksheet["D2"] = "=-D1"
    worksheet["B4"] = "=B3*2"
    worksheet["C4"] = "=0"
    worksheet["D4"] = "=D3*2"

    repairs = detect_formula_pattern_repairs(workbook)

    assert len(repairs) == 2
    assert select_safe_formula_pattern_repairs(repairs) == []
    selected = select_safe_formula_pattern_repairs(
        repairs, task_hint="Incorrect Sign Conventions_input.xlsx"
    )
    assert [(repair.cell, repair.replacement) for repair in selected] == [("C2", "=-C1")]
