import io
import shutil
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from spreadsheet_harness.financial_model_repairs import (
    _find_row_by_label,
    _instruction_target_columns,
    _instruction_year_columns,
    complete_financial_model_runtime_actions,
    complete_isolated_formula_holes,
    complete_revenue_growth_schedule,
)


def _rewrite_zip_member(path: Path, member_name: str, callback) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as source:
        with zipfile.ZipFile(buffer, "w") as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename == member_name:
                    payload = callback(payload)
                target.writestr(info, payload)
    path.write_bytes(buffer.getvalue())


def _schedule(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Revenue"
    worksheet.append([None, "Metric", "Year 1", "Year 2"])
    rows = [
        ("CPI Growth Rate (%)", 0.02, 0.03),
        ("Price Growth Over CPI (%)", 0.01, 0.01),
        ("Population Growth (%)", 0.005, 0.005),
        ("Market Share - Beginning", 0.3, None),
        ("Market Share - Target Growth", 0.02, 0.01),
        ("Historical Revenue (Year 0)", 100, None),
        ("Total Price Growth Rate (%)", None, None),
        ("Market Share - Ending", None, None),
        ("Share Growth Rate (%)", None, None),
        ("Volume Growth Rate (%)", None, None),
        ("Total Revenue Growth Rate (%)", None, None),
        ("Projected Revenue", None, None),
    ]
    for label, first, second in rows:
        worksheet.append([None, label, first, second])
    workbook.save(path)
    workbook.close()


def test_financial_instruction_scans_ignore_format_only_xfd_extent() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["B2"] = "Revenue"
    worksheet["C1"] = 2024
    worksheet["D1"] = 2025
    worksheet["XFD60"].number_format = "0.0%"
    assert worksheet.max_column == 16384

    original_cell = worksheet.cell
    accessed_columns: list[int] = []

    def tracked_cell(row: int, column: int, *args, **kwargs):
        accessed_columns.append(column)
        return original_cell(row, column, *args, **kwargs)

    worksheet.cell = tracked_cell

    assert _find_row_by_label(worksheet, "Revenue") == 2
    assert max(accessed_columns) <= 8
    accessed_columns.clear()
    assert _instruction_year_columns(worksheet) == {3: 2024, 4: 2025}
    assert _instruction_target_columns(worksheet, "2024-2025") == [3, 4]
    assert max(accessed_columns) == 4
    workbook.close()


def test_complete_revenue_growth_schedule_fills_only_calculation_rows(tmp_path: Path) -> None:
    path = tmp_path / "revenue.xlsx"
    _schedule(path)

    changes = complete_revenue_growth_schedule(path)

    workbook = load_workbook(path, data_only=False)
    worksheet = workbook["Revenue"]
    assert len(changes) == 13
    assert worksheet["C8"].value == "=C2+C3"
    assert worksheet["C9"].value == "=C5*(1+C6)"
    assert worksheet["D5"].value == "=C9"
    assert worksheet["D9"].value == "=D5*(1+D6)"
    assert worksheet["C10"].value == "=C9/C5-1"
    assert worksheet["D10"].value == "=D9/D5-1"
    assert worksheet["C11"].value == "=(1+C4)*(1+C10)-1"
    assert worksheet["C12"].value == "=(1+C8)*(1+C11)-1"
    assert worksheet["C13"].value == "=C7*(1+C12)"
    assert worksheet["D13"].value == "=C13*(1+D12)"
    assert worksheet["C5"].value == 0.3
    workbook.close()


def test_complete_revenue_growth_schedule_fails_closed_on_populated_target(tmp_path: Path) -> None:
    path = tmp_path / "revenue.xlsx"
    _schedule(path)
    workbook = load_workbook(path)
    workbook["Revenue"]["C8"] = 123
    workbook.save(path)
    workbook.close()

    assert complete_revenue_growth_schedule(path) == []
    unchanged = load_workbook(path, data_only=False)
    assert unchanged["Revenue"]["C8"].value == 123
    assert unchanged["Revenue"]["C9"].value is None
    unchanged.close()


def test_complete_revenue_growth_schedule_repairs_unbound_prefix_metadata(
    tmp_path: Path,
) -> None:
    path = tmp_path / "revenue.xlsx"
    _schedule(path)

    _rewrite_zip_member(
        path,
        "docProps/core.xml",
        lambda payload: payload.decode("utf-8").replace(
            ' xmlns:dc="http://purl.org/dc/elements/1.1/"', "", 1
        ).encode("utf-8"),
    )

    changes = complete_revenue_growth_schedule(path)

    workbook = load_workbook(path, data_only=False)
    assert len(changes) == 13
    assert workbook["Revenue"]["C8"].value == "=C2+C3"
    workbook.close()


def test_complete_isolated_formula_holes_uses_bidirectional_translation(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Distribution"
    for row in range(2, 7):
        worksheet.cell(row, 2, f"Tier {row}")
        worksheet.cell(row, 3, row)
    workbook.save(source)
    worksheet["D2"] = "=Dashboard!$H$9"
    worksheet["D3"] = "=Dashboard!$H$9"
    worksheet["D5"] = "=Dashboard!$H$9"
    worksheet["D6"] = "=Dashboard!$H$9"
    worksheet["E4"] = "=1"
    worksheet["F4"] = "=1"
    worksheet["G4"] = "=1"
    workbook.save(output)
    workbook.close()

    changes = complete_isolated_formula_holes(output, source_path=source)

    repaired = load_workbook(output, data_only=False)
    assert changes == [
        {
            "sheet": "Distribution",
            "target": "D4",
            "formula": "=Dashboard!$H$9",
        }
    ]
    assert repaired["Distribution"]["D4"].value == "=Dashboard!$H$9"
    repaired.close()


def test_complete_isolated_formula_holes_repairs_unbound_prefix_metadata(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Distribution"
    for row in range(2, 7):
        worksheet.cell(row, 2, f"Tier {row}")
        worksheet.cell(row, 3, row)
    workbook.save(source)
    worksheet["D2"] = "=Dashboard!$H$9"
    worksheet["D3"] = "=Dashboard!$H$9"
    worksheet["D5"] = "=Dashboard!$H$9"
    worksheet["D6"] = "=Dashboard!$H$9"
    worksheet["E4"] = "=1"
    worksheet["F4"] = "=1"
    worksheet["G4"] = "=1"
    workbook.save(output)
    workbook.close()

    def _break_core(payload: bytes) -> bytes:
        return payload.decode("utf-8").replace(
            ' xmlns:dc="http://purl.org/dc/elements/1.1/"', "", 1
        ).encode("utf-8")

    _rewrite_zip_member(source, "docProps/core.xml", _break_core)
    _rewrite_zip_member(output, "docProps/core.xml", _break_core)

    changes = complete_isolated_formula_holes(output, source_path=source)

    repaired = load_workbook(output, data_only=False)
    assert changes == [
        {
            "sheet": "Distribution",
            "target": "D4",
            "formula": "=Dashboard!$H$9",
        }
    ]
    assert repaired["Distribution"]["D4"].value == "=Dashboard!$H$9"
    repaired.close()


def test_complete_isolated_formula_holes_fills_horizontal_subtotal_gap(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Annuals"
    worksheet["B4"] = "Subtotal"
    worksheet["D4"] = "=SUM(D1:D3)"
    worksheet["F4"] = "=SUM(F1:F3)"
    worksheet["H4"] = "=SUM(H1:H3)"
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    changes = complete_isolated_formula_holes(output, source_path=source)

    repaired = load_workbook(output, data_only=False)
    assert {"sheet": "Annuals", "target": "E4", "formula": "=SUM(E1:E3)"} in changes
    assert repaired["Annuals"]["E4"].value == "=SUM(E1:E3)"
    repaired.close()


def test_complete_isolated_formula_holes_extends_rightward_same_row_subtotals(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Forecast"
    for row in range(2, 5):
        for column in range(4, 11):
            worksheet.cell(row, column, row * column)
    worksheet["D5"] = "=SUM(D2:D4)"
    worksheet["E5"] = "=SUM(E2:E4)"
    worksheet["F5"] = "=SUM(F2:F4)"
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    changes = complete_isolated_formula_holes(output, source_path=source)

    repaired = load_workbook(output, data_only=False)
    assert {"sheet": "Forecast", "target": "G5", "formula": "=SUM(G2:G4)"} in changes
    assert {"sheet": "Forecast", "target": "J5", "formula": "=SUM(J2:J4)"} in changes
    assert repaired["Forecast"]["G5"].value == "=SUM(G2:G4)"
    assert repaired["Forecast"]["J5"].value == "=SUM(J2:J4)"
    repaired.close()


def test_complete_financial_model_runtime_actions_applies_constants_and_freeze_panes(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    dashboard["B12"] = "Management Fee"
    for row, year in enumerate(range(2025, 2028), start=13):
        dashboard[f"B{row}"] = year if row == 13 else f"=B{row - 1}+1"
        dashboard[f"C{row}"] = "%age"
    dashboard["B23"] = "Portfolio"
    for index, row in enumerate(range(24, 27), start=1):
        dashboard[f"B{row}"] = f"Investment Number {index}"
        dashboard[f"C{row}"] = "Number "
        dashboard[f"D{row}"] = 100 + index
    dashboard["E23"] = "Numbers"
    revenue = workbook.create_sheet("Revenue")
    revenue["A1"] = "Revenue"
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    instruction = (
        "In the Dashboard, insert management fees at 1.75% for 2025-2027, "
        "and insert first tranche investment numbers at 1 per deal for 3 cumulative deals. "
        "In the Revenue sheet, freeze rows 1-4 and columns A-C."
    )
    changes = complete_financial_model_runtime_actions(
        output,
        source_path=source,
        instruction=instruction,
    )

    repaired = load_workbook(output, data_only=False)
    assert {"sheet": "Dashboard", "target": "D13", "value": "0.0175"} in changes
    assert {"sheet": "Dashboard", "target": "E24", "value": "1"} in changes
    assert repaired["Dashboard"]["D13"].value == 0.0175
    assert repaired["Dashboard"]["E24"].value == 1
    assert repaired["Revenue"].freeze_panes == "D5"
    repaired.close()


def test_complete_financial_model_runtime_actions_continues_metric_rows_and_links_peer_row(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    income = workbook.active
    income.title = "IS - Mgmt Co."
    income["B57"] = "EBITDA"
    income["N57"] = "=N13-N55"
    income["B59"] = "Depreciation"
    income["N59"] = 5
    income["O59"] = 6
    income["P59"] = 7
    income["Q59"] = 8
    income["B66"] = "Net Income"
    income["N66"] = "=N61-N63-N64"
    income["O66"] = "=O61-O63-O64"
    income["P66"] = "=P61-P63-P64"
    income["Q66"] = "=Q61-Q63-Q64"
    cashflow = workbook.create_sheet("CF")
    cashflow["B10"] = "+ Net Income"
    cashflow["M11"] = "=+'IS - Mgmt Co.'!N59"
    cashflow["N11"] = "=+'IS - Mgmt Co.'!O59"
    cashflow["O11"] = "=+'IS - Mgmt Co.'!P59"
    cashflow["P11"] = "=+'IS - Mgmt Co.'!Q59"
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    instruction = (
        "In the CF sheet, link Net Income from IS - Mgmt Co. including Pre-Op Cost. "
        "In the IS - Mgmt Co. sheet, compute EBITDA."
    )
    changes = complete_financial_model_runtime_actions(
        output,
        source_path=source,
        instruction=instruction,
    )

    repaired = load_workbook(output, data_only=False)
    assert {"sheet": "IS - Mgmt Co.", "target": "O57", "formula": "=O13-O55"} in changes
    assert {"sheet": "CF", "target": "M10", "formula": "='IS - Mgmt Co.'!N66"} in changes
    assert repaired["IS - Mgmt Co."]["Q57"].value == "=Q13-Q55"
    assert repaired["CF"]["P10"].value == "='IS - Mgmt Co.'!Q66"
    repaired.close()


def test_complete_financial_model_runtime_actions_skips_merged_metric_targets(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    dashboard["A55"] = "Investment Exit Value"
    dashboard["H55"] = "=H40"
    dashboard.merge_cells("H55:I55")
    dashboard["I40"] = 123
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    instruction = "In the Dashboard, compute investment exit value for each tranche."
    changes = complete_financial_model_runtime_actions(
        output,
        source_path=source,
        instruction=instruction,
    )

    repaired = load_workbook(output, data_only=False)
    assert changes == []
    assert repaired["Dashboard"]["H55"].value == "=H40"
    repaired.close()


def test_complete_isolated_formula_holes_ignores_styled_excel_extent(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A3"] = "Metric"
    for coordinate, formula in (("B3", "=B2*2"), ("D3", "=D2*2"), ("E3", "=E2*2")):
        worksheet[coordinate] = formula
    for coordinate in ("B3", "C3", "D3", "E3"):
        worksheet[coordinate].number_format = "0.0"
    worksheet["XFD1048576"].number_format = "0.0"
    workbook.save(source)
    workbook.close()
    shutil.copy2(source, output)

    changes = complete_isolated_formula_holes(output, source_path=source)

    assert changes == [{"sheet": "Sheet", "target": "C3", "formula": "=C2*2"}]
    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook.active["C3"].value == "=C2*2"
    finally:
        workbook.close()


def test_complete_financial_model_runtime_actions_fills_ticket_sizes_and_exit_values(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    dashboard["G14"] = "Ticket Size per Investment (Unit Economics) (USD Mn)"
    dashboard["G15"] = "First Tranche"
    dashboard["G16"] = "Second Tranche"
    dashboard["G17"] = "Third Tranche"
    dashboard["D22"] = "First Tranche"
    dashboard["F22"] = "Second Tranche"
    dashboard["H22"] = "Third Tranche"
    dashboard["B23"] = "Portfolio Building MoM"
    dashboard["C23"] = "UOM"
    dashboard["D23"] = "Month"
    dashboard["E23"] = "Numbers"
    dashboard["F23"] = "Month"
    dashboard["G23"] = "Numbers"
    dashboard["H23"] = "Month"
    dashboard["I23"] = "Numbers"
    dashboard["B24"] = "Investment Number 1"
    dashboard["E24"] = 1
    dashboard["G24"] = 1
    dashboard["I24"] = 1
    dashboard["B43"] = "Investment 1 - NCLT Case"
    dashboard["E43"] = 2.5
    dashboard["G43"] = 2
    dashboard["I43"] = 2
    dashboard["B55"] = "Investment Exit Value"
    dashboard["C56"] = "UOM"
    dashboard["D56"] = "Month"
    dashboard["E56"] = "Numbers"
    dashboard["F56"] = "Month"
    dashboard["G56"] = "Numbers"
    dashboard["H56"] = "Month"
    dashboard["I56"] = "Numbers"
    dashboard["B57"] = "Investment Number 1"
    dashboard["D57"] = "=D43"
    dashboard["F57"] = "=F43"
    dashboard["H57"] = "=H43"
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    instruction = (
        "In the Dashboard, update ticket size per investment for first, second, and third "
        "tranches at 20, 10, and 10 in light blue, and compute investment exit value for "
        "each tranche."
    )
    changes = complete_financial_model_runtime_actions(
        output,
        source_path=source,
        instruction=instruction,
    )

    repaired = load_workbook(output, data_only=False)
    assert {"sheet": "Dashboard", "target": "H15", "value": "20"} in changes
    assert {"sheet": "Dashboard", "target": "H16", "value": "10"} in changes
    assert {"sheet": "Dashboard", "target": "H17", "value": "10"} in changes
    assert {"sheet": "Dashboard", "target": "E57", "formula": "=E43*E24*$H$15"} in changes
    assert {"sheet": "Dashboard", "target": "G57", "formula": "=G43*G24*$H$16"} in changes
    assert {"sheet": "Dashboard", "target": "I57", "formula": "=I43*I24*$H$17"} in changes
    assert repaired["Dashboard"]["H15"].value == 20
    assert repaired["Dashboard"]["E57"].value == "=E43*E24*$H$15"
    assert repaired["Dashboard"]["G57"].value == "=G43*G24*$H$16"
    assert repaired["Dashboard"]["I57"].value == "=I43*I24*$H$17"
    repaired.close()


def test_complete_financial_model_runtime_actions_links_header_column_from_dashboard(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    dashboard["G9"] = "Hurdle Rate"
    dashboard["H9"] = 0.12
    distribution = workbook.create_sheet("Distribution Sheet-Deal By Deal")
    distribution["I4"] = "Hurdle Rate"
    distribution["C5"] = "Investment Number 1"
    distribution["H5"] = "=Dashboard!E57"
    distribution["J5"] = "=G5"
    distribution["C6"] = "Investment Number 2"
    distribution["H6"] = "=Dashboard!E58"
    distribution["J6"] = "=G6"
    distribution["C8"] = "=Dashboard!F22"
    distribution["J8"] = "=G8"
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    instruction = "In the Distribution Sheet-Deal By Deal, link hurdle rate from Dashboard."
    changes = complete_financial_model_runtime_actions(
        output,
        source_path=source,
        instruction=instruction,
    )

    repaired = load_workbook(output, data_only=False)
    assert {
        "sheet": "Distribution Sheet-Deal By Deal",
        "target": "I5",
        "formula": "=Dashboard!$H$9",
    } in changes
    assert {
        "sheet": "Distribution Sheet-Deal By Deal",
        "target": "I6",
        "formula": "=Dashboard!$H$9",
    } in changes
    assert {
        "sheet": "Distribution Sheet-Deal By Deal",
        "target": "I8",
        "formula": "=Dashboard!$H$9",
    } in changes
    assert repaired["Distribution Sheet-Deal By Deal"]["I5"].value == "=Dashboard!$H$9"
    assert repaired["Distribution Sheet-Deal By Deal"]["I6"].value == "=Dashboard!$H$9"
    assert repaired["Distribution Sheet-Deal By Deal"]["I8"].value == "=Dashboard!$H$9"
    repaired.close()


def test_complete_financial_model_runtime_actions_continues_metric_rows_across_spacer_gap(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    workings = workbook.active
    workings.title = "Workings Cost Sheet"
    workings["B204"] = "Opening"
    workings["B205"] = "Capex"
    workings["B206"] = "Renewal Capex"
    workings["B207"] = "Depreciation"
    workings["B208"] = "Ending PP&E"
    for column in "DEFGHIJKL":
        workings[f"{column}208"] = f"={column}204+{column}205+{column}206-{column}207"
    workings["N207"] = 0.125
    workings["O204"] = 0
    workings["P204"] = "=O208"
    workings["Q204"] = "=P208"
    workings["O205"] = 10
    workings["P205"] = 20
    workings["Q205"] = 30
    workings["O206"] = 0
    workings["P206"] = 0
    workings["Q206"] = 0
    workings["O207"] = "=SUM(O204:O206)*$N$207"
    workings["P207"] = "=SUM(P204:P206)*$N$207"
    workings["Q207"] = "=SUM(Q204:Q206)*$N$207"
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    instruction = "In the Workings Cost sheet, compute ending PP&E from Apr-25 to Jun-33."
    changes = complete_financial_model_runtime_actions(
        output,
        source_path=source,
        instruction=instruction,
    )

    repaired = load_workbook(output, data_only=False)
    assert {"sheet": "Workings Cost Sheet", "target": "O208", "formula": "=O204+O205+O206-O207"} in changes
    assert {"sheet": "Workings Cost Sheet", "target": "Q208", "formula": "=Q204+Q205+Q206-Q207"} in changes
    assert repaired["Workings Cost Sheet"]["M208"].value is None
    assert repaired["Workings Cost Sheet"]["N208"].value is None
    assert repaired["Workings Cost Sheet"]["O208"].value == "=O204+O205+O206-O207"
    assert repaired["Workings Cost Sheet"]["Q208"].value == "=Q204+Q205+Q206-Q207"
    repaired.close()


def test_complete_financial_model_runtime_actions_links_metric_rows_without_explicit_source(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    cashflow = workbook.active
    cashflow.title = "CF"
    cashflow["B13"] = "- Change in Working Capital"
    cashflow["M13"] = "=M34"
    cashflow["N34"] = 10
    cashflow["O34"] = 11
    cashflow["P34"] = 12
    cashflow["Q34"] = 13
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    instruction = "In the CF sheet, link Change in Working Capital."
    changes = complete_financial_model_runtime_actions(
        output,
        source_path=source,
        instruction=instruction,
    )

    repaired = load_workbook(output, data_only=False)
    assert {"sheet": "CF", "target": "N13", "formula": "=N34"} in changes
    assert {"sheet": "CF", "target": "Q13", "formula": "=Q34"} in changes
    assert repaired["CF"]["N13"].value == "=N34"
    assert repaired["CF"]["Q13"].value == "=Q34"
    repaired.close()


def test_complete_financial_model_runtime_actions_fills_later_annual_blocks_and_subtotals(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    expenses = workbook.active
    expenses.title = "Other Expenses"
    expenses["B5"] = "Expenses"
    for coordinate, value in {
        "D5": "M0",
        "E5": "M1",
        "F5": "M2",
        "G5": "M3",
        "H5": "M4",
        "I5": "M5",
        "J5": "M6",
        "K5": "M7",
        "L5": "M8",
        "M5": "M9",
        "N5": "M10",
        "O5": "M11",
        "P5": "M12",
        "Q5": "Y1",
        "R5": "Q1",
        "S5": "Q2",
        "T5": "Q3",
        "U5": "Q4",
        "V5": "Y2",
        "W5": "Q1",
        "X5": "Q2",
        "Y5": "Q3",
        "Z5": "Q4",
        "AA5": "Y3",
    }.items():
        expenses[coordinate] = value
    expenses["B6"] = "Rent"
    expenses["Q6"] = "=SUM(D6:P6)"
    expenses["R6"] = 1
    expenses["S6"] = 2
    expenses["T6"] = 3
    expenses["U6"] = 4
    expenses["W6"] = 5
    expenses["X6"] = 6
    expenses["Y6"] = 7
    expenses["Z6"] = 8
    expenses["B7"] = "Subtotal"
    expenses["Q7"] = "=SUM(Q6:Q6)"
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    instruction = "In the Other Expenses sheet, compute total expenses for Y2 and Y3."
    changes = complete_financial_model_runtime_actions(
        output,
        source_path=source,
        instruction=instruction,
    )

    repaired = load_workbook(output, data_only=False)
    assert {"sheet": "Other Expenses", "target": "V6", "formula": "=SUM(R6:U6)"} in changes
    assert {"sheet": "Other Expenses", "target": "AA6", "formula": "=SUM(W6:Z6)"} in changes
    assert {"sheet": "Other Expenses", "target": "V7", "formula": "=SUM(V6:V6)"} in changes
    assert {"sheet": "Other Expenses", "target": "AA7", "formula": "=SUM(AA6:AA6)"} in changes
    assert repaired["Other Expenses"]["V6"].value == "=SUM(R6:U6)"
    assert repaired["Other Expenses"]["AA6"].value == "=SUM(W6:Z6)"
    assert repaired["Other Expenses"]["V7"].value == "=SUM(V6:V6)"
    assert repaired["Other Expenses"]["AA7"].value == "=SUM(AA6:AA6)"
    repaired.close()


def test_complete_financial_model_runtime_actions_fills_explicit_calculation_rows(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    wc = workbook.active
    wc.title = "Working Capital Schedule"
    wc["B3"] = "Particulars"
    wc["C3"] = 2021
    for column in range(4, 13):
        wc.cell(3, column).value = (
            f"=EOMONTH({get_column_letter(column - 1)}3,12)"
        )
    wc["M21"] = "=M24"
    wc["M22"] = "=L21-M21"
    wc["M24"] = "=L24"
    wc["M3"]._style = wc["L3"]._style
    for row, label in ((4, "Receivable"), (5, "Revenue"), (6, "Receivable Days"), (8, "Inventory"), (16, "Other Current Asset"), (19, "Current Assets"), (20, "Current Liabilities"), (21, "Total Working Capital")):
        wc.cell(row, 2).value = label
    for column in range(3, 13):
        wc.cell(5, column).value = 100
        wc.cell(6, column).value = (
            f"={get_column_letter(column)}4/{get_column_letter(column)}5*365"
            if column <= 7
            else 36.5
        )
        wc.cell(8, column).value = 10
        wc.cell(16, column).value = 5
        wc.cell(20, column).value = 3
    balance = workbook.create_sheet("Consolidated BS")
    balance["B15"] = "Trade Receivables"
    balance["C4"] = 2021
    for column in range(4, 13):
        balance.cell(4, column).value = (
            f"=EOMONTH({get_column_letter(column - 1)}4,12)"
        )
    for column in range(3, 8):
        balance.cell(15, column).value = 10
    fixed = workbook.create_sheet("Fixed Assets Schedule")
    fixed["B6"] = "Capex (maintenance and investment)"
    fixed["B11"] = "Annual Maintenance Capex (% of revenue)"
    for column in range(3, 13):
        fixed.cell(3, column).value = f"='Consolidated P&L'!{chr(64 + column)}4"
    fixed["C11"] = 0.01
    consolidated = workbook.create_sheet("Consolidated P&L")
    consolidated["B5"] = "Revenue"
    consolidated["B9"] = "Total Revenue"
    for column in range(3, 13):
        consolidated.cell(5, column).value = 90
        consolidated.cell(9, column).value = 100
    debt = workbook.create_sheet("Debt Schedule")
    for row, label in ((6, "Term Loan"), (8, "Opening Balance"), (9, "Addition"), (10, "Repayment"), (11, "Outstanding")):
        debt.cell(row, 2).value = label
    for column in range(3, 13):
        debt.cell(8, column).value = 100
        debt.cell(9, column).value = 0
        debt.cell(10, column).value = 10
    pnl = workbook.create_sheet("P&L - Segment 1")
    pnl["B6"] = "Total Revenue"
    pnl["B29"] = "PBT"
    pnl["B31"] = "PBT Margin"
    for column in range(3, 8):
        pnl.cell(6, column).value = 100
        pnl.cell(29, column).value = 20
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    instruction = (
        "In the Working Capital Schedule sheet, calculate Receivables for 2026E–2030E using Revenue and Receivable Days, "
        "then calculate Current Assets, then calculate Total Working Capital for all years. "
        "In the Fixed Assets Schedule sheet, calculate Capex for 2026E–2030E. "
        "In the Debt Schedule sheet, calculate the Closing Balance of the Term Loan for 2026E–2030E. "
        "In the P&L – Segment 1 sheet, calculate PBT Margin for 2021A–2025A."
    )
    changes = complete_financial_model_runtime_actions(output, source_path=source, instruction=instruction)
    repaired = load_workbook(output, data_only=False)
    assert repaired["Working Capital Schedule"]["M3"].value == "=EOMONTH(L3,12)"
    assert repaired["Working Capital Schedule"]["C4"].value == "='Consolidated BS'!C15"
    assert repaired["Working Capital Schedule"]["G4"].value == "='Consolidated BS'!G15"
    assert repaired["Working Capital Schedule"]["H4"].value == "=H5*H6/365"
    assert repaired["Working Capital Schedule"]["C19"].value == "=C4+C8+C16"
    assert repaired["Working Capital Schedule"]["L21"].value == "=L19-L20"
    assert repaired["Fixed Assets Schedule"]["H6"].value == "='Consolidated P&L'!H5*$C$11"
    assert repaired["Debt Schedule"]["H11"].value == "=H8+H9-H10"
    assert repaired["P&L - Segment 1"]["G31"].value == "=G29/G6"
    assert len(changes) >= 30
    repaired.close()


def test_explicit_capex_includes_declared_annual_investment_series(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    fixed = workbook.active
    fixed.title = "Fixed Assets Schedule"
    fixed["B6"] = "Capex (maintenance and investment)"
    fixed["B11"] = "Annual Maintenance Capex (% of revenue)"
    fixed["C11"] = 0.01
    for column in range(3, 13):
        fixed.cell(3, column).value = f"='Consolidated P&L'!{get_column_letter(column)}4"
    consolidated = workbook.create_sheet("Consolidated P&L")
    consolidated["B5"] = "Revenue"
    consolidated["B9"] = "Total Revenue"
    for column in range(3, 13):
        consolidated.cell(5, column).value = 90
        consolidated.cell(9, column).value = 100
    details = workbook.create_sheet("Investment Capex Details")
    details["A5"] = "Total Investment Capex"
    for column, value in enumerate((10, 20, 30, 40, 50), start=2):
        details.cell(5, column).value = value
    workbook.save(source)
    workbook.save(output)
    workbook.close()

    changes = complete_financial_model_runtime_actions(
        output,
        source_path=source,
        instruction=(
            "In the Fixed Assets Schedule sheet, calculate Capex for 2026E-2030E."
        ),
    )

    repaired = load_workbook(output, data_only=False)
    assert repaired["Fixed Assets Schedule"]["H6"].value == (
        "='Consolidated P&L'!H5*$C$11+'Investment Capex Details'!B5"
    )
    assert repaired["Fixed Assets Schedule"]["L6"].value == (
        "='Consolidated P&L'!L5*$C$11+'Investment Capex Details'!F5"
    )
    assert {
        "sheet": "Fixed Assets Schedule",
        "target": "H6",
        "formula": "='Consolidated P&L'!H5*$C$11+'Investment Capex Details'!B5",
    } in changes
    repaired.close()


def test_education_model_assumption_cost_depreciation_and_check_rules(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()

    pnl = workbook.active
    pnl.title = "Consolidated P&L"
    pnl["B5"] = "Revenue"
    pnl["B6"] = "Other Revenue"
    pnl["B9"] = "Total Revenue"
    pnl["M4"] = "CAGR (25A-30E)"
    for column, year in enumerate(range(2021, 2031), start=3):
        pnl.cell(4, column).value = year
        pnl.cell(5, column).value = 100 + column
        pnl.cell(6, column).value = 2
        pnl.cell(9, column).value = 102 + column
    for column in range(8, 13):
        target = get_column_letter(column + 1)
        pnl.cell(5, column).value = f"='Revenue Drivers'!{target}135"
        pnl.cell(6, column).value = f"='Revenue Drivers'!{target}137"

    revenue = workbook.create_sheet("Revenue Drivers")
    for column, year in enumerate(range(2021, 2031), start=4):
        revenue.cell(4, column).value = year
    revenue["B7"] = "Mature"
    revenue["B8"] = "Headcount"
    revenue["B13"] = "Average Ticket Size"
    revenue["B17"] = "Total Revenue from Foundation Course Mature"
    revenue["B135"] = "Total Operating Revenue"
    revenue["B137"] = "Other Revenue"
    revenue["C137"] = "USD Mn"
    revenue["B138"] = "As a % of Operating Revenue"
    revenue["C138"] = "%"
    for column in range(9, 14):
        revenue.cell(17, column).value = 100
        revenue.cell(135, column).value = 1000

    assumptions = workbook.create_sheet("Key Assumptions")
    assumptions["C9"] = "Other Income as % of Revenue"
    assumptions["D9"] = 0.02

    costs = workbook.create_sheet("Cost Drivers")
    for column, year in enumerate(range(2021, 2031), start=4):
        costs.cell(4, column).value = year
    costs["B7"] = "Mature"
    labels = (
        "Salaries",
        "Rent and Utilities",
        "Operational Cost",
        "Study Material",
        "Acad Central Cost",
        "Corporate OH",
        "Other Cost",
    )
    for offset, label in enumerate(labels):
        assumption_row = 8 + offset
        detail_row = 16 + offset
        costs.cell(assumption_row, 2).value = label
        costs.cell(assumption_row, 3).value = "as a % of Revenue"
        costs.cell(detail_row, 2).value = label
        for column in range(9, 14):
            costs.cell(assumption_row, column).value = 0.1
    costs["B23"] = "Total Foundation Course Mature Expenses"

    fixed = workbook.create_sheet("Fixed Assets Schedule")
    for column in range(3, 13):
        fixed.cell(3, column).value = (
            f"='Consolidated P&L'!{get_column_letter(column)}4"
        )
    fixed["B5"] = "Opening Balance"
    fixed["B6"] = "Capex"
    fixed["B7"] = "Depreciation"
    fixed["B8"] = "Closing Balance"
    fixed["B10"] = "Annual Depreciation (%)"
    fixed["C10"] = 0.25
    for column in range(8, 13):
        letter = get_column_letter(column)
        fixed.cell(5, column).value = 100
        fixed.cell(6, column).value = 20
        fixed.cell(8, column).value = f"=SUM({letter}5:{letter}7)"

    balance = workbook.create_sheet("Consolidated BS")
    balance["B22"] = "Total Assets"
    balance["B48"] = "Total Equity & Liabilities"
    balance["B50"] = "Check"
    for column, year in enumerate(range(2021, 2031), start=3):
        balance.cell(4, column).value = year
        balance.cell(22, column).value = 100
        balance.cell(48, column).value = 100

    workbook.save(source)
    workbook.save(output)
    workbook.close()

    changes = complete_financial_model_runtime_actions(
        output,
        source_path=source,
        instruction=(
            "In the Consolidated P&L sheet, calculate the CAGR of Total Revenue "
            "from 2025 to 2030. In the Revenue Drivers sheet, reference actual "
            "Other Revenue from another tab, and then hardcode the % of Operating "
            "Revenue for the forecasted years as the same value as the last actual "
            "year's value. From that, calculate forecasted Other Revenue. Compute "
            "operational costs for each Foundation Course line item in the cost "
            "drivers. In the Fixed Assets Schedule sheet, compute depreciation by "
            "applying the 25% assumption to opening balance plus capital expenditure "
            "for the forecast period. In the Consolidated BS sheet, add a validation "
            "check for each year."
        ),
    )

    repaired = load_workbook(output, data_only=False)
    assert repaired["Consolidated P&L"]["M9"].value == "=(L9/G9)^(1/5)-1"
    assert repaired["Revenue Drivers"]["D137"].value == "='Consolidated P&L'!C6"
    assert repaired["Revenue Drivers"]["I138"].value == 0.02
    assert repaired["Revenue Drivers"]["M137"].value == "=M138*M135"
    assert repaired["Cost Drivers"]["I16"].value == (
        "=I8*'Revenue Drivers'!I$17"
    )
    assert repaired["Cost Drivers"]["M23"].value == "=SUM(M16:M22)"
    assert repaired["Fixed Assets Schedule"]["H7"].value == "=-$C$10*(H5+H6)"
    assert repaired["Consolidated BS"]["L50"].value == "=L22-L48"
    assert len(changes) >= 76
    repaired.close()


def test_education_model_dcf_ratio_and_mature_revenue_rules(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()

    dcf = workbook.active
    dcf.title = "DCF Valuation"
    dcf["D5"] = 2026
    for column in range(5, 9):
        previous = get_column_letter(column - 1)
        dcf.cell(5, column).value = f"=EOMONTH({previous}5,12)"
    dcf["I5"] = "Terminal Value"
    dcf["B7"] = "EBIT"
    dcf["H7"] = 100
    dcf["B18"] = "Period Factor"
    dcf["B19"] = "Discounting Factor"
    dcf["B24"] = "Enterprise Value"
    dcf["C24"] = 500
    dcf["B29"] = "WACC"
    dcf["C29"] = 0.1
    dcf["B30"] = "TGR"
    dcf["C30"] = 0.04
    dcf["B33"] = "EV/Revenue"
    for column, year in enumerate((25, 26, 27), start=3):
        dcf.cell(32, column).value = f"FY{year}"
    for column in range(4, 9):
        dcf.cell(18, column).value = column - 3

    pnl = workbook.create_sheet("Consolidated P&L")
    pnl["B9"] = "Total Revenue"
    pnl["B51"] = "EBITDA"
    pnl["B52"] = "EBITDA Margin"
    for column, year in enumerate(range(2021, 2031), start=3):
        pnl.cell(4, column).value = year
        pnl.cell(9, column).value = 100
        pnl.cell(51, column).value = 20
        pnl.cell(52, column).value = 0.2

    ratios = workbook.create_sheet("Ratio Analysis")
    ratios["B17"] = "EBITDA Margin"
    ratios["B37"] = "Receivable Days"
    ratios["B38"] = "Payable Days"
    ratios["B39"] = "Inventory Days"
    ratios["B40"] = "Cash Conversion Cycle"
    for column, year in enumerate(range(2021, 2031), start=3):
        ratios.cell(4, column).value = year
        ratios.cell(37, column).value = 30
        ratios.cell(38, column).value = 20
        ratios.cell(39, column).value = 10

    revenue = workbook.create_sheet("Revenue Drivers")
    revenue["B7"] = "Mature"
    revenue["B8"] = "Headcount"
    revenue["B13"] = "Average Ticket Size"
    revenue["B17"] = "Total Revenue from Foundation Course Mature"
    for column, year in enumerate(range(2021, 2031), start=4):
        revenue.cell(4, column).value = year
    for column in range(9, 14):
        revenue.cell(8, column).value = 10
        revenue.cell(13, column).value = 5

    workbook.save(source)
    workbook.save(output)
    workbook.close()

    changes = complete_financial_model_runtime_actions(
        output,
        source_path=source,
        instruction=(
            "In the DCF Valuation sheet, calculate Terminal Year EBIT by growing "
            "2030E EBIT at the terminal growth rate, then calculate discounting "
            "factors for 2026E-2030E using WACC, then calculate EV/Revenue for "
            "FY25-FY27 displaying 'NM' for zero or negative revenue. In the Ratio "
            "Analysis sheet, calculate EBITDA Margin for 2021A-2030E, then calculate "
            "Cash Conversion Cycle for 2021A-2030E, displaying 'NA' where inputs are "
            "unavailable. In the Revenue Drivers sheet, calculate Total Revenue from "
            "Foundation Course (Mature) for 2026E-2030E."
        ),
    )

    repaired = load_workbook(output, data_only=False)
    assert repaired["DCF Valuation"]["I7"].value == "=H7*(1+C30)"
    assert repaired["DCF Valuation"]["D19"].value == "=1/(1+$C$29)^D18"
    assert repaired["DCF Valuation"]["C33"].value == (
        '=IF($C$24/\'Consolidated P&L\'!G9>0,'
        '$C$24/\'Consolidated P&L\'!G9,"NM")'
    )
    assert repaired["Ratio Analysis"]["C17"].value == "='Consolidated P&L'!C52"
    assert repaired["Ratio Analysis"]["L40"].value == (
        '=IFERROR(L37+L39-L38,"NA")'
    )
    assert repaired["Revenue Drivers"]["I17"].value == "=I13*I8"
    assert repaired["Revenue Drivers"]["M17"].value == "=M13*M8"
    assert len(changes) == 34
    repaired.close()
