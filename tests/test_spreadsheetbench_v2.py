from __future__ import annotations

import io
import json
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from spreadsheet_harness.config import ProviderConfig
from spreadsheet_harness.errors import AgentExecutionFailure, HarnessError
from spreadsheet_harness.skills import SkillRegistry
from spreadsheet_harness.spreadsheetbench_v2 import (
    DEFAULT_V2_EVALUATOR,
    SpreadsheetBenchV2Task,
    _balanced_arm_orders,
    _load_official_evaluator,
    _official_score,
    _parse_answer_position_segment,
    _restore_unchanged_input_formula_caches,
    _seal_interrupted_v2_row,
    _summarize_results,
    audit_spreadsheetbench_v2_comparison,
    load_spreadsheetbench_v2_tasks,
    run_spreadsheetbench_v2_comparison,
    select_spreadsheetbench_v2_tasks,
)


def _rewrite_zip_member(path: Path, member_name: str, callback) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as source:
        with zipfile.ZipFile(buffer, "w") as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename == member_name:
                    payload = callback(payload)
                target.writestr(info, payload)
    path.write_bytes(buffer.getvalue())


def _workbook(path: Path, value: int) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Model"
    sheet["A1"] = value
    sheet["A2"] = "unchanged"
    workbook.save(path)
    workbook.close()


def _set_formula_caches(path: Path, values: dict[str, str]) -> None:
    def rewrite(payload: bytes) -> bytes:
        root = ET.fromstring(payload)
        namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        for cell in root.findall(".//x:c", namespace):
            coordinate = cell.attrib.get("r")
            if coordinate not in values or cell.find("x:f", namespace) is None:
                continue
            cached = cell.find("x:v", namespace)
            if cached is None:
                cached = ET.SubElement(cell, f"{{{namespace['x']}}}v")
            cached.text = values[coordinate]
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    _rewrite_zip_member(path, "xl/worksheets/sheet1.xml", rewrite)


def _dataset(tmp_path: Path, category: str = "Template") -> Path:
    root = tmp_path / "spreadsheetbench-v2"
    category_root = root / category
    spreadsheet = category_root / "spreadsheet"
    spreadsheet.mkdir(parents=True)
    _workbook(spreadsheet / "input.xlsx", 1)
    _workbook(spreadsheet / "golden.xlsx", 2)
    (category_root / "dataset.json").write_text(
        json.dumps(
            [
                {
                    "id": "01_01",
                    "instruction": "Update the model.",
                    "spreadsheet_path": "spreadsheet/input.xlsx",
                    "golden_response_path": "spreadsheet/golden.xlsx",
                    "answer_position": "'Model'!A1:A2",
                }
            ]
        ),
        encoding="utf-8",
    )
    return root


def test_load_and_select_spreadsheetbench_v2_tasks(tmp_path: Path) -> None:
    root = _dataset(tmp_path)
    tasks = load_spreadsheetbench_v2_tasks(root, categories=("Template",))

    assert [task.task_id for task in tasks] == ["Template/01_01"]
    assert select_spreadsheetbench_v2_tasks(tasks, ("01_01",)) == tasks
    assert select_spreadsheetbench_v2_tasks(tasks, ("Template/01_01",)) == tasks


def test_answer_position_parser_preserves_quoted_sheet_whitespace() -> None:
    assert _parse_answer_position_segment(
        "'Ex 10 - Forecast Assumptions '!B2:H50",
        default_sheet="Model",
    ) == ("Ex 10 - Forecast Assumptions ", "B2:H50")


def test_unchanged_formula_cache_restore_uses_input_without_golden(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    output = tmp_path / "output.xlsx"
    for path, second_formula in ((source, "=2+2"), (output, "=3+3")):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Model"
        sheet["A1"] = "=1+1"
        sheet["A2"] = second_formula
        if path == output:
            sheet["A1"].font = Font(bold=True)
        workbook.save(path)
        workbook.close()
    _set_formula_caches(source, {"A1": "2", "A2": "4"})
    _set_formula_caches(output, {"A1": "99", "A2": "6"})
    task = SpreadsheetBenchV2Task(
        category="Template",
        item_id="01_01",
        instruction="Update the model.",
        category_root=tmp_path,
        input_path=source,
        golden_path=tmp_path / "missing-golden.xlsx",
        answer_position="'Model'!A1:A2",
        source_row={},
    )

    restored = _restore_unchanged_input_formula_caches(task, output)

    assert restored == {"selected_cells": 1, "restored_cells": 1, "skipped": False}
    formulas = load_workbook(output, data_only=False)
    values = load_workbook(output, data_only=True)
    try:
        assert formulas["Model"]["A1"].value == "=1+1"
        assert formulas["Model"]["A1"].font.bold is True
        assert formulas["Model"]["A2"].value == "=3+3"
        assert values["Model"]["A1"].value == 2
        assert values["Model"]["A2"].value == 6
    finally:
        formulas.close()
        values.close()


def test_v2_runner_does_not_restore_formula_caches_after_calculate_all(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _dataset(tmp_path)
    task = load_spreadsheetbench_v2_tasks(root, categories=("Template",))[0]

    class Evidence:
        def to_dict(self) -> dict[str, object]:
            return {"observed_terminal_tool": "completed"}

    monkeypatch.setattr(
        "spreadsheet_harness.spreadsheetbench_v2.run_arm",
        lambda **_kwargs: Evidence(),
    )
    monkeypatch.setattr(
        "spreadsheet_harness.spreadsheetbench_v2.recalculate_workbook",
        lambda *_args, **_kwargs: {"calculation_mode": "uno-calculate-all"},
    )
    monkeypatch.setattr(
        "spreadsheet_harness.spreadsheetbench_v2._restore_unchanged_input_formula_caches",
        lambda *_args, **_kwargs: pytest.fail(
            "dependency-consistent calculateAll caches must not be overwritten"
        ),
    )
    monkeypatch.setattr(
        "spreadsheet_harness.spreadsheetbench_v2._official_score",
        lambda *_args, **_kwargs: {
            "accuracy": 1.0,
            "modification_accuracy": 1.0,
            "regression_accuracy": 1.0,
        },
    )

    run_spreadsheetbench_v2_comparison(
        config=ProviderConfig("https://example.test/v1", "secret", "test-model"),
        dataset_root=root,
        evaluator_path=DEFAULT_V2_EVALUATOR,
        output_dir=tmp_path / "results-no-cache-restore",
        skill_registry=SkillRegistry([]),
        tasks=(task,),
        arms=("ours",),
        max_model_calls=1,
        max_turns_per_arm=1,
    )
    rows = json.loads(
        (tmp_path / "results-no-cache-restore" / "results.json").read_text()
    )

    assert rows[0]["recalculation"]["unchanged_input_formula_cache_restore"] == {
        "selected_cells": 0,
        "restored_cells": 0,
        "skipped": True,
        "reason": "preserve_dependency_consistent_calculate_all_values",
    }


def test_spreadsheetbench_v2_short_id_must_be_unambiguous(tmp_path: Path) -> None:
    template = load_spreadsheetbench_v2_tasks(
        _dataset(tmp_path / "template", "Template"), categories=("Template",)
    )[0]
    debugging = load_spreadsheetbench_v2_tasks(
        _dataset(tmp_path / "debugging", "Debugging"), categories=("Debugging",)
    )[0]

    with pytest.raises(HarnessError, match="ambiguous"):
        select_spreadsheetbench_v2_tasks((template, debugging), ("01_01",))


def test_pinned_official_v2_evaluator_scores_selected_output(tmp_path: Path) -> None:
    root = _dataset(tmp_path)
    task = load_spreadsheetbench_v2_tasks(root, categories=("Template",))[0]
    evaluator, _, digest = _load_official_evaluator(DEFAULT_V2_EVALUATOR)

    score = _official_score(
        evaluator,
        task,
        task.golden_path,
        tmp_path / "official-output",
        model_calls=3,
    )

    assert len(digest) == 64
    assert score["regression_accuracy"] == 1.0
    assert score["modification_accuracy"] == 1.0
    assert score["accuracy"] == 1.0
    assert score["interaction_turns"] == 3


def test_pinned_official_v2_evaluator_repairs_unbound_prefix_metadata(tmp_path: Path) -> None:
    root = _dataset(tmp_path)
    task = load_spreadsheetbench_v2_tasks(root, categories=("Template",))[0]

    def _break_core(payload: bytes) -> bytes:
        return payload.decode("utf-8").replace(
            ' xmlns:dc="http://purl.org/dc/elements/1.1/"', "", 1
        ).encode("utf-8")

    _rewrite_zip_member(task.input_path, "docProps/core.xml", _break_core)
    evaluator, _, _ = _load_official_evaluator(DEFAULT_V2_EVALUATOR)

    score = _official_score(
        evaluator,
        task,
        task.golden_path,
        tmp_path / "official-output-broken-input",
        model_calls=1,
    )

    assert score["accuracy"] == 1.0
    assert score["regression_accuracy"] == 1.0


def test_v2_arm_order_is_deterministic_and_balanced(tmp_path: Path) -> None:
    tasks = []
    for index in range(6):
        task = load_spreadsheetbench_v2_tasks(
            _dataset(tmp_path / str(index)), categories=("Template",)
        )[0]
        tasks.append(
            type(task)(
                category=task.category,
                item_id=f"01_{index + 1:02d}",
                instruction=task.instruction,
                category_root=task.category_root,
                input_path=task.input_path,
                golden_path=task.golden_path,
                answer_position=task.answer_position,
                source_row=task.source_row,
            )
        )

    orders = _balanced_arm_orders(tasks, 20260820)

    assert orders == _balanced_arm_orders(tasks, 20260820)
    assert sum(order[0] == "bare" for order in orders.values()) == 3
    assert sum(order[0] == "ours" for order in orders.values()) == 3


def test_v2_incomplete_results_do_not_turn_not_scored_into_zero() -> None:
    rows = [
        {
            "arm": "bare",
            "status": "completed",
            "passed": False,
            "outcome_kind": "scored",
            "official_score": {
                "accuracy": 0.0,
                "modification_accuracy": 0.5,
                "regression_accuracy": 1.0,
            },
            "budget": {"used": {"total_tokens": 10, "model_calls": 1}},
        },
        {
            "arm": "ours",
            "status": "error",
            "passed": False,
            "budget": {"used": {"total_tokens": 5, "model_calls": 1}},
        },
    ]

    summary, complete = _summarize_results(rows, task_count=1, arms=("bare", "ours"))

    assert complete is False
    assert summary["arms"]["bare"]["accuracy"] == 0.0
    assert summary["arms"]["ours"]["accuracy"] is None
    assert summary["arms"]["ours"]["scored_accuracy"] is None
    assert summary["paired_delta"]["accuracy"] is None


def test_v2_summary_counts_scored_model_execution_failures() -> None:
    rows = []
    for arm in ("bare", "ours"):
        rows.append(
            {
                "arm": arm,
                "status": "completed",
                "passed": False,
                "outcome_kind": ("model_execution_failure" if arm == "ours" else "scored"),
                "official_score": {
                    "accuracy": 0.0,
                    "modification_accuracy": 0.5,
                    "regression_accuracy": 1.0,
                },
                "budget": {"used": {"total_tokens": 10, "model_calls": 1}},
            }
        )

    summary, complete = _summarize_results(rows, task_count=1, arms=("bare", "ours"))

    assert complete is True
    assert summary["arms"]["ours"]["model_execution_failures"] == 1
    assert summary["arms"]["ours"]["accuracy"] == 0.0
    assert summary["paired_delta"]["accuracy"] == 0.0


def test_v2_audit_authentic_not_scored_row_is_valid_but_incomplete(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _dataset(tmp_path)
    task = load_spreadsheetbench_v2_tasks(root, categories=("Template",))[0]
    output = tmp_path / "results"
    output.mkdir()
    import spreadsheet_harness.spreadsheetbench_v2 as v2

    _, evaluator_path, evaluator_sha256 = _load_official_evaluator(DEFAULT_V2_EVALUATOR)
    manifest = {
        "schema_version": v2.SPREADSHEETBENCH_V2_MANIFEST_SCHEMA,
        "protocol": v2.SPREADSHEETBENCH_V2_PROTOCOL,
        "official_evaluator": {"sha256": evaluator_sha256},
        "implementation": v2._implementation_record(SkillRegistry([]).freeze()),
        "arms": ["ours"],
        "tasks": [
            {
                "task_id": task.task_id,
                "category": task.category,
                "item_id": task.item_id,
                "instruction_sha256": v2.hashlib.sha256(task.instruction.encode()).hexdigest(),
                "input_sha256": v2._sha256(task.input_path),
                "golden_sha256": v2._sha256(task.golden_path),
                "answer_position_sha256": v2.hashlib.sha256(
                    task.answer_position.encode()
                ).hexdigest(),
            }
        ],
    }
    manifest["manifest_sha256"] = v2._manifest_sha256(manifest)
    (output / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    (output / "results.json").write_text(
        json.dumps(
            [
                {
                    "task_id": task.task_id,
                    "arm": "ours",
                    "status": "error",
                    "outcome_kind": "not_scored",
                    "manifest_sha256": manifest["manifest_sha256"],
                }
            ]
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(v2, "_atomic_write_json", lambda *_args: None)
    report = audit_spreadsheetbench_v2_comparison(
        output,
        dataset_root=root,
        evaluator_path=evaluator_path,
    )
    assert report["audit_valid"] is True
    assert report["study_complete"] is False
    assert report["not_scored_arm_tasks"] == 1


def test_v2_runner_officially_scores_known_model_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _dataset(tmp_path)
    task = load_spreadsheetbench_v2_tasks(root, categories=("Template",))[0]

    class Evidence:
        def to_dict(self) -> dict[str, object]:
            return {"observed_terminal_tool": "budget_exhausted"}

    def fail_with_partial_evidence(**_kwargs: object) -> None:
        raise AgentExecutionFailure(
            "token budget exhausted",
            reason="budget_exhausted",
            agent_result=Evidence(),
        )

    monkeypatch.setattr(
        "spreadsheet_harness.spreadsheetbench_v2.run_arm", fail_with_partial_evidence
    )
    monkeypatch.setattr(
        "spreadsheet_harness.spreadsheetbench_v2.recalculate_workbook",
        lambda *_args, **_kwargs: {"ok": True},
    )
    monkeypatch.setattr(
        "spreadsheet_harness.spreadsheetbench_v2._official_score",
        lambda *_args, **_kwargs: {
            "accuracy": 1.0,
            "modification_accuracy": 1.0,
            "regression_accuracy": 1.0,
        },
    )

    summary = run_spreadsheetbench_v2_comparison(
        config=ProviderConfig("https://example.test/v1", "secret", "test-model"),
        dataset_root=root,
        evaluator_path=DEFAULT_V2_EVALUATOR,
        output_dir=tmp_path / "results",
        skill_registry=SkillRegistry([]),
        tasks=(task,),
        max_model_calls=1,
        max_turns_per_arm=1,
    )
    rows = json.loads((tmp_path / "results" / "results.json").read_text())

    assert summary["study_complete"] is True
    assert summary["arms"]["bare"]["accuracy"] == 1.0
    assert summary["arms"]["ours"]["accuracy"] == 1.0
    assert all(row["outcome_kind"] == "model_execution_failure" for row in rows)
    assert all(row["model_failure_reason"] == "budget_exhausted" for row in rows)


def test_v2_visual_generation_stages_official_windows_filename(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _dataset(tmp_path, category="Visualization")
    task = load_spreadsheetbench_v2_tasks(root, categories=("Visualization",))[0]

    class Evidence:
        def to_dict(self) -> dict[str, object]:
            return {"observed_terminal_tool": "completed"}

    monkeypatch.setattr(
        "spreadsheet_harness.spreadsheetbench_v2.run_arm",
        lambda **_kwargs: Evidence(),
    )
    output = tmp_path / "visual-results"
    summary = run_spreadsheetbench_v2_comparison(
        config=ProviderConfig("https://example.test/v1", "secret", "test-model"),
        dataset_root=root,
        evaluator_path=DEFAULT_V2_EVALUATOR,
        output_dir=output,
        skill_registry=SkillRegistry([]),
        tasks=(task,),
        arms=("ours",),
        max_model_calls=1,
        max_turns_per_arm=1,
        visual_generation_only=True,
    )
    rows = json.loads((output / "results.json").read_text())
    staged = output / "visual_outputs" / "ours" / "1_01_01_output.xlsx"

    assert summary["generation_complete"] is True
    assert summary["study_complete"] is False
    assert staged.is_file()
    assert rows[0]["status"] == "generated"
    assert rows[0]["outcome_kind"] == "pending_official_visual_evaluation"
    assert rows[0]["visual_output_sha256"]


def test_v2_color_only_task_skips_libreoffice_recalculation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _dataset(tmp_path, category="Debugging")
    category_root = root / "Debugging"
    original = category_root / "spreadsheet" / "input.xlsx"
    # Give the detector actual color-neighborhood evidence.  The task filename
    # is intentionally generic: semantic routing must not depend on a hidden
    # fixture label encoded in the basename.
    workbook = load_workbook(original)
    sheet = workbook.active
    for row in range(1, 16):
        sheet.cell(row, 2).value = f"=A{row}"
        sheet.cell(row, 2).font = Font(color="7030A0" if row in {2, 5, 8, 11, 14} else "70AD47")
    workbook.save(original)
    workbook.close()
    manifest_path = category_root / "dataset.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    task = load_spreadsheetbench_v2_tasks(root, categories=("Debugging",))[0]

    class Evidence:
        def to_dict(self) -> dict[str, object]:
            return {"observed_terminal_tool": "completed"}

    monkeypatch.setattr(
        "spreadsheet_harness.spreadsheetbench_v2.run_arm",
        lambda **_kwargs: Evidence(),
    )
    monkeypatch.setattr(
        "spreadsheet_harness.spreadsheetbench_v2.recalculate_workbook",
        lambda *_args, **_kwargs: pytest.fail("color-only task must not be recalculated"),
    )
    monkeypatch.setattr(
        "spreadsheet_harness.spreadsheetbench_v2._official_score",
        lambda *_args, **_kwargs: {
            "accuracy": 0.0,
            "modification_accuracy": 0.0,
            "regression_accuracy": 1.0,
        },
    )

    run_spreadsheetbench_v2_comparison(
        config=ProviderConfig("https://example.test/v1", "secret", "test-model"),
        dataset_root=root,
        evaluator_path=DEFAULT_V2_EVALUATOR,
        output_dir=tmp_path / "color-results",
        skill_registry=SkillRegistry([]),
        tasks=(task,),
        arms=("bare",),
        max_model_calls=1,
        max_turns_per_arm=1,
    )
    rows = json.loads((tmp_path / "color-results" / "results.json").read_text())

    assert rows[0]["recalculation"] == {
        "ok": True,
        "skipped": True,
        "reason": "preserve_ooxml_formula_caches_and_font_theme_for_color_only_task",
    }


def test_v2_seals_ambiguous_interrupted_request_without_replay(tmp_path: Path) -> None:
    root = _dataset(tmp_path)
    task = load_spreadsheetbench_v2_tasks(root, categories=("Template",))[0]
    output = tmp_path / "results"
    run_dir = output / "runs" / "Template" / "01_01" / "bare"
    (run_dir / "artifacts").mkdir(parents=True)
    _workbook(run_dir / "artifacts" / "output.xlsx", 3)
    events = [
        {
            "timestamp": "2026-08-20T00:00:00+00:00",
            "event": "model.requested",
            "payload": {},
        },
        {
            "timestamp": "2026-08-20T00:00:01+00:00",
            "event": "model.responded",
            "payload": {"usage": {"total_tokens": 12}},
        },
        {
            "timestamp": "2026-08-20T00:00:02+00:00",
            "event": "model.requested",
            "payload": {},
        },
    ]
    (run_dir / "trajectory.jsonl").write_text(
        "\n".join(json.dumps(event) for event in events) + "\n", encoding="utf-8"
    )

    row = _seal_interrupted_v2_row(
        output=output,
        task=task,
        arm="bare",
        manifest_sha256="a" * 64,
        model="test-model",
        max_model_calls=20,
        max_total_tokens=200_000,
        task_timeout_seconds=1_800,
    )

    assert row["status"] == "error"
    assert row["outcome_kind"] == "not_scored"
    assert row["error_type"] == "InterruptedAmbiguousRequest"
    assert row["ambiguous_inflight_requests"] == 1
    assert row["known_completed_model_calls"] == 1
    assert row["budget"]["used"]["model_calls"] == 2
    assert row["budget"]["used"]["total_tokens"] == 12
    assert len(row["output_sha256"]) == 64
