from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from spreadsheet_harness.template_repairs import complete_template_schedules


def quarterly_fixture(path: Path, row_shift: int, col_shift: int, year: int, days: int) -> None:
    book = Workbook()
    ws = book.active
    ws.title = "Unseen Layout"
    labels = {
        9: "Revenue", 10: "qoq growth", 11: "Cost of goods sold (incl. Depr)",
        12: "% of revenue", 13: "Gross Margin", 15: "Cash & Equivalents",
        17: "Accounts Receivable", 18: "DSO (Days)", 20: "Inventory",
        21: "Inventory Days", 23: "Other Current Assets", 25: "Current Liabilities",
        30: "Changes in Working Capital",
    }
    for row, label in labels.items():
        ws.cell(row + row_shift, 2 + col_shift, label)
    ws.cell(6 + row_shift, 3 + col_shift, f"FY {year-1}")
    for q in range(1, 5):
        ws.cell(6 + row_shift, 3 + q + col_shift, f"Q{q} {year}E")
    ws.cell(6 + row_shift, 8 + col_shift, f"FY {year}E")
    ws.cell(32 + row_shift, 2 + col_shift, f"All calculations use {days} days per quarter")
    for row in (15, 17, 20, 23, 25):
        ws.cell(row + row_shift, 3 + col_shift, row * 31)
    for col in range(4 + col_shift, 8 + col_shift):
        for row, value in ((10, .07), (12, .43), (18, 53), (21, 144), (23, 670), (25, 1600)):
            ws.cell(row + row_shift, col, value)
    ws.cell(9 + row_shift, 4 + col_shift, 937.25)
    book.save(path)
    book.close()


@pytest.mark.parametrize("row_shift,col_shift,year,days", [(0, 0, 2034, 90), (19, 4, 2041, 91)])
def test_quarterly_completion_is_layout_year_and_day_count_independent(
    tmp_path: Path, row_shift: int, col_shift: int, year: int, days: int,
) -> None:
    path = tmp_path / "unseen.xlsx"
    quarterly_fixture(path, row_shift, col_shift, year, days)
    changes = complete_template_schedules(path)
    assert changes
    book = load_workbook(path)
    ws = book.active
    quarter, total = get_column_letter(6 + col_shift), get_column_letter(8 + col_shift)
    assert ws.cell(17 + row_shift, 6 + col_shift).value == f"=({quarter}{9+row_shift}/{days})*{quarter}{18+row_shift}"
    assert ws.cell(12 + row_shift, 8 + col_shift).value == f"={total}{11+row_shift}/{total}{9+row_shift}"
    assert ws.cell(9 + row_shift, 4 + col_shift).value == 937.25
    assert ws.cell(16 + row_shift, 6 + col_shift).value is None
    book.close()
    assert complete_template_schedules(path) == []


@pytest.mark.parametrize("defect", ["missing-day-count", "missing-input", "duplicate-label", "inconsistent-year"])
def test_quarterly_completion_abstains_on_ambiguous_contract(tmp_path: Path, defect: str) -> None:
    path = tmp_path / "ambiguous.xlsx"
    quarterly_fixture(path, 0, 0, 2034, 90)
    book = load_workbook(path)
    ws = book.active
    if defect == "missing-day-count":
        ws["B32"] = None
    elif defect == "missing-input":
        ws["G21"] = None
    elif defect == "duplicate-label":
        ws["B40"] = "Revenue"
    else:
        ws["F6"] = "Q3 2037E"
    book.save(path)
    book.close()
    before = path.read_bytes()
    assert complete_template_schedules(path) == []
    assert path.read_bytes() == before
