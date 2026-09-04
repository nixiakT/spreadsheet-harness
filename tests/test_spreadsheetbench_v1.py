from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from spreadsheet_harness.spreadsheetbench_v1 import (
    SpreadsheetBenchV1Case,
    SpreadsheetBenchV1Instruction,
    audit_spreadsheetbench_v1_comparison,
    official_compare_v1,
    replay_v1_calls,
    score_v1_instruction,
    successful_v1_replay_calls,
    summarize_v1_scores,
    v1_planner_replay_plan,
)


def _book(path: Path, value: object, *, title: str = "Sheet") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet["A1"] = value
    workbook.save(path)
    workbook.close()


def test_official_compare_v1_uses_first_sheet_for_unqualified_range(tmp_path: Path) -> None:
    golden = tmp_path / "golden.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    _book(golden, 1.234)
    _book(candidate, 1.23)
    assert official_compare_v1(golden, candidate, "A1")


def test_v1_instruction_keeps_provider_failure_not_scored(tmp_path: Path) -> None:
    golden = tmp_path / "golden.xlsx"
    source = tmp_path / "input.xlsx"
    output = tmp_path / "output.xlsx"
    _book(golden, 5)
    _book(source, 0)
    _book(output, 5)
    task = SpreadsheetBenchV1Instruction(
        task_id="x",
        instruction="edit",
        instruction_type="Cell-Level Manipulation",
        answer_position="A1",
        answer_sheet=None,
        cases=(
            SpreadsheetBenchV1Case(1, source, golden),
            SpreadsheetBenchV1Case(2, source, golden),
            SpreadsheetBenchV1Case(3, None, None),
        ),
    )
    row = score_v1_instruction(task, {1: output})
    assert row["status"] == "not_scored"
    assert row["soft"] is None
    assert row["hard"] is None
    assert [item["status"] for item in row["case_results"]] == [
        "scored",
        "not_scored",
        "dataset_missing",
    ]


def test_v1_missing_dataset_case_matches_official_three_case_denominator(
    tmp_path: Path,
) -> None:
    golden = tmp_path / "golden.xlsx"
    source = tmp_path / "input.xlsx"
    output = tmp_path / "output.xlsx"
    _book(golden, 5)
    _book(source, 0)
    _book(output, 5)
    task = SpreadsheetBenchV1Instruction(
        task_id="x",
        instruction="edit",
        instruction_type="Cell-Level Manipulation",
        answer_position="A1",
        answer_sheet=None,
        cases=(
            SpreadsheetBenchV1Case(1, source, golden),
            SpreadsheetBenchV1Case(2, source, golden),
            SpreadsheetBenchV1Case(3, None, None),
        ),
    )
    row = score_v1_instruction(task, {1: output, 2: output})
    assert row["status"] == "completed"
    assert row["soft"] == 2 / 3
    assert row["hard"] == 0


def test_v1_summary_refuses_partial_study() -> None:
    summary = summarize_v1_scores(
        [{"task_id": "x", "status": "completed", "soft": 1.0, "hard": 1}]
    )
    assert summary["study_complete"] is False
    assert summary["soft"] is None
    assert summary["hard"] is None


def test_v1_extracts_and_replays_only_successful_mutations(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _book(source, 0)
    trajectory = tmp_path / "trajectory.jsonl"
    events = [
        {
            "event": "tool.called",
            "payload": {
                "name": "inspect_range",
                "arguments": {"sheet": "Sheet", "range_ref": "A1"},
            },
        },
        {
            "event": "tool.returned",
            "payload": {"name": "inspect_range", "result": {"ok": True}},
        },
        {
            "event": "tool.called",
            "payload": {
                "name": "write_range",
                "arguments": {"sheet": "Sheet", "start_cell": "A1", "values": [[7]]},
            },
        },
        {
            "event": "tool.returned",
            "payload": {"name": "write_range", "result": {"ok": True}},
        },
    ]
    trajectory.write_text(
        "".join(f"{json.dumps(event)}\n" for event in events),
        encoding="utf-8",
    )
    calls = successful_v1_replay_calls(trajectory)
    assert calls == [
        {
            "name": "write_range",
            "arguments": {"sheet": "Sheet", "start_cell": "A1", "values": [[7]]},
        }
    ]
    result = replay_v1_calls(source, tmp_path / "replay", calls)
    assert result["status"] == "scored"
    workbook = load_workbook(result["output_workbook"], data_only=False)
    assert workbook["Sheet"]["A1"].value == 7
    workbook.close()


def test_v1_replays_applied_planner_before_tool_calls(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _book(source, 0)
    trajectory = tmp_path / "trajectory.jsonl"
    trajectory.write_text(
        json.dumps(
            {
                "event": "harness.planner_actions.applied",
                "payload": {"count": 1},
            }
        )
        + "\n",
        encoding="utf-8",
    )
    plan = """actions:
- action: write_value
  target: Sheet!A1
  value: 9
provenance:
- sheet: Sheet
  range: A1
"""
    frozen = v1_planner_replay_plan(
        {"stages": [{"name": "plan", "agent": {"final_text": plan}}]},
        trajectory,
    )
    result = replay_v1_calls(
        source,
        tmp_path / "planner-replay",
        [],
        instruction="Set A1 to 9",
        planner_plan=frozen,
    )
    assert result["planner_actions_replayed"] == 1
    workbook = load_workbook(result["output_workbook"], data_only=False)
    assert workbook["Sheet"]["A1"].value == 9
    workbook.close()


def test_v1_audit_keeps_authentic_not_scored_run_incomplete(
    tmp_path: Path, monkeypatch,
) -> None:
    dataset = tmp_path / "dataset"
    dataset.mkdir()
    output = tmp_path / "output"
    output.mkdir()
    source = dataset / "source.xlsx"
    golden = dataset / "golden.xlsx"
    _book(source, 0)
    _book(golden, 1)
    task = SpreadsheetBenchV1Instruction(
        task_id="x",
        instruction="edit",
        instruction_type="Cell-Level Manipulation",
        answer_position="A1",
        answer_sheet=None,
        cases=(SpreadsheetBenchV1Case(1, source, golden),),
    )
    import spreadsheet_harness.spreadsheetbench_v1 as v1

    monkeypatch.setattr(v1, "load_spreadsheetbench_v1", lambda _root: [task])
    manifest = {
        "schema_version": v1.V1_RUN_SCHEMA,
        "protocol": v1.V1_RUN_PROTOCOL,
        "arms": ["ours"],
        "tasks": [
            {
                "task_id": "x",
                "instruction_sha256": v1.hashlib.sha256(b"edit").hexdigest(),
                "answer_position_sha256": v1.hashlib.sha256(b"A1").hexdigest(),
                "cases": [
                    {
                        "case_index": 1,
                        "input_sha256": v1._sha256(source),
                        "golden_sha256": v1._sha256(golden),
                    }
                ],
            }
        ],
    }
    manifest["manifest_sha256"] = v1._manifest_sha256(manifest)
    (output / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    (output / "results.json").write_text(
        json.dumps(
            [
                {
                    "task_id": "x",
                    "arm": "ours",
                    "status": "not_scored",
                    "manifest_sha256": manifest["manifest_sha256"],
                }
            ]
        ),
        encoding="utf-8",
    )
    report = audit_spreadsheetbench_v1_comparison(output, dataset_root=dataset)
    assert report["audit_valid"] is True
    assert report["study_complete"] is False
    assert report["completed_arm_instructions"] == 0
    assert report["not_scored_arm_instructions"] == 1
