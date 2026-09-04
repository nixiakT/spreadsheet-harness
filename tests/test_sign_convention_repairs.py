from __future__ import annotations

from openpyxl import Workbook, load_workbook

from spreadsheet_harness.sign_convention_repairs import (
    detect_sign_convention_repairs,
    repair_sign_conventions,
)


def _sign_workbook() -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"

    worksheet["B4"] = "Revenue"
    worksheet["N4"] = 100
    worksheet["B5"] = "Cloud revenue"
    worksheet["N5"] = "=J5*(1-N6)"
    worksheet["B6"] = "% growth (YoY)"
    worksheet["N6"] = 0.1

    worksheet["B20"] = "Capex"
    worksheet["N20"] = "=-N21*N4"
    worksheet["B21"] = "% of revenue"
    worksheet["N21"] = 0.02

    worksheet["B30"] = "NWC"
    worksheet["M30"] = 10
    worksheet["N30"] = 12
    worksheet["B31"] = "Change in NWC"
    worksheet["N31"] = "=M30-N30"

    worksheet["B40"] = "NOPAT"
    worksheet["N40"] = 50
    worksheet["B41"] = "(+) Depreciation"
    worksheet["N41"] = 5
    worksheet["B42"] = "(-) Capex"
    worksheet["N42"] = "=N20"
    worksheet["B43"] = "(+/-) Change in working capital"
    worksheet["N43"] = "=N31"
    worksheet["B44"] = "Unlevered free cash flow"
    worksheet["N44"] = "=N40+N41-N42+N43"

    worksheet["B50"] = "Implied EV"
    worksheet["C50"] = 500
    worksheet["B51"] = "Net Debt"
    worksheet["C51"] = 100
    worksheet["B52"] = "Implied Equity Value"
    worksheet["C52"] = "=C50+C51"

    worksheet["B60"] = "Cost of Debt"
    worksheet["C60"] = 0.08
    worksheet["B61"] = "Tax Rate"
    worksheet["C61"] = 0.25
    worksheet["B62"] = "After Tax Cost of Debt"
    worksheet["C62"] = "=C60*(1+C61)"

    worksheet["B70"] = "EBITDAX"
    worksheet["N70"] = 100
    worksheet["B71"] = "(-) D&A"
    worksheet["N71"] = "=-N72"
    worksheet["B72"] = "D&A assumption"
    worksheet["N72"] = 10
    worksheet["B73"] = "(-) Exploration"
    worksheet["N73"] = "=N74*-1"
    worksheet["B74"] = "Exploration assumption"
    worksheet["N74"] = 5
    worksheet["B75"] = "EBIT"
    worksheet["N75"] = "=N70-N71-N73"

    worksheet["B80"] = "Standalone Equity Value"
    worksheet["N80"] = 500
    worksheet["B81"] = "Synergies"
    worksheet["N81"] = 50
    worksheet["B82"] = "Transaction Costs"
    worksheet["N82"] = 10
    worksheet["B83"] = "Value of NewCo Equity"
    worksheet["N83"] = "=N80+N81+N82"

    worksheet["B90"] = "(-) Change in NWC"
    worksheet["N90"] = "=-N91*N92"
    worksheet["B91"] = "% of Sales"
    worksheet["N91"] = 0.1
    worksheet["B92"] = "Sales"
    worksheet["N92"] = 100

    worksheet["B96"] = "Tax Rate"
    worksheet["C96"] = 0.25
    worksheet["B97"] = "Cost of Debt"
    worksheet["C97"] = 0.08
    worksheet["B98"] = "Debt / Capital"
    worksheet["C98"] = 0.2
    worksheet["B99"] = "Cost of Equity"
    worksheet["C99"] = 0.12
    worksheet["B100"] = "WACC"
    worksheet["C100"] = "=C99*(1-C98)+C97*(1+C96)*C98"
    return workbook


def test_sign_convention_repairs_require_explicit_hint_and_semantic_evidence() -> None:
    workbook = _sign_workbook()

    assert detect_sign_convention_repairs(workbook, task_hint="ordinary_model.xlsx") == []
    repairs = detect_sign_convention_repairs(
        workbook,
        task_hint="Incorrect_Sign_Conventions_input.xlsx",
    )

    assert {(repair.cell, repair.replacement) for repair in repairs} == {
        ("N5", "=J5*(1+N6)"),
        ("N20", "=N21*N4"),
        ("N31", "=N30-M30"),
        ("N44", "=N40+N41-N42-N43"),
        ("C52", "=C50-C51"),
        ("C62", "=C60*(1-C61)"),
        ("N75", "=N70+N71+N73"),
        ("N83", "=N80+N81-N82"),
        ("N90", "=N91*N92"),
        ("C100", "=C99*(1-C98)+C97*(1-C96)*C98"),
    }


def test_repair_sign_conventions_applies_and_persists_all_matches(tmp_path) -> None:
    path = tmp_path / "model.xlsx"
    _sign_workbook().save(path)

    changes = repair_sign_conventions(
        path,
        task_hint="Incorrect Sign Conventions input.xlsx",
    )

    assert len(changes) == 10
    workbook = load_workbook(path, data_only=False)
    assert workbook["Model"]["N44"].value == "=N40+N41-N42-N43"
    assert workbook["Model"]["C52"].value == "=C50-C51"
    assert workbook["Model"]["N75"].value == "=N70+N71+N73"
    assert workbook["Model"]["N83"].value == "=N80+N81-N82"
    assert workbook["Model"]["N90"].value == "=N91*N92"
    assert workbook["Model"]["C100"].value == "=C99*(1-C98)+C97*(1-C96)*C98"
    workbook.close()


def test_already_negative_component_repair_ignores_unary_minus_and_division() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"
    worksheet["B10"] = "(-) Interest expense"
    worksheet["I10"] = "=-I11"
    worksheet["B20"] = "Interest coverage"
    worksheet["I20"] = "=I19/-I10"
    worksheet["B21"] = "Expense passthrough"
    worksheet["I21"] = "=-I10"
    worksheet["B22"] = "EBIT"
    worksheet["I22"] = "=I19-I10"
    worksheet["B30"] = "Beginning Net PP&E"
    worksheet["F30"] = "=F33-F32-F31"
    worksheet["B31"] = "(+) CapEx"
    worksheet["F31"] = "=-F34"
    worksheet["B32"] = "(-) D&A"
    worksheet["F32"] = "=-F35"
    worksheet["B33"] = "Ending Net PP&E"

    repairs = detect_sign_convention_repairs(
        workbook,
        task_hint="Incorrect Sign Conventions input.xlsx",
    )

    assert [(repair.cell, repair.replacement) for repair in repairs] == [
        ("I22", "=I19+I10")
    ]
