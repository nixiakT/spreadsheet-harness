#!/usr/bin/env python3
"""Build a deterministic, task-level representative subset of SpreadsheetBench v1."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import random
import re
import statistics
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_TASKS = 912
CASE_INDICES = (1, 2, 3)
OPERATION_PATTERNS = {
    "find": r"\b(?:find|search|locate|look for)\b",
    "extract": r"\b(?:extract|retrieve|pull out|pull)\b",
    "sum": r"\b(?:sum|summ?ation|subtotal|total(?:ing|led)?|aggregate)\b",
    "highlight": r"\b(?:highlight|conditional formatting)\b",
    "remove": r"\bremove\b",
    "modify": r"\b(?:modify|edit|change|adjust|update)\b",
    "count": r"\b(?:count|number of)\b",
    "delete": r"\bdelete\w*\b",
    "calculate": r"\b(?:calculate|calculation|compute|formula|derive)\w*\b",
    "display": r"\b(?:display|show|output|present)\w*\b",
    "filter": r"\bfilter\w*\b",
    "sort": r"\bsort\w*\b",
    "create": r"\b(?:create|make|build|add|insert)\w*\b",
    "format": r"\bformat\w*\b",
    "copy": r"\b(?:copy|duplicate)\w*\b",
    "link": r"\b(?:link|reference|referencing)\w*\b",
    "lookup": r"\b(?:vlookup|xlookup|lookup|match)\w*\b",
    "chart": r"\b(?:chart|graph)\w*\b",
    "pivot": r"\bpivot\w*\b",
    "transpose": r"\btranspose\w*\b",
    "macro": r"\b(?:macro|vba)\b",
}
COMPILED_OPERATIONS = {k: re.compile(v, re.IGNORECASE) for k, v in OPERATION_PATTERNS.items()}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def quantile_bins(values: list[float], bins: int = 3) -> tuple[list[str], list[float]]:
    """Assign equal-frequency bins, preserving ties deterministically."""
    ordered = sorted(values)
    cuts = [ordered[min(len(ordered) - 1, math.ceil(len(ordered) * i / bins) - 1)] for i in range(1, bins)]
    labels = []
    for value in values:
        index = 0
        while index < len(cuts) and value > cuts[index]:
            index += 1
        labels.append(("low", "medium", "high")[index])
    return labels, cuts


def labels_from_cuts(values: list[float], cuts: list[float]) -> list[str]:
    labels = []
    for value in values:
        index = 0
        while index < len(cuts) and value > cuts[index]:
            index += 1
        labels.append(("low", "medium", "high")[index])
    return labels


def summary(values: list[float]) -> dict[str, float | None]:
    if not values:
        return {"min": None, "p25": None, "median": None, "p75": None, "max": None, "mean": None}
    ordered = sorted(values)
    return {
        "min": min(values),
        "p25": statistics.quantiles(ordered, n=4, method="inclusive")[0] if len(values) > 1 else ordered[0],
        "median": statistics.median(values),
        "p75": statistics.quantiles(ordered, n=4, method="inclusive")[2] if len(values) > 1 else ordered[0],
        "max": max(values),
        "mean": statistics.fmean(values),
    }


def distribution(values: list[Any], *, order: list[str] | None = None) -> dict[str, dict[str, float | int]]:
    counts = Counter(str(value) for value in values)
    keys = order or sorted(counts)
    total = len(values)
    return {key: {"count": counts.get(key, 0), "proportion": counts.get(key, 0) / total if total else 0.0} for key in keys}


def task_complexity(dataset_root: Path, row: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    task_id = str(row["id"])
    task_dir = dataset_root / str(row["spreadsheet_path"])
    cases: list[dict[str, Any]] = []
    for case_index in CASE_INDICES:
        path = task_dir / f"{case_index}_{task_id}_input.xlsx"
        if not path.is_file():
            cases.append({"case": case_index, "input_present": False})
            continue
        record: dict[str, Any] = {
            "case": case_index,
            "input_present": True,
            "file_size_bytes": path.stat().st_size,
        }
        try:
            workbook = load_workbook(path, read_only=True, data_only=False)
            dimensions = [(ws.max_row or 0, ws.max_column or 0) for ws in workbook.worksheets]
            workbook.close()
            record.update(
                {
                    "sheet_count": len(dimensions),
                    "max_rows": max((item[0] for item in dimensions), default=0),
                    "max_columns": max((item[1] for item in dimensions), default=0),
                    "total_rows": sum(item[0] for item in dimensions),
                    "total_columns": sum(item[1] for item in dimensions),
                }
            )
        except Exception as exc:  # keep one malformed workbook from losing its task
            record.update({"metadata_error": type(exc).__name__, "sheet_count": 0, "max_rows": 0, "max_columns": 0, "total_rows": 0, "total_columns": 0})
        cases.append(record)

    available = [item for item in cases if item.get("input_present") and "metadata_error" not in item]
    def mean(key: str) -> float:
        return statistics.fmean(float(item[key]) for item in available) if available else 0.0
    def maximum(key: str) -> float:
        return float(max((item[key] for item in available), default=0))
    metrics = {
        "task_id": task_id,
        "instruction_type": "Cell-Level" if "Cell-Level" in str(row.get("instruction_type", "")) else "Sheet-Level",
        "instruction": str(row.get("instruction", "")),
        "instruction_word_count": len(re.findall(r"[A-Za-z0-9]+(?:['-][A-Za-z0-9]+)?", str(row.get("instruction", "")))),
        "instruction_char_count": len(str(row.get("instruction", ""))),
        "instruction_sentence_count": max(1, len(re.findall(r"[.!?]+", str(row.get("instruction", ""))))),
        "instruction_clause_count": len(re.findall(r"\b(?:and|also|additionally|then|while|before|after|so that|as well as)\b", str(row.get("instruction", "")), re.IGNORECASE)) + 1,
        "case_input_count": sum(bool(item.get("input_present")) for item in cases),
        "case_answer_count": sum((task_dir / f"{i}_{task_id}_answer.xlsx").is_file() for i in CASE_INDICES),
        "sheet_count_mean": mean("sheet_count"),
        "max_rows_mean": mean("max_rows"),
        "max_columns_mean": mean("max_columns"),
        "total_rows_mean": mean("total_rows"),
        "total_columns_mean": mean("total_columns"),
        "file_size_bytes_mean": mean("file_size_bytes"),
        "sheet_count_max": maximum("sheet_count"),
        "max_rows_max": maximum("max_rows"),
        "max_columns_max": maximum("max_columns"),
        "file_size_bytes_max": maximum("file_size_bytes"),
        "cases": cases,
    }
    text = metrics["instruction"]
    operations = [name for name, pattern in COMPILED_OPERATIONS.items() if pattern.search(text)]
    metrics["operations"] = operations or ["other"]
    return metrics, cases


def proportional_quotas(groups: dict[Any, list[dict[str, Any]]], target: int) -> dict[Any, int]:
    """Allocate a fixed sample size by largest remainder, respecting capacities."""
    if not groups:
        return {}
    total = sum(len(items) for items in groups.values())
    raw = {key: target * len(items) / total for key, items in groups.items()}
    quotas = {key: min(len(groups[key]), math.floor(value)) for key, value in raw.items()}
    left = target - sum(quotas.values())
    ranking = sorted(groups, key=lambda key: (raw[key] - math.floor(raw[key]), len(groups[key]), repr(key)), reverse=True)
    while left > 0:
        progressed = False
        for key in ranking:
            if quotas[key] < len(groups[key]):
                quotas[key] += 1
                left -= 1
                progressed = True
                if left == 0:
                    break
        if not progressed:
            raise RuntimeError("unable to satisfy proportional quota")
    return quotas


def allocate_strata(records: list[dict[str, Any]], sample_size: int, seed: int) -> list[dict[str, Any]]:
    strata: dict[tuple[str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        key = (record["instruction_type"], record["primary_operation"], record["workbook_complexity_bin"], record["instruction_complexity_bin"])
        strata[key].append(record)
    by_type: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        by_type[record["instruction_type"]].append(record)
    type_quotas = proportional_quotas(by_type, sample_size)
    strata_by_type: dict[str, dict[tuple[str, str, str, str], list[dict[str, Any]]]] = defaultdict(dict)
    for key, items in strata.items():
        strata_by_type[key[0]][key] = items
    quotas: dict[tuple[str, str, str, str], int] = {}
    for instruction_type, groups in strata_by_type.items():
        quotas.update(proportional_quotas(groups, type_quotas[instruction_type]))
    rng = random.Random(seed)
    selected: list[dict[str, Any]] = []
    for key in sorted(strata, key=repr):
        pool = list(strata[key])
        rng.shuffle(pool)
        selected.extend(pool[:quotas[key]])
    if len(selected) != sample_size:
        raise RuntimeError(f"selected {len(selected)} records, expected {sample_size}")
    rng.shuffle(selected)
    return selected


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset", type=Path, default=Path("benchmarks/data/spreadsheetbench_912_v0.1"))
    parser.add_argument("--output", type=Path, default=Path("benchmarks/results/spreadsheetbench-v1-representative-200-seed42"))
    parser.add_argument("--sample-size", type=int, default=200)
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()
    dataset = args.dataset.resolve()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    rows = json.loads((dataset / "dataset.json").read_text(encoding="utf-8"))
    if len(rows) != EXPECTED_TASKS:
        raise SystemExit(f"expected {EXPECTED_TASKS} tasks, found {len(rows)}")
    if not 1 <= args.sample_size <= len(rows):
        raise SystemExit("sample size must be between 1 and the full task count")

    records: list[dict[str, Any]] = []
    for index, row in enumerate(rows, 1):
        record, _ = task_complexity(dataset, row)
        record["dataset_index"] = index - 1
        records.append(record)
        if index % 100 == 0:
            print(f"scanned {index}/{len(rows)} tasks", flush=True)

    operation_frequency = Counter(op for record in records for op in record["operations"])
    top_operations = {name for name, _ in operation_frequency.most_common(12)}
    for record in records:
        record["primary_operation"] = max(record["operations"], key=lambda op: (operation_frequency[op], op))
        if record["primary_operation"] not in top_operations:
            record["primary_operation"] = "other"
        workbook_score = (
            math.log1p(record["sheet_count_mean"])
            + math.log1p(record["max_rows_mean"])
            + math.log1p(record["max_columns_mean"])
            + math.log1p(record["file_size_bytes_mean"]) / 4
        )
        instruction_score = (
            math.log1p(record["instruction_word_count"])
            + math.log1p(record["instruction_sentence_count"])
            + math.log1p(record["instruction_clause_count"])
            + math.log1p(len(record["operations"]))
        )
        record["workbook_complexity_score"] = workbook_score
        record["instruction_complexity_score"] = instruction_score
    workbook_bins, workbook_cuts = quantile_bins([r["workbook_complexity_score"] for r in records])
    instruction_bins, instruction_cuts = quantile_bins([r["instruction_complexity_score"] for r in records])
    for record, wb_bin, inst_bin in zip(records, workbook_bins, instruction_bins):
        record["workbook_complexity_bin"] = wb_bin
        record["instruction_complexity_bin"] = inst_bin

    sampling_population = [record for record in records if record["case_input_count"] == len(CASE_INDICES)]
    if len(sampling_population) < args.sample_size:
        raise SystemExit(f"only {len(sampling_population)} tasks have all {len(CASE_INDICES)} input cases")
    selected = allocate_strata(sampling_population, args.sample_size, args.seed)
    selected_ids = sorted(record["task_id"] for record in selected)
    (output / "task_ids.txt").write_text("".join(task_id + "\n" for task_id in selected_ids), encoding="utf-8")
    selected_rows = [row for row in rows if str(row["id"]) in set(selected_ids)]
    (output / "selected_dataset.json").write_text(json.dumps(selected_rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    sampled_records = [r for r in records if r["task_id"] in set(selected_ids)]
    case_manifest = {
        r["task_id"]: {
            str(case["case"]): {
                "input": bool(case.get("input_present")),
                "answer": (dataset / str(next(row["spreadsheet_path"] for row in rows if str(row["id"]) == r["task_id"])) / f"{case['case']}_{r['task_id']}_answer.xlsx").is_file(),
            }
            for case in r["cases"]
        }
        for r in sampled_records
    }
    (output / "case_manifest.json").write_text(json.dumps(case_manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def dist_for(items: list[dict[str, Any]], key: str, order: list[str] | None = None) -> dict[str, dict[str, float | int]]:
        return distribution([item[key] for item in items], order=order)

    operation_order = list(OPERATION_PATTERNS) + ["other"]
    full_case_manifest = {r["task_id"]: [c["case"] for c in r["cases"] if c.get("input_present")] for r in records}
    report = {
        "schema_version": "spreadsheetbench-v1-representative-sampling-v1",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "seed": args.seed,
        "sample_size": args.sample_size,
        "sampling_unit": "instruction/task; all three input cases must exist and are retained per selected task",
        "dataset": {
            "root": str(dataset),
            "dataset_json_sha256": sha256(dataset / "dataset.json"),
            "full_task_count": len(records),
            "eligible_task_count_all_input_cases_present": len(sampling_population),
            "excluded_incomplete_input_case_task_count": len(records) - len(sampling_population),
            "expected_case_indices": list(CASE_INDICES),
        },
        "method": {
            "description": "Hierarchical proportional stratified sampling: first allocate by instruction_type, then allocate within each type across joint strata using largest-remainder quotas and deterministic per-stratum shuffles.",
            "joint_strata": ["instruction_type", "primary_operation", "workbook_complexity_bin", "instruction_complexity_bin"],
            "workbook_complexity": "Mean across all available input cases of sheet count, maximum rows, maximum columns, total dimensions, and file size; score is log1p(sheet)+log1p(rows)+log1p(columns)+log1p(bytes)/4, then equal-frequency low/medium/high bins.",
            "instruction_complexity": "log1p(word count)+log1p(sentence count)+log1p(clause count)+log1p(number of detected operation labels), then equal-frequency low/medium/high bins.",
            "operation_detection": OPERATION_PATTERNS,
            "primary_operation": "Highest-frequency detected operation for the task; operations outside the 12 most frequent labels are grouped as other for joint strata. Marginal operation distributions retain every detected label.",
        },
        "selected_task_ids": selected_ids,
        "case_preservation": {
            "selected_task_count": len(sampled_records),
            "selected_input_case_count": sum(r["case_input_count"] for r in sampled_records),
            "selected_answer_case_count": sum(r["case_answer_count"] for r in sampled_records),
            "expected_case_count_if_complete": len(sampled_records) * len(CASE_INDICES),
            "input_cases_by_task": {r["task_id"]: full_case_manifest[r["task_id"]] for r in sampled_records},
        },
        "distributions": {
            "instruction_type": {"full_912": dist_for(records, "instruction_type", ["Cell-Level", "Sheet-Level"]), "sampled_200": dist_for(sampled_records, "instruction_type", ["Cell-Level", "Sheet-Level"])},
            "primary_operation": {"full_912": dist_for(records, "primary_operation"), "sampled_200": dist_for(sampled_records, "primary_operation")},
            "workbook_complexity_bin": {"full_912": dist_for(records, "workbook_complexity_bin", ["low", "medium", "high"]), "sampled_200": dist_for(sampled_records, "workbook_complexity_bin", ["low", "medium", "high"])},
            "instruction_complexity_bin": {"full_912": dist_for(records, "instruction_complexity_bin", ["low", "medium", "high"]), "sampled_200": dist_for(sampled_records, "instruction_complexity_bin", ["low", "medium", "high"])},
            "workbook_metric_bins": {},
            "operation_labels": {
                "full_912": {op: {"count": sum(op in r["operations"] for r in records), "proportion": sum(op in r["operations"] for r in records) / len(records)} for op in operation_order},
                "sampled_200": {op: {"count": sum(op in r["operations"] for r in sampled_records), "proportion": sum(op in r["operations"] for r in sampled_records) / len(sampled_records)} for op in operation_order},
            },
        },
        "numeric_metrics": {
            "workbook": {metric: {"full_912": summary([float(r[metric]) for r in records]), "sampled_200": summary([float(r[metric]) for r in sampled_records])} for metric in ["sheet_count_mean", "max_rows_mean", "max_columns_mean", "file_size_bytes_mean"]},
            "instruction": {metric: {"full_912": summary([float(r[metric]) for r in records]), "sampled_200": summary([float(r[metric]) for r in sampled_records])} for metric in ["instruction_word_count", "instruction_char_count", "instruction_sentence_count", "instruction_clause_count"]},
        },
        "joint_strata": {
            "full_912": dist_for(records, "workbook_complexity_bin"),
            "stratum_counts_full": {"|".join(key): len(items) for key, items in sorted(((key, items) for key, items in defaultdict(list).items()), key=repr)},
            "sampled_task_ids": selected_ids,
        },
        "quality_checks": {
            "unique_selected_ids": len(selected_ids) == len(set(selected_ids)),
            "selected_ids_are_in_full_dataset": set(selected_ids).issubset({r["task_id"] for r in records}),
            "all_three_input_cases_present_for_selected": all(r["case_input_count"] == len(CASE_INDICES) for r in sampled_records),
            "not_first_200_in_dataset_order": selected_ids != [r["task_id"] for r in records[:args.sample_size]],
            "workbook_metadata_errors_full": sum(any(c.get("metadata_error") for c in r["cases"]) for r in records),
            "workbook_metadata_errors_sample": sum(any(c.get("metadata_error") for c in r["cases"]) for r in sampled_records),
            "missing_input_cases_full": sum(3 - r["case_input_count"] for r in records),
            "missing_input_cases_sample": sum(3 - r["case_input_count"] for r in sampled_records),
        },
    }
    for metric in ["sheet_count_mean", "max_rows_mean", "max_columns_mean", "file_size_bytes_mean"]:
        _, cuts = quantile_bins([float(r[metric]) for r in records])
        full_labels = labels_from_cuts([float(r[metric]) for r in records], cuts)
        sample_labels = labels_from_cuts([float(r[metric]) for r in sampled_records], cuts)
        report["distributions"]["workbook_metric_bins"][metric] = {
            "bin_boundaries_from_full_912": cuts,
            "full_912": distribution(full_labels, order=["low", "medium", "high"]),
            "sampled_200": distribution(sample_labels, order=["low", "medium", "high"]),
        }
    # Replace the compact placeholder with actual joint-stratum distributions.
    def stratum_key(record: dict[str, Any]) -> str:
        return "|".join(record[field] for field in ("instruction_type", "primary_operation", "workbook_complexity_bin", "instruction_complexity_bin"))
    report["joint_strata"] = {
        "definition": ["instruction_type", "primary_operation", "workbook_complexity_bin", "instruction_complexity_bin"],
        "full_912": distribution([stratum_key(r) for r in records]),
        "sampled_200": distribution([stratum_key(r) for r in sampled_records]),
    }
    (output / "sampling_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(output), "selected_tasks": len(selected_ids), "selected_input_cases": report["case_preservation"]["selected_input_case_count"], "joint_strata": len(report["joint_strata"]["full_912"])}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
